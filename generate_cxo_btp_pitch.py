import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_presentation(output_path="Lumbini_SAP_BTP_Strategic_CXO_Pitch.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # -------------------------------------------------------------
    # Premium Executive Design Tokens
    # -------------------------------------------------------------
    DARK_BG = RGBColor(11, 19, 43)           # #0B132B Deep Executive Navy
    DARK_CARD = RGBColor(21, 31, 60)         # #151F3C Elevated Navy Card
    DARK_CARD_INNER = RGBColor(16, 25, 48)   # #101930 Deeper navy container
    DARK_BORDER = RGBColor(40, 56, 95)       # Subtle midnight border
    
    LIGHT_BG = RGBColor(248, 250, 252)       # #F8FAFC Crisp Slate/Ice White
    CARD_BG = RGBColor(255, 255, 255)        # Pure White Card
    CARD_BG_MUTED = RGBColor(241, 245, 249)  # Light Slate Tint
    CARD_BORDER = RGBColor(226, 232, 240)    # Clean Card Outline
    
    SAP_BLUE = RGBColor(0, 102, 204)         # #0066CC SAP Primary Blue
    SAP_DEEP = RGBColor(10, 40, 95)          # #0A285F Sapphire Blue
    CYAN_ACCENT = RGBColor(0, 180, 216)      # #00B4D8 High-Tech Cyan
    PURPLE_ACCENT = RGBColor(124, 58, 237)   # #7C3AED AI / Data Violet
    EMERALD_GREEN = RGBColor(16, 185, 129)   # #10B981 Verified Emerald
    AMBER_ACCENT = RGBColor(217, 119, 6)     # #D97706 Context / Amber
    
    TEXT_DARK = RGBColor(15, 23, 42)         # #0F172A Executive Dark
    TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B Secondary Slate
    TEXT_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC Primary Light
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184) # Secondary Light
    
    TOTAL_SLIDES = 9

    def add_base_slide(is_dark=False):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG if is_dark else LIGHT_BG
        bg.line.fill.background()
        return slide

    def add_header(slide, title_text, category_text, is_dark=False):
        # Category Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.38), Inches(3.8), Inches(0.26))
        badge.fill.solid()
        badge.fill.fore_color.rgb = DARK_CARD if is_dark else RGBColor(238, 242, 255)
        badge.line.color.rgb = CYAN_ACCENT if is_dark else SAP_BLUE
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.margin_left = Inches(0.12)
        tf_b.margin_top = Inches(0.02)
        p_b = tf_b.paragraphs[0]
        p_b.text = category_text.upper()
        p_b.font.name = "Arial"
        p_b.font.size = Pt(8.5)
        p_b.font.bold = True
        p_b.font.color.rgb = CYAN_ACCENT if is_dark else SAP_BLUE
        
        # Slide Title
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.733), Inches(0.48))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT if is_dark else TEXT_DARK

    def add_footer(slide, slide_num, is_dark=False):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = DARK_BORDER if is_dark else CARD_BORDER
        line.line.fill.background()
        
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.98), Inches(9.5), Inches(0.3))
        tf = tx_box.text_frame
        tf.margin_left = tf.margin_top = 0
        p = tf.paragraphs[0]
        p.text = "Lumbini Elite Solutions & Services  |  Strategic SAP BTP Architecture & Transformation  |  Confidential"
        p.font.name = "Arial"
        p.font.size = Pt(8.5)
        p.font.color.rgb = TEXT_LIGHT_MUTED if is_dark else TEXT_MUTED
        
        tx_num = slide.shapes.add_textbox(Inches(11.533), Inches(6.98), Inches(1.0), Inches(0.3))
        tf_num = tx_num.text_frame
        tf_num.margin_right = tf_num.margin_top = 0
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"{slide_num:02d} / {TOTAL_SLIDES:02d}"
        p_num.font.name = "Arial"
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = CYAN_ACCENT if is_dark else SAP_BLUE

    # =============================================================
    # SLIDE 1: ECOSYSTEM ALIGNMENT & COLLABORATIVE ARCHITECTURE
    # =============================================================
    s1 = add_base_slide(is_dark=True)
    
    top_glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_glow.fill.solid()
    top_glow.fill.fore_color.rgb = CYAN_ACCENT
    top_glow.line.fill.background()
    
    add_header(s1, "Ecosystem Alignment & Collaborative Architecture", "Strategic Architecture & Vision", is_dark=True)
    
    # Executive Banner Statement
    exec_banner = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.22), Inches(11.733), Inches(0.6))
    exec_banner.fill.solid()
    exec_banner.fill.fore_color.rgb = DARK_CARD
    exec_banner.line.color.rgb = CYAN_ACCENT
    exec_banner.line.width = Pt(1)
    tf_eb = exec_banner.text_frame
    tf_eb.margin_top = Inches(0.08)
    p = tf_eb.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "“BTP does not replace the SAP core — it connects, extends and innovates around it.”"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    
    # Visual Architecture Diagram (Top-to-Bottom Flow across enterprise ecosystem)
    # Box 1: Manufacturing Plants & Operations
    b_plant = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.05), Inches(5.7), Inches(1.15))
    b_plant.fill.solid()
    b_plant.fill.fore_color.rgb = DARK_CARD
    b_plant.line.color.rgb = DARK_BORDER
    tf = b_plant.text_frame
    tf.margin_top = Inches(0.1)
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "MANUFACTURING PLANTS & SHOP FLOOR (OPERATIONAL EDGE)"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "Multi-Plant Discrete Operations  |  MES  |  PLC / SCADA  |  Quality Stations  |  PLM"
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_LIGHT_MUTED
    
    # Box 2: Enterprise Ecosystem (Dealers, OEMs, Suppliers)
    b_eco = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(2.05), Inches(5.7), Inches(1.15))
    b_eco.fill.solid()
    b_eco.fill.fore_color.rgb = DARK_CARD
    b_eco.line.color.rgb = DARK_BORDER
    tf = b_eco.text_frame
    tf.margin_top = Inches(0.1)
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "BUSINESS ECOSYSTEM & EXTERNAL TOUCHPOINTS"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "B2B Dealers  |  Tier-1 OEMs  |  Supply Chain Vendors  |  Field Service  |  Customers"
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_LIGHT_MUTED
    
    # Central SAP Core Box (S/4HANA & ECC)
    b_core = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.45), Inches(11.733), Inches(1.2))
    b_core.fill.solid()
    b_core.fill.fore_color.rgb = RGBColor(16, 28, 58)
    b_core.line.color.rgb = SAP_BLUE
    b_core.line.width = Pt(1.5)
    tf = b_core.text_frame
    tf.margin_top = Inches(0.1)
    tf.margin_left = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "EXISTING ENTERPRISE CORE: SAP S/4HANA & SAP ECC (SYSTEMS OF RECORD)"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p = tf.add_paragraph()
    p.text = "Mission-Critical Ledger  |  Standard Sales & Distribution (SD)  |  Materials Management (MM)  |  Production Planning (PP)  |  Quality (QM)  |  Controlling (CO)\nKept Stable, Governed, and Unaltered — Zero Replacement Required"
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_LIGHT
    
    # Central BTP Platform Layer
    b_btp = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.88), Inches(11.733), Inches(1.85))
    b_btp.fill.solid()
    b_btp.fill.fore_color.rgb = DARK_CARD
    b_btp.line.color.rgb = CYAN_ACCENT
    b_btp.line.width = Pt(1.5)
    
    tf = b_btp.text_frame
    tf.margin_top = Inches(0.1)
    tf.margin_left = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = "SAP BUSINESS TECHNOLOGY PLATFORM (SAP BTP) — COLLABORATIVE INNOVATION FABRIC"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    
    # 5 Functional Columns inside BTP
    btp_cols = [
        ("INTEGRATION", "SAP Integration Suite\nEvent Mesh & B2B/EDI\nAPI Lifecycle Gateway"),
        ("DATA FABRIC", "SAP Datasphere\nHANA Cloud & BW\nUnified Semantics"),
        ("BUSINESS AI", "SAP Joule & Copilots\nAgentic Workflows\nContextual LLMs"),
        ("AUTOMATION", "SAP Build Process\nDynamic Rules Engine\nCentral Task Center"),
        ("APP EXTENSION", "SAP CAP (Cloud Foundry)\nSAP RAP (ABAP Cloud)\nFiori Portals & Mobile")
    ]
    for ci, (c_tit, c_desc) in enumerate(btp_cols):
        c_left = Inches(1.05 + ci * 2.25)
        c_sub = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(5.3), Inches(2.1), Inches(1.25))
        c_sub.fill.solid()
        c_sub.fill.fore_color.rgb = DARK_CARD_INNER
        c_sub.line.color.rgb = DARK_BORDER
        tf_c = c_sub.text_frame
        tf_c.margin_top = Inches(0.08)
        tf_c.margin_left = Inches(0.1)
        p = tf_c.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = c_tit
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT
        p = tf_c.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = c_desc
        p.font.name = "Calibri"
        p.font.size = Pt(8.5)
        p.font.color.rgb = TEXT_LIGHT_MUTED
        
    add_footer(s1, 1, is_dark=True)

    # =============================================================
    # SLIDE 2: BTP FIT, CLEAN CORE & AI-READY ENTERPRISE
    # =============================================================
    s2 = add_base_slide(is_dark=False)
    add_header(s2, "Modernizing the Enterprise: Clean Core & AI-Ready", "Foundational Architecture Paradigm", is_dark=False)
    
    # Left Card: Core vs BTP Roles & Clean Core Philosophy
    c_cc = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(5.4))
    c_cc.fill.solid()
    c_cc.fill.fore_color.rgb = CARD_BG
    c_cc.line.color.rgb = CARD_BORDER
    
    # Header tag
    t_cc = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.05), Inches(1.5), Inches(4.0), Inches(0.28))
    t_cc.fill.solid()
    t_cc.fill.fore_color.rgb = RGBColor(238, 242, 255)
    t_cc.line.color.rgb = SAP_BLUE
    p = t_cc.text_frame.paragraphs[0]
    p.text = "THE CLEAN CORE EXTENSION PHILOSOPHY"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    tx_cc = s2.shapes.add_textbox(Inches(1.05), Inches(1.88), Inches(5.2), Inches(4.6))
    tf = tx_cc.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "System of Record vs. System of Innovation"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP
    
    cc_items = [
        ("SAP S/4HANA & ECC as System of Record", "Preserve standard transactional logic, financial ledgers, and inventory integrity without custom modifications inside the core."),
        ("SAP BTP as Innovation Platform", "Provide the dedicated side-by-side runtime for integration, custom portals, automations, and intelligent services."),
        ("Minimize Inside-the-Core Customization", "Eliminate new user-exits, implicit enhancements, and non-standard Z-tables in core transactions."),
        ("Standard APIs & Event-Driven Decoupling", "Communicate exclusively via standard SAP OData/REST APIs and asynchronous business events."),
        ("Reduced Upgrade & Patch Impact", "Allow S/4HANA release updates and patches to execute predictably without breaking custom business apps."),
        ("Enterprise Service Reusability", "Build reusable microservices and workflows accessible across both S/4HANA and ECC environments.")
    ]
    for t, d in cc_items:
        p1 = tf.add_paragraph()
        p1.text = f"- {t}: "
        p1.font.name = "Arial"
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p2 = tf.add_paragraph()
        p2.text = d
        p2.font.name = "Calibri"
        p2.font.size = Pt(9)
        p2.font.color.rgb = TEXT_MUTED

    # Right Card: The AI-Ready Enterprise Foundation
    c_ai = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.3), Inches(5.7), Inches(5.4))
    c_ai.fill.solid()
    c_ai.fill.fore_color.rgb = CARD_BG
    c_ai.line.color.rgb = RGBColor(224, 231, 255)
    
    t_ai = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.083), Inches(1.5), Inches(4.3), Inches(0.28))
    t_ai.fill.solid()
    t_ai.fill.fore_color.rgb = RGBColor(245, 243, 255)
    t_ai.line.color.rgb = PURPLE_ACCENT
    p = t_ai.text_frame.paragraphs[0]
    p.text = "THE AI-READY ENTERPRISE FOUNDATION"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    
    tx_ai = s2.shapes.add_textbox(Inches(7.083), Inches(1.88), Inches(5.2), Inches(4.6))
    tf = tx_ai.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "Business AI Powered by Trusted Context & Governance"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    
    p = tf.add_paragraph()
    p.text = "Enterprise AI succeeds only when grounded in operational truth. SAP BTP serves as the enabling technological foundation connecting data, processes, and AI models:"
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    
    # 4-Stage Horizontal Flow Diagram
    # Business Data -> Business Context -> AI -> Action
    flow_steps = [
        ("Business Data", "Clean master data, transactional records, and telemetry"),
        ("Business Context", "Process rules, semantic models, and organization hierarchies"),
        ("Enterprise AI", "Joule assistants, generative models, and agentic workflows"),
        ("Decisive Action", "Automated ERP transactions, recommendations, and execution")
    ]
    for fi, (f_step, f_sub) in enumerate(flow_steps):
        f_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.083), Inches(2.9 + fi * 0.8), Inches(5.2), Inches(0.68))
        f_box.fill.solid()
        f_box.fill.fore_color.rgb = CARD_BG_MUTED
        f_box.line.color.rgb = RGBColor(209, 213, 219)
        tf_f = f_box.text_frame
        tf_f.margin_top = Inches(0.06)
        tf_f.margin_left = Inches(0.15)
        p = tf_f.paragraphs[0]
        p.text = f"Step {fi+1}: {f_step.upper()}"
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = PURPLE_ACCENT if fi >= 2 else SAP_DEEP
        p = tf_f.add_paragraph()
        p.text = f_sub
        p.font.name = "Calibri"
        p.font.size = Pt(8.5)
        p.font.color.rgb = TEXT_DARK

    add_footer(s2, 2, is_dark=False)

    # =============================================================
    # SLIDE 3: USE CASE 1 - BUSINESS DATA CLOUD
    # =============================================================
    s3 = add_base_slide(is_dark=False)
    add_header(s3, "From Fragmented Data to Trusted Enterprise Intelligence", "Use Case 1: SAP Business Data Cloud", is_dark=False)
    
    # Architecture Flow Banner at Top
    # SAP S/4HANA + ECC + MES + Other Systems -> SAP Business Data Cloud -> Trusted Data -> SAC / Planning / Decisions
    af_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.22), Inches(11.733), Inches(0.8))
    af_box.fill.solid()
    af_box.fill.fore_color.rgb = RGBColor(238, 242, 255)
    af_box.line.color.rgb = SAP_BLUE
    af_box.line.width = Pt(1)
    tf_af = af_box.text_frame
    tf_af.margin_top = Inches(0.08)
    tf_af.margin_left = Inches(0.2)
    p = tf_af.paragraphs[0]
    p.text = "INTEGRATED ENTERPRISE DATA ARCHITECTURE FLOW"
    p.font.name = "Arial"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP
    p = tf_af.add_paragraph()
    p.text = "SAP S/4HANA & ECC + MES / Shop Floor + External Systems  —►  SAP Business Data Cloud (Federation & Governance)  —►  Trusted Semantics  —►  SAP Analytics Cloud & Executive Decisions"
    p.font.name = "Calibri"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    # Left Column: Five Visual Capability Groups
    c_dc = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.15), Inches(5.7), Inches(4.65))
    c_dc.fill.solid()
    c_dc.fill.fore_color.rgb = CARD_BG
    c_dc.line.color.rgb = CARD_BORDER
    
    tx_dc = s3.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(5.3), Inches(4.4))
    tf = tx_dc.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "Five Core Data Pillars"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP
    
    data_pillars = [
        ("DATA FOUNDATION", "SAP HANA Cloud & SAP Datasphere", "In-memory multi-model engine and business data fabric connecting distributed manufacturing data without replication."),
        ("ANALYTICS & PLANNING", "SAP Analytics Cloud (SAC)", "Unified BI, predictive forecasting, agentic analytics, and enterprise planning across finance and plant operations."),
        ("DATA WAREHOUSING", "SAP Business Warehouse (BW / BW/4HANA)", "Proven enterprise data consolidation model preserving historic manufacturing trends and statutory reporting."),
        ("AI / DATA SCIENCE", "SAP Databricks Open Data Ecosystem", "Zero-copy bi-directional data sharing between SAP Datasphere and Databricks lakehouse for advanced industrial ML."),
        ("DATA GOVERNANCE", "SAP Master Data Governance (MDG)", "Centralized governance, duplicate detection, and golden-record stewardship for materials, BOMs, and vendor entities.")
    ]
    for cat, prod, desc in data_pillars:
        p1 = tf.add_paragraph()
        p1.text = f"- {cat} [{prod}]: "
        p1.font.name = "Arial"
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = SAP_BLUE
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_MUTED

    # Right Column: Industry Experience References (Carefully Framed)
    # Ref 1: Healthium Medtech
    c_h1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(2.15), Inches(5.7), Inches(2.25))
    c_h1.fill.solid()
    c_h1.fill.fore_color.rgb = CARD_BG
    c_h1.line.color.rgb = CARD_BORDER
    
    t1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(2.3), Inches(3.2), Inches(0.24))
    t1.fill.solid()
    t1.fill.fore_color.rgb = RGBColor(241, 245, 249)
    t1.line.color.rgb = RGBColor(148, 163, 184)
    p = t1.text_frame.paragraphs[0]
    p.text = "RELEVANT EXPERIENCE REFERENCE"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    tx1 = s3.shapes.add_textbox(Inches(7.05), Inches(2.6), Inches(5.3), Inches(1.7))
    tf1 = tx1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = 0
    p = tf1.paragraphs[0]
    p.text = "Healthium Medtech Pvt. Ltd."
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf1.add_paragraph()
    p.text = "Domain Context: Enterprise data integration, analytics, and multi-hub manufacturing reporting.\n"
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.color.rgb = SAP_BLUE
    p = tf1.add_paragraph()
    p.text = "Illustrative BTP Application Scenario:\nConsolidating disparate operational data sources into a unified reporting layer to provide management with reliable visibility across plant operations, production output, and commercial performance."
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    # Ref 2: Hical Technologies
    c_h2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(4.55), Inches(5.7), Inches(2.25))
    c_h2.fill.solid()
    c_h2.fill.fore_color.rgb = CARD_BG
    c_h2.line.color.rgb = CARD_BORDER
    
    t2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(4.7), Inches(3.2), Inches(0.24))
    t2.fill.solid()
    t2.fill.fore_color.rgb = RGBColor(241, 245, 249)
    t2.line.color.rgb = RGBColor(148, 163, 184)
    p = t2.text_frame.paragraphs[0]
    p.text = "RELEVANT EXPERIENCE REFERENCE"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    tx2 = s3.shapes.add_textbox(Inches(7.05), Inches(5.0), Inches(5.3), Inches(1.7))
    tf2 = tx2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = 0
    p = tf2.paragraphs[0]
    p.text = "Hical Technologies Private Limited"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf2.add_paragraph()
    p.text = "Domain Context: Master data governance, reporting accuracy, and enterprise visibility.\n"
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.color.rgb = SAP_BLUE
    p = tf2.add_paragraph()
    p.text = "Illustrative BTP Application Scenario:\nAddressing material and supplier master data inconsistencies across multi-plant manufacturing units, establishing verified data quality to accelerate executive decision-making."
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    add_footer(s3, 3, is_dark=False)

    # =============================================================
    # SLIDE 4: USE CASE 2 - SAP BUSINESS AI
    # =============================================================
    s4 = add_base_slide(is_dark=False)
    add_header(s4, "From Information Access to Intelligent Business Assistance", "Use Case 2: SAP Business AI", is_dark=False)
    
    # Evolution Flow Banner
    # User asks -> AI understands business context -> AI retrieves information -> AI recommends/assists -> Business executes
    ai_flow = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.22), Inches(11.733), Inches(0.75))
    ai_flow.fill.solid()
    ai_flow.fill.fore_color.rgb = RGBColor(245, 243, 255)
    ai_flow.line.color.rgb = PURPLE_ACCENT
    ai_flow.line.width = Pt(1)
    tf_f = ai_flow.text_frame
    tf_f.margin_top = Inches(0.06)
    tf_f.margin_left = Inches(0.2)
    p = tf_f.paragraphs[0]
    p.text = "THE ENTERPRISE AI ENGAGEMENT CYCLE"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    p = tf_f.add_paragraph()
    p.text = "User Asks  —►  AI Understands Business Context  —►  AI Retrieves Enterprise Information  —►  AI Recommends / Assists  —►  Business Executes"
    p.font.name = "Calibri"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    # Left Card: SAP Business AI Capabilities & Enterprise Use Scenarios
    c_aicap = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), Inches(5.7), Inches(4.7))
    c_aicap.fill.solid()
    c_aicap.fill.fore_color.rgb = CARD_BG
    c_aicap.line.color.rgb = CARD_BORDER
    
    tx_aicap = s4.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(4.5))
    tf = tx_aicap.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "SAP Joule Suite & Practical Scenarios"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    
    ai_caps = [
        ("Joule Core & Joule Work", "Contextual conversational copilot built directly into SAP applications for natural language queries and transactional navigation."),
        ("Joule Assistants & Specialized Agents", "Autonomous agents that monitor production schedules, flag component delays, and recommend rescheduling actions."),
        ("Joule Studio", "Low-code environment to fine-tune enterprise models, create custom business prompts, and ground generative AI in specific manufacturing context."),
        ("Joule for Consultants & Developers", "AI-assisted code development, automated test generation, and BTP API discovery to accelerate project execution."),
        ("Practical Enterprise Use Scenarios", "Manufacturing plant insights, procurement assistance, sales order exception analysis, warranty claims adjudication, and technical knowledge retrieval.")
    ]
    for tit, desc in ai_caps:
        p1 = tf.add_paragraph()
        p1.text = f"- {tit}: "
        p1.font.name = "Arial"
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_MUTED

    # Right Card: Experience References (SAIL & Taurani Holdings)
    # Ref 1: SAIL
    c_s1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(2.1), Inches(5.7), Inches(2.28))
    c_s1.fill.solid()
    c_s1.fill.fore_color.rgb = CARD_BG
    c_s1.line.color.rgb = CARD_BORDER
    
    ts1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(2.25), Inches(3.2), Inches(0.24))
    ts1.fill.solid()
    ts1.fill.fore_color.rgb = RGBColor(241, 245, 249)
    ts1.line.color.rgb = RGBColor(148, 163, 184)
    p = ts1.text_frame.paragraphs[0]
    p.text = "RELEVANT EXPERIENCE REFERENCE"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    tx_s1 = s4.shapes.add_textbox(Inches(7.05), Inches(2.55), Inches(5.3), Inches(1.75))
    tf1 = tx_s1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = 0
    p = tf1.paragraphs[0]
    p.text = "Steel Authority of India Limited (SAIL)"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf1.add_paragraph()
    p.text = "Domain Context: Heavy-industry manufacturing providing relevant context for operational insights and enterprise decision support.\n"
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.color.rgb = PURPLE_ACCENT
    p = tf1.add_paragraph()
    p.text = "Illustrative AI Opportunity Scenario:\nEquipping engineering and operations teams with intelligent query interfaces that synthesize complex plant telemetry, maintenance history, and operational parameters to resolve bottlenecks rapidly."
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    # Ref 2: Taurani Holdings
    c_s2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(4.52), Inches(5.7), Inches(2.28))
    c_s2.fill.solid()
    c_s2.fill.fore_color.rgb = CARD_BG
    c_s2.line.color.rgb = CARD_BORDER
    
    ts2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(4.67), Inches(3.2), Inches(0.24))
    ts2.fill.solid()
    ts2.fill.fore_color.rgb = RGBColor(241, 245, 249)
    ts2.line.color.rgb = RGBColor(148, 163, 184)
    p = ts2.text_frame.paragraphs[0]
    p.text = "RELEVANT EXPERIENCE REFERENCE"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    tx_s2 = s4.shapes.add_textbox(Inches(7.05), Inches(4.97), Inches(5.3), Inches(1.75))
    tf2 = tx_s2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = 0
    p = tf2.paragraphs[0]
    p.text = "Taurani Holdings"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf2.add_paragraph()
    p.text = "Domain Context: Multi-division enterprise context applicable to intelligent analytics and commercial decision support.\n"
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.color.rgb = PURPLE_ACCENT
    p = tf2.add_paragraph()
    p.text = "Illustrative AI Opportunity Scenario:\nApplying predictive modeling and automated variance detection to highlight commodity and component cost trends, giving commercial leaders early warning to adjust procurement and pricing strategies."
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    add_footer(s4, 4, is_dark=False)

    # =============================================================
    # SLIDE 5: USE CASE 3 - SAP INTEGRATION SUITE
    # =============================================================
    s5 = add_base_slide(is_dark=False)
    add_header(s5, "Connect Every Plant, Application, Partner and Process", "Use Case 3: SAP Integration Suite", is_dark=False)
    
    # Left Card: Integration Suite Architectural Capabilities
    c_int = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(5.4))
    c_int.fill.solid()
    c_int.fill.fore_color.rgb = CARD_BG
    c_int.line.color.rgb = CARD_BORDER
    
    t_int = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.05), Inches(1.5), Inches(4.0), Inches(0.26))
    t_int.fill.solid()
    t_int.fill.fore_color.rgb = RGBColor(238, 242, 255)
    t_int.line.color.rgb = SAP_BLUE
    p = t_int.text_frame.paragraphs[0]
    p.text = "ENTERPRISE CONNECTIVITY CAPABILITIES"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    tx_int = s5.shapes.add_textbox(Inches(1.05), Inches(1.85), Inches(5.2), Inches(4.7))
    tf = tx_int.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "From Point-to-Point to Governed Fabric"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP
    
    int_caps = [
        ("Application Integration", "Connect heterogeneous SAP and non-SAP enterprise systems using pre-built integration adapters and packaged content."),
        ("API Lifecycle Management", "Design, govern, secure, and monitor enterprise APIs with rate limiting, OAuth authentication, and catalog discoverability."),
        ("Event-Driven Architecture", "Leverage SAP Event Mesh to decouple transactional events (e.g. goods receipt, quality check) asynchronously without burdening the core."),
        ("B2B & Partner Integration", "Standardize electronic communication with suppliers, Tier-1 automotive OEMs, and distributor networks across EDI and modern REST protocols."),
        ("Enterprise Integration Strategy", "Migrate away from fragile point-to-point custom integrations toward a centralized, governed, and highly observable integration backbone."),
        ("Agentic AI Integration", "Provide secure, governed API access layers enabling intelligent agents to interact with transactional ERP systems safely.")
    ]
    for tit, desc in int_caps:
        p1 = tf.add_paragraph()
        p1.text = f"- {tit}: "
        p1.font.name = "Arial"
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_MUTED

    # Right Column: Verified References (TEPL & TSAT)
    # Ref 1: TATA Electronics (TEPL)
    c_t1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.3), Inches(5.7), Inches(2.6))
    c_t1.fill.solid()
    c_t1.fill.fore_color.rgb = CARD_BG
    c_t1.line.color.rgb = CARD_BORDER
    
    tt1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(1.45), Inches(3.2), Inches(0.24))
    tt1.fill.solid()
    tt1.fill.fore_color.rgb = RGBColor(236, 253, 245)
    tt1.line.color.rgb = EMERALD_GREEN
    p = tt1.text_frame.paragraphs[0]
    p.text = "PROVEN EXPERIENCE REFERENCE"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    
    tx_t1 = s5.shapes.add_textbox(Inches(7.05), Inches(1.75), Inches(5.3), Inches(2.1))
    tf1 = tx_t1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = 0
    p = tf1.paragraphs[0]
    p.text = "TATA Electronics Pvt. Ltd. (TEPL)"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf1.add_paragraph()
    p.text = "Verified Context: SAP / MES Integration + Manufacturing Data + Analytics"
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    p = tf1.add_paragraph()
    p.text = "Architectural Flow Demonstrated:\nMES / Manufacturing Systems  —►  Integration Layer  —►  SAP S/4HANA & ECC  —►  Business Analytics (HANA / Tableau)\n\nRelevance to Target Enterprise: Demonstrates concrete Lumbini capability in bridging plant-floor execution systems to enterprise ERP for unified production and yield visibility."
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    # Ref 2: Tata Semiconductor Assembly and Test (TSAT)
    c_t2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(4.1), Inches(5.7), Inches(2.6))
    c_t2.fill.solid()
    c_t2.fill.fore_color.rgb = CARD_BG
    c_t2.line.color.rgb = CARD_BORDER
    
    tt2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(4.25), Inches(3.2), Inches(0.24))
    tt2.fill.solid()
    tt2.fill.fore_color.rgb = RGBColor(236, 253, 245)
    tt2.line.color.rgb = EMERALD_GREEN
    p = tt2.text_frame.paragraphs[0]
    p.text = "PROVEN EXPERIENCE REFERENCE"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    
    tx_t2 = s5.shapes.add_textbox(Inches(7.05), Inches(4.55), Inches(5.3), Inches(2.1))
    tf2 = tx_t2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = 0
    p = tf2.paragraphs[0]
    p.text = "Tata Semiconductor Assembly & Test (TSAT)"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf2.add_paragraph()
    p.text = "Verified Context: Manufacturing-System Connectivity & Enterprise Integration"
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    p = tf2.add_paragraph()
    p.text = "Architectural Flow Demonstrated:\nPlant Systems  —►  Secure Integration Gateway  —►  SAP Core  —►  Operational Analytics\n\nRelevance to Target Enterprise: Demonstrates hands-on capability in architecting high-reliability, secure data exchanges between plant-floor environments and central ERP."
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    add_footer(s5, 5, is_dark=False)

    # =============================================================
    # SLIDE 6: USE CASE 4 - WORKFLOW & AUTOMATION
    # =============================================================
    s6 = add_base_slide(is_dark=False)
    add_header(s6, "Automate the Work Around SAP — Not Just the Transaction", "Use Case 4: Workflow, Automation & Signavio", is_dark=False)
    
    # Manufacturing Process Flow Banner
    # Trigger -> Business Rule -> Workflow -> Human Approval -> SAP Transaction -> Notification -> Audit Trail
    wf_banner = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.22), Inches(11.733), Inches(0.8))
    wf_banner.fill.solid()
    wf_banner.fill.fore_color.rgb = RGBColor(238, 242, 255)
    wf_banner.line.color.rgb = SAP_BLUE
    wf_banner.line.width = Pt(1)
    tf_w = wf_banner.text_frame
    tf_w.margin_top = Inches(0.08)
    tf_w.margin_left = Inches(0.2)
    p = tf_w.paragraphs[0]
    p.text = "TYPICAL MANUFACTURING AUTOMATION FLOW ON SAP BTP"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP
    p = tf_w.add_paragraph()
    p.text = "Trigger (Plant / Sensor / Order)  —►  Business Rule Evaluated  —►  Orchestrated Workflow  —►  Human Approval (Mobile/Task Center)  —►  SAP Core Posting  —►  Audit Trail & Notification"
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    # Left Card: SAP Build Process Automation & Signavio Capabilities
    c_wf = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.15), Inches(5.7), Inches(4.65))
    c_wf.fill.solid()
    c_wf.fill.fore_color.rgb = CARD_BG
    c_wf.line.color.rgb = CARD_BORDER
    
    tx_wf = s6.shapes.add_textbox(Inches(1.05), Inches(2.28), Inches(5.2), Inches(4.4))
    tf = tx_wf.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "Automation & Process Intelligence Suite"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP
    
    wf_items = [
        ("SAP Build Process Automation", "Unified low-code/no-code environment combining workflow orchestration, business rules, form builders, and robotic process automation (RPA)."),
        ("Dynamic Business Rules Engine", "Decouples decision logic (discount thresholds, engineering tolerance checks, warranty validation) from core ABAP code, editable directly by process owners."),
        ("SAP Central Task Center", "Consolidates approval inboxes across SAP S/4HANA, ECC, Ariba, and custom applications into one unified mobile and desktop inbox."),
        ("SAP Signavio Process Intelligence", "Continuously mines operational event logs to reveal hidden bottlenecks, rework loops in manufacturing, and process deviations before automating."),
        ("Audit-Ready Traceability", "Ensures end-to-end electronic records, digital signatures, and immutable logging across plant and supply chain processes.")
    ]
    for tit, desc in wf_items:
        p1 = tf.add_paragraph()
        p1.text = f"- {tit}: "
        p1.font.name = "Arial"
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_MUTED

    # Right Card: Reference Case (Dr. Reddy's Laboratories Ltd.)
    c_dr = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(2.15), Inches(5.7), Inches(4.65))
    c_dr.fill.solid()
    c_dr.fill.fore_color.rgb = CARD_BG
    c_dr.line.color.rgb = CARD_BORDER
    
    t_dr = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(2.35), Inches(3.2), Inches(0.24))
    t_dr.fill.solid()
    t_dr.fill.fore_color.rgb = RGBColor(241, 245, 249)
    t_dr.line.color.rgb = RGBColor(148, 163, 184)
    p = t_dr.text_frame.paragraphs[0]
    p.text = "RELEVANT EXPERIENCE REFERENCE"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    tx_dr = s6.shapes.add_textbox(Inches(7.05), Inches(2.7), Inches(5.3), Inches(3.9))
    tf = tx_dr.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "Dr. Reddy's Laboratories Ltd."
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf.add_paragraph()
    p.text = "Domain Context: Relevant regulated-industry experience providing a strong reference point for workflow, approval, compliance, and process automation opportunities.\n"
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    dr_scenarios = [
        ("Multi-Tier Plant Approvals", "Managing complex delegation of authority and multi-stage sign-offs across distributed manufacturing sites without manual delays."),
        ("Compliance & Traceability", "Replacing paper-based checklists and manual logs with automated, audit-ready digital workflows integrated into ERP."),
        ("Exception Routing", "Automatically detecting and routing quality holds, material variance alerts, and non-conformance records to authorized specialists."),
        ("Cross-Functional Coordination", "Synchronizing handoffs between production, quality assurance, regulatory affairs, and inventory teams seamlessly.")
    ]
    p = tf.add_paragraph()
    p.text = "Illustrative BTP Application Opportunities in Discrete Manufacturing:"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    for st, sd in dr_scenarios:
        p1 = tf.add_paragraph()
        p1.text = f"* {st}: "
        p1.font.name = "Arial"
        p1.font.size = Pt(8.8)
        p1.font.bold = True
        p1.font.color.rgb = SAP_DEEP
        p2 = tf.add_paragraph()
        p2.text = sd
        p2.font.name = "Calibri"
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_MUTED

    add_footer(s6, 6, is_dark=False)

    # =============================================================
    # SLIDE 7: USE CASE 5 - CUSTOM APPLICATION DEVELOPMENT
    # =============================================================
    s7 = add_base_slide(is_dark=False)
    add_header(s7, "Build New Business Experiences Without Overloading the SAP Core", "Use Case 5: Custom App Development (CAP vs. RAP)", is_dark=False)
    
    # Left Card: Direct Comparison Table (CAP vs RAP)
    c_cmp = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(5.4))
    c_cmp.fill.solid()
    c_cmp.fill.fore_color.rgb = CARD_BG
    c_cmp.line.color.rgb = CARD_BORDER
    
    tx_cmp = s7.shapes.add_textbox(Inches(1.0), Inches(1.45), Inches(5.3), Inches(0.45))
    tf = tx_cmp.text_frame
    p = tf.paragraphs[0]
    p.text = "Extension Architecture Frameworks: CAP vs. RAP"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP
    
    # Comparison Table
    table_shape = s7.shapes.add_table(6, 2, Inches(1.0), Inches(1.95), Inches(5.3), Inches(4.5))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.65)
    tbl.columns[1].width = Inches(2.65)
    
    headers = ["SAP CAP (Cloud-Native)", "SAP RAP (ABAP-Centric)"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = SAP_DEEP if j == 0 else PURPLE_ACCENT
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = h
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        
    cmp_rows = [
        ("Technology Stack", "Node.js / Java on Cloud Foundry & Kyma", "ABAP Cloud / RESTful ABAP on S/4HANA & BTP ABAP"),
        ("Architectural Model", "Side-by-side, fully decoupled cloud extensions", "In-app extensions & tightly coupled ledger logic"),
        ("Runtime Environment", "BTP Cloud Foundry / Kubernetes (Kyma)", "S/4HANA Cloud / BTP ABAP Environment"),
        ("Primary Use Cases", "Dealer portals, customer mobile apps, multi-tenant SaaS", "Transactional business objects, core extensions, heavy calculations"),
        ("Core Benefit", "Zero ERP release lock-in; open web developer ecosystem", "Direct CDS data model reuse & native ABAP performance")
    ]
    for row_idx, (r_dim, c_cap, c_rap) in enumerate(cmp_rows, start=1):
        cell_cap = tbl.cell(row_idx, 0)
        cell_cap.fill.solid()
        cell_cap.fill.fore_color.rgb = RGBColor(248, 250, 252)
        p = cell_cap.text_frame.paragraphs[0]
        p.text = f"{r_dim}\n"
        p.font.name = "Arial"
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = SAP_DEEP
        p2 = cell_cap.text_frame.add_paragraph()
        p2.text = c_cap
        p2.font.name = "Calibri"
        p2.font.size = Pt(8)
        p2.font.color.rgb = TEXT_DARK
        
        cell_rap = tbl.cell(row_idx, 1)
        cell_rap.fill.solid()
        cell_rap.fill.fore_color.rgb = RGBColor(248, 250, 252)
        p = cell_rap.text_frame.paragraphs[0]
        p.text = f"{r_dim}\n"
        p.font.name = "Arial"
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = PURPLE_ACCENT
        p2 = cell_rap.text_frame.add_paragraph()
        p2.text = c_rap
        p2.font.name = "Calibri"
        p2.font.size = Pt(8)
        p2.font.color.rgb = TEXT_DARK

    # Right Column: Concrete Proof Point (L&T DEMANS) + Context (IB Group)
    # Ref 1: L&T Construction & Mining Machinery (DEMANS)
    c_lt = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.3), Inches(5.7), Inches(2.9))
    c_lt.fill.solid()
    c_lt.fill.fore_color.rgb = CARD_BG
    c_lt.line.color.rgb = CARD_BORDER
    
    t_lt = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(1.45), Inches(3.2), Inches(0.24))
    t_lt.fill.solid()
    t_lt.fill.fore_color.rgb = RGBColor(236, 253, 245)
    t_lt.line.color.rgb = EMERALD_GREEN
    p = t_lt.text_frame.paragraphs[0]
    p.text = "PROVEN ENTERPRISE PROOF POINT"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    
    tx_lt = s7.shapes.add_textbox(Inches(7.05), Inches(1.75), Inches(5.3), Inches(2.35))
    tf = tx_lt.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "L&T Construction & Mining Machinery: DEMANS"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf.add_paragraph()
    p.text = "Dealer Management System (DEMANS) — Custom Application Development"
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    p = tf.add_paragraph()
    p.text = "Architectural Flow:\nDealer  —►  DEMANS Custom Portal  —►  Integration Layer  —►  SAP Core  —►  Sales / Inventory / Orders / Business Processes\n\nLumbini Value Demonstrated:\nArchitected an enterprise-grade dealer management application that unifies dealer spare parts ordering, inventory visibility, and sales workflows connected directly to SAP backend without customizing standard core code."
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    # Ref 2: Indian Agro and Food Industries Limited (IB Group)
    c_ib = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(4.35), Inches(5.7), Inches(2.35))
    c_ib.fill.solid()
    c_ib.fill.fore_color.rgb = CARD_BG
    c_ib.line.color.rgb = CARD_BORDER
    
    t_ib = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(4.5), Inches(3.2), Inches(0.24))
    t_ib.fill.solid()
    t_ib.fill.fore_color.rgb = RGBColor(241, 245, 249)
    t_ib.line.color.rgb = RGBColor(148, 163, 184)
    p = t_ib.text_frame.paragraphs[0]
    p.text = "RELEVANT EXPERIENCE REFERENCE"
    p.font.name = "Arial"
    p.font.size = Pt(7.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    tx_ib = s7.shapes.add_textbox(Inches(7.05), Inches(4.78), Inches(5.3), Inches(1.85))
    tf = tx_ib.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "Indian Agro & Food Industries Limited (IB Group)"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p = tf.add_paragraph()
    p.text = "Domain Context: Multi-plant agro/food industry experience applicable to enterprise process digitization, operational visibility, and business integration."
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    p = tf.add_paragraph()
    p.text = "Illustrative BTP Opportunity Scenario:\nExtending ERP with responsive, side-by-side web and mobile apps to handle multi-plant batch tracking, dispatch schedules, and inventory transfers across geographically dispersed sites."
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_MUTED

    add_footer(s7, 7, is_dark=False)

    # =============================================================
    # SLIDE 8: LUMBINI ELITIA PRODUCT SUITE
    # =============================================================
    s8 = add_base_slide(is_dark=False)
    add_header(s8, "Beyond Standard SAP: Lumbini Elitia Accelerators", "Proprietary Enterprise Microservices Portfolio", is_dark=False)
    
    # Positioning Header Banner
    el_pos = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.22), Inches(11.733), Inches(0.72))
    el_pos.fill.solid()
    el_pos.fill.fore_color.rgb = RGBColor(238, 242, 255)
    el_pos.line.color.rgb = SAP_BLUE
    el_pos.line.width = Pt(1)
    tf_ep = el_pos.text_frame
    tf_ep.margin_top = Inches(0.06)
    tf_ep.margin_left = Inches(0.2)
    p = tf_ep.paragraphs[0]
    p.text = "COMPLEMENTING SAP — NOT REPLACING IT"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP
    p = tf_ep.add_paragraph()
    p.text = "A portfolio of proprietary enterprise microservices and digital accelerators designed to complement SAP landscapes.\nArchitecture: Elitia Microservices  ↕  APIs / Integration  ↕  SAP BTP Integration Layer  ↕  SAP S/4HANA & ECC"
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    # 5 Product Cards across the slide
    # 11.733 inches total width. 5 cards with 0.18 spacing -> width ~ 2.2 inches
    elitia_modules = [
        ("Supplier Collaboration Portal",
         "Digital vendor onboarding, purchase order acknowledgment, advance shipment notices (ASN), digital gate passes, and delivery scheduling.",
         "SAP MM / S/4HANA Sourcing",
         CYAN_ACCENT),
         
        ("Data Governance Suite",
         "Enterprise-wide data quality management, cross-plant deduplication rules, automated approval workflows, and audit-ready data tracking.",
         "SAP MDG / Datasphere",
         SAP_DEEP),
         
        ("Master Data Validator",
         "Automated pre-posting validation engine that checks material specs, customer records, and tax configurations against business rules before ERP posting.",
         "Pre-Posting Rule Engine",
         PURPLE_ACCENT),
         
        ("Customer Experience Portal",
         "Modern B2B/B2C self-service portal for dealers and end-customers providing real-time order status, pricing, catalog search, and warranty registration.",
         "SAP SD / Digital Commerce",
         EMERALD_GREEN),
         
        ("Workforce Safety Management",
         "Comprehensive plant safety management covering incident logging, hazardous work permits, EHS compliance tracking, and safety training audits.",
         "SAP EHS / Plant Maintenance",
         AMBER_ACCENT)
    ]
    
    card_w = Inches(2.2)
    spacing = Inches(0.18)
    for idx, (em_title, em_desc, em_conn, em_color) in enumerate(elitia_modules):
        c_x = Inches(0.8 + idx * (2.2 + 0.18))
        card_shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, Inches(2.05), Inches(2.2), Inches(4.7))
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = CARD_BG
        card_shape.line.color.rgb = CARD_BORDER
        
        # Color accent stripe
        c_bar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x + Inches(0.15), Inches(2.2), Inches(1.9), Inches(0.06))
        c_bar.fill.solid()
        c_bar.fill.fore_color.rgb = em_color
        c_bar.line.fill.background()
        
        # Pill Tag
        ptag = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.15), Inches(2.38), Inches(1.9), Inches(0.24))
        ptag.fill.solid()
        ptag.fill.fore_color.rgb = CARD_BG_MUTED
        ptag.line.fill.background()
        p = ptag.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = f"ACCELERATOR {idx+1}"
        p.font.name = "Arial"
        p.font.size = Pt(7.5)
        p.font.bold = True
        p.font.color.rgb = em_color
        
        # Title and Description
        tx_e = s8.shapes.add_textbox(c_x + Inches(0.15), Inches(2.72), Inches(1.9), Inches(3.8))
        tf_e = tx_e.text_frame
        tf_e.word_wrap = True
        tf_e.margin_left = tf_e.margin_top = 0
        p = tf_e.paragraphs[0]
        p.text = em_title
        p.font.name = "Arial"
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        p = tf_e.add_paragraph()
        p.text = "\n" + em_desc
        p.font.name = "Calibri"
        p.font.size = Pt(8.8)
        p.font.color.rgb = TEXT_MUTED
        
        p = tf_e.add_paragraph()
        p.text = f"\nIntegration Layer:\n{em_conn}"
        p.font.name = "Arial"
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = em_color

    add_footer(s8, 8, is_dark=False)

    # =============================================================
    # SLIDE 9: LUMBINI ADVANTAGE & NEXT STEP
    # =============================================================
    s9 = add_base_slide(is_dark=True)
    add_header(s9, "Clean Core. Connected Enterprise. Accelerated Innovation.", "Delivery Methodology & Next Steps", is_dark=True)
    
    # 4-Stage Delivery Methodology Across Top Half
    methodology_stages = [
        ("01 — DISCOVER", "Identify Landscape Friction", "Analyze core customizations, point-to-point integration pain points, data fragmentation, and manual plant approval processes."),
        ("02 — PRIORITIZE", "Strategic Value Matrix", "Evaluate candidate opportunities based on:\nBusiness Value × Complexity × Reusability × Time-to-Value."),
        ("03 — PILOT", "Single High-Impact Use Case", "Implement one focused lighthouse pilot (e.g. process automation, dealer portal, unified data view) to validate value rapidly."),
        ("04 — SCALE", "Reusable Foundation", "Standardize enterprise APIs, integration patterns, semantic data models, and deploy relevant Elitia accelerators across divisions.")
    ]
    for mi, (m_num, m_sub, m_desc) in enumerate(methodology_stages):
        m_x = Inches(0.8 + mi * 2.98)
        m_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, m_x, Inches(1.3), Inches(2.8), Inches(1.9))
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = DARK_CARD
        m_card.line.color.rgb = DARK_BORDER
        
        tx_m = s9.shapes.add_textbox(m_x + Inches(0.15), Inches(1.4), Inches(2.5), Inches(1.7))
        tf = tx_m.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0
        p = tf.paragraphs[0]
        p.text = m_num
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT
        p = tf.add_paragraph()
        p.text = m_sub
        p.font.name = "Arial"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p = tf.add_paragraph()
        p.text = m_desc
        p.font.name = "Calibri"
        p.font.size = Pt(8.5)
        p.font.color.rgb = TEXT_LIGHT_MUTED

    # Middle Section: 4 Lumbini Differentiation Pillars
    diff_pillars = [
        ("SAP EXPERTISE", "Deep architectural mastery across S/4HANA, ECC, and SAP BTP.", SAP_BLUE),
        ("MANUFACTURING FOCUS", "Extensive experience across plants, MES, quality, supply chain, and dealers.", CYAN_ACCENT),
        ("DIGITAL ENGINEERING", "Modern cloud-native engineering across APIs, integrations, data, and analytics.", PURPLE_ACCENT),
        ("PROPRIETARY ACCELERATORS", "Proven Elitia microservices and pre-built connectors to shorten time-to-value.", EMERALD_GREEN)
    ]
    for di, (d_tit, d_desc, d_col) in enumerate(diff_pillars):
        d_x = Inches(0.8 + di * 2.98)
        d_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, d_x, Inches(3.35), Inches(2.8), Inches(1.3))
        d_card.fill.solid()
        d_card.fill.fore_color.rgb = DARK_CARD_INNER
        d_card.line.color.rgb = d_col
        
        tx_d = s9.shapes.add_textbox(d_x + Inches(0.15), Inches(3.45), Inches(2.5), Inches(1.1))
        tf = tx_d.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0
        p = tf.paragraphs[0]
        p.text = d_tit
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = d_col
        p = tf.add_paragraph()
        p.text = d_desc
        p.font.name = "Calibri"
        p.font.size = Pt(8.5)
        p.font.color.rgb = TEXT_LIGHT_MUTED

    # Bottom Callout Box: 2-Week Joint Discovery & Closing Mantra
    cta_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.8), Inches(11.733), Inches(1.95))
    cta_box.fill.solid()
    cta_box.fill.fore_color.rgb = RGBColor(16, 32, 68)
    cta_box.line.color.rgb = CYAN_ACCENT
    cta_box.line.width = Pt(1.5)
    
    tx_cta = s9.shapes.add_textbox(Inches(1.05), Inches(4.92), Inches(11.2), Inches(1.75))
    tf = tx_cta.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "PROPOSED IMMEDIATE NEXT STEP: 2-WEEK JOINT DISCOVERY WORKSHOP"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    
    p = tf.add_paragraph()
    p.text = "“Let's identify the first three BTP opportunities where the enterprise can create measurable business value without disrupting the SAP core.”"
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    
    p = tf.add_paragraph()
    p.text = "• Week 1 (Discover → Assess → Map): Review existing customization hotspots, plant connectivity bottlenecks, and dealer integration flows.\n• Week 2 (Prioritize → Design → Recommend): Formulate the BTP target architecture, evaluate business impact, and deliver a concrete Pilot Implementation Roadmap.\nDeliverable: BTP Opportunity Roadmap + Priority Use Cases + Target Architecture + Pilot Recommendation."
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_LIGHT_MUTED
    
    p = tf.add_paragraph()
    p.text = "\n“Keep the Core Clean. Connect the Enterprise. Empower the Business. Build for Continuous Innovation.”"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    add_footer(s9, 9, is_dark=True)

    # -------------------------------------------------------------
    # Save Presentation
    # -------------------------------------------------------------
    prs.save(output_path)
    print(f"Successfully created presentation at: {output_path}")

if __name__ == "__main__":
    build_presentation()
