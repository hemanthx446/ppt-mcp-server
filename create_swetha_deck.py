import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_presentation(output_path="Swetha_GRC_Speech_Deck.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # -------------------------------------------------------------
    # Design Tokens & Palettes
    # -------------------------------------------------------------
    DARK_BG = RGBColor(18, 18, 18)           # #121212 Dark Charcoal
    DARK_CARD = RGBColor(28, 28, 30)         # #1C1C1E
    DARK_BORDER = RGBColor(45, 45, 48)       # #2D2D30
    
    LIGHT_BG = RGBColor(248, 249, 250)       # #F8F9FA Ice White
    CARD_BG = RGBColor(255, 255, 255)        # Clean White
    BORDER_GRAY = RGBColor(226, 232, 240)    # Soft Slate Border
    
    PRIMARY_RED = RGBColor(227, 6, 19)       # #E30613 Brand Accent Red
    ACCENT_RED_BG = RGBColor(254, 242, 242)  # Subtle Red Tint for key pills
    
    TEXT_DARK = RGBColor(26, 32, 44)         # #1A202C
    TEXT_MUTED = RGBColor(100, 116, 139)     # Slate Muted
    TEXT_LIGHT = RGBColor(248, 250, 252)     # Off-white
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184)
    
    TOTAL_SLIDES = 4

    def create_base_slide(is_dark=False):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG if is_dark else LIGHT_BG
        bg.line.fill.background()
        return slide

    def add_header(slide, title_text, subtitle_text):
        # Red Accent Bar
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.55), Inches(0.08), Inches(0.55))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = PRIMARY_RED
        stripe.line.fill.background()
        
        # Title
        tx_box = slide.shapes.add_textbox(Inches(1.05), Inches(0.48), Inches(11.4), Inches(0.45))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        # Subtitle
        tx_sub = slide.shapes.add_textbox(Inches(1.05), Inches(0.98), Inches(11.4), Inches(0.35))
        tf_sub = tx_sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = "Calibri"
        p_sub.font.size = Pt(12)
        p_sub.font.italic = True
        p_sub.font.color.rgb = TEXT_MUTED

    def add_footer(slide, slide_num):
        # Footer Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_GRAY
        line.line.fill.background()
        
        # Left Text
        tx_l = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(8.5), Inches(0.3))
        tf_l = tx_l.text_frame
        tf_l.margin_left = tf_l.margin_top = tf_l.margin_bottom = tf_l.margin_right = 0
        p_l = tf_l.paragraphs[0]
        p_l.text = "Lumbini Elite  |  Executive Briefing with Swetha, CPO of SAP GRC"
        p_l.font.name = "Calibri"
        p_l.font.size = Pt(9)
        p_l.font.color.rgb = TEXT_MUTED
        
        # Right Page Number
        tx_r = slide.shapes.add_textbox(Inches(11.033), Inches(6.95), Inches(1.5), Inches(0.3))
        tf_r = tx_r.text_frame
        tf_r.margin_left = tf_r.margin_top = tf_r.margin_bottom = tf_r.margin_right = 0
        p_r = tf_r.paragraphs[0]
        p_r.text = f"Slide {slide_num} of {TOTAL_SLIDES}"
        p_r.font.name = "Calibri"
        p_r.font.size = Pt(9)
        p_r.font.color.rgb = TEXT_MUTED
        p_r.alignment = PP_ALIGN.RIGHT

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme)
    # =========================================================================
    s1 = create_base_slide(is_dark=True)
    
    # Top Tag Pill
    tag_bg = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(3.2), Inches(0.38))
    tag_bg.fill.solid()
    tag_bg.fill.fore_color.rgb = DARK_CARD
    tag_bg.line.color.rgb = PRIMARY_RED
    tag_bg.line.width = Pt(1)
    
    tx_tag = s1.shapes.add_textbox(Inches(1.1), Inches(1.26), Inches(3.0), Inches(0.3))
    tf_tag = tx_tag.text_frame
    tf_tag.margin_top = tf_tag.margin_left = tf_tag.margin_bottom = tf_tag.margin_right = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "EXECUTIVE STRATEGIC BRIEFING"
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = PRIMARY_RED
    
    # Main Title
    tx_title = s1.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(11.3), Inches(1.8))
    tf_title = tx_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_top = tf_title.margin_left = tf_title.margin_bottom = tf_title.margin_right = 0
    p1 = tf_title.paragraphs[0]
    p1.text = "Co-Innovating the Future of SAP GRC"
    p1.font.name = "Arial"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.space_after = Pt(8)
    
    p2 = tf_title.add_paragraph()
    p2.text = "Lumbini Elite & SAP CPO Office Collaboration"
    p2.font.name = "Arial"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY_RED
    
    # Red Horizontal Divider Accent
    div_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(3.85), Inches(2.2), Inches(0.04))
    div_line.fill.solid()
    div_line.fill.fore_color.rgb = PRIMARY_RED
    div_line.line.fill.background()
    
    # Subtitle / Objective
    tx_sub = s1.shapes.add_textbox(Inches(1.0), Inches(4.1), Inches(11.3), Inches(0.9))
    tf_sub = tx_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_top = tf_sub.margin_left = tf_sub.margin_bottom = tf_sub.margin_right = 0
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "A targeted 30-minute dialogue on SAP Clean Core alignment, accelerated Cloud IAG adoption, and proven enterprise execution in regulated global environments."
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_LIGHT_MUTED
    
    # Bottom 3 Highlight Badges (Cards)
    pills = [
        ("Clean Core Alignment", "Zero-modification compliant governance models"),
        ("Cloud IAG Acceleration", "Frictionless hybrid migration for on-prem clients"),
        ("Field-Tested Scale", "20k+ user scale, 90% cycle time reduction")
    ]
    card_w = 3.5
    gap = 0.4
    start_x = 1.0
    for i, (p_title, p_desc) in enumerate(pills):
        x = start_x + i * (card_w + gap)
        c_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.3), Inches(card_w), Inches(1.3))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = DARK_CARD
        c_shape.line.color.rgb = DARK_BORDER
        c_shape.line.width = Pt(1)
        
        tx_c = s1.shapes.add_textbox(Inches(x + 0.2), Inches(5.45), Inches(card_w - 0.4), Inches(1.0))
        tf_c = tx_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_top = tf_c.margin_left = tf_c.margin_bottom = tf_c.margin_right = 0
        
        pc1 = tf_c.paragraphs[0]
        pc1.text = p_title
        pc1.font.name = "Arial"
        pc1.font.size = Pt(13)
        pc1.font.bold = True
        pc1.font.color.rgb = TEXT_LIGHT
        pc1.space_after = Pt(4)
        
        pc2 = tf_c.add_paragraph()
        pc2.text = p_desc
        pc2.font.name = "Calibri"
        pc2.font.size = Pt(10.5)
        pc2.font.color.rgb = TEXT_LIGHT_MUTED

    # =========================================================================
    # SLIDE 2: Strategic Vision Alignment (Light Theme)
    # =========================================================================
    s2 = create_base_slide(is_dark=False)
    add_header(
        s2,
        "Strategic Vision Alignment with SAP Clean Core",
        "Bridging enterprise compliance to SAP Cloud Identity Services, BTP, and AI-driven risk frameworks."
    )
    
    # 3 Strategic Pillar Cards
    pillars = [
        {
            "num": "01",
            "title": "SAP Clean Core Alignment",
            "subtitle": "Decoupled Extensibility on BTP",
            "bullets": [
                "Eliminate custom core modifications in S/4HANA transformations.",
                "Enforce native standard APIs for provisioning and compliance.",
                "Maintain continuous upgradeability with zero downtime risk."
            ]
        },
        {
            "num": "02",
            "title": "Cloud IAG & Identity Services",
            "subtitle": "Hybrid Identity Bridge",
            "bullets": [
                "Frictionless bridge from on-prem Access Control to Cloud IAG.",
                "Native integration with SAP Cloud Identity Services (IAS/IPS).",
                "Full cross-application identity governance & SCIM connectors."
            ]
        },
        {
            "num": "03",
            "title": "AI-Driven Risk Frameworks",
            "subtitle": "Predictive & Continuous Compliance",
            "bullets": [
                "AI-assisted SoD conflict analysis and proactive simulation.",
                "Contextual anomaly detection across firefighter / emergency roles.",
                "Autonomous remediation recommendations before audit cycles."
            ]
        }
    ]
    
    col_w = 3.65
    col_gap = 0.39
    left_start = 0.8
    card_y = 1.55
    card_h = 4.2
    
    for i, p_info in enumerate(pillars):
        x = left_start + i * (col_w + col_gap)
        
        # Container Shape
        c_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(card_y), Inches(col_w), Inches(card_h))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = BORDER_GRAY
        c_shape.line.width = Pt(1)
        
        # Red Header Bar on top of each card
        top_bar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(card_y), Inches(col_w), Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = PRIMARY_RED
        top_bar.line.fill.background()
        
        # Card Text Box
        tx_box = s2.shapes.add_textbox(Inches(x + 0.25), Inches(card_y + 0.25), Inches(col_w - 0.5), Inches(card_h - 0.5))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_bottom = tf.margin_right = 0
        
        # Pillar Number
        p_num = tf.paragraphs[0]
        p_num.text = p_info["num"]
        p_num.font.name = "Arial"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = PRIMARY_RED
        p_num.space_after = Pt(3)
        
        # Pillar Title
        p_title = tf.add_paragraph()
        p_title.text = p_info["title"]
        p_title.font.name = "Arial"
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.space_after = Pt(2)
        
        # Pillar Subtitle
        p_subt = tf.add_paragraph()
        p_subt.text = p_info["subtitle"]
        p_subt.font.name = "Calibri"
        p_subt.font.size = Pt(10.5)
        p_subt.font.bold = True
        p_subt.font.color.rgb = PRIMARY_RED
        p_subt.space_after = Pt(14)
        
        # Bullets
        for b_text in p_info["bullets"]:
            pb = tf.add_paragraph()
            pb.text = "•  " + b_text
            pb.font.name = "Calibri"
            pb.font.size = Pt(10.5)
            pb.font.color.rgb = TEXT_DARK
            pb.space_after = Pt(8)
            
    # Bottom Takeaway Callout Card
    takeaway_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.75))
    takeaway_shape.fill.solid()
    takeaway_shape.fill.fore_color.rgb = ACCENT_RED_BG
    takeaway_shape.line.color.rgb = PRIMARY_RED
    takeaway_shape.line.width = Pt(1)
    
    tx_tk = s2.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.55))
    tf_tk = tx_tk.text_frame
    tf_tk.word_wrap = True
    tf_tk.margin_top = tf_tk.margin_left = tf_tk.margin_bottom = tf_tk.margin_right = 0
    ptk = tf_tk.paragraphs[0]
    ptk.text = "Key Takeaway for CPO Office: "
    ptk.font.name = "Arial"
    ptk.font.size = Pt(11)
    ptk.font.bold = True
    ptk.font.color.rgb = PRIMARY_RED
    
    run_tk = ptk.add_run()
    run_tk.text = "Lumbini Elite actively accelerates customer transitions from legacy on-prem custom code to SAP's standard Cloud IAG platform, protecting the Clean Core imperative."
    run_tk.font.name = "Calibri"
    run_tk.font.size = Pt(11)
    run_tk.font.bold = False
    run_tk.font.color.rgb = TEXT_DARK
    
    add_footer(s2, 2)

    # =========================================================================
    # SLIDE 3: Proven Field Execution & Scale (Light Theme)
    # =========================================================================
    s3 = create_base_slide(is_dark=False)
    add_header(
        s3,
        "Proven Field Execution & Scale",
        "Validating enterprise reliability, zero-touch governance, and verified business outcomes."
    )
    
    # 3 High-Impact Stat Containers
    stats = [
        {
            "val": "90%",
            "label": "Provisioning Cycle Reduction",
            "desc": "Transformed 5-day manual ticketing bottlenecks into automated, policy-compliant 4-hour approvals."
        },
        {
            "val": "20,000+",
            "label": "Enterprise Users Managed",
            "desc": "Scaled across highly regulated multi-national SAP estates with 100% audit accuracy and zero findings."
        },
        {
            "val": "< 2 Hours",
            "label": "Emergency Access SLA",
            "desc": "Frictionless firefighter role elevation with automated continuous logging and post-session risk analysis."
        }
    ]
    
    stat_w = 3.65
    stat_gap = 0.39
    stat_y = 1.55
    stat_h = 2.1
    
    for i, s_item in enumerate(stats):
        x = left_start + i * (stat_w + stat_gap)
        
        c_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(stat_y), Inches(stat_w), Inches(stat_h))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = BORDER_GRAY
        c_shape.line.width = Pt(1)
        
        tx_box = s3.shapes.add_textbox(Inches(x + 0.25), Inches(stat_y + 0.2), Inches(stat_w - 0.5), Inches(stat_h - 0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_bottom = tf.margin_right = 0
        
        p_val = tf.paragraphs[0]
        p_val.text = s_item["val"]
        p_val.font.name = "Arial"
        p_val.font.size = Pt(28)
        p_val.font.bold = True
        p_val.font.color.rgb = PRIMARY_RED
        p_val.space_after = Pt(2)
        
        p_lbl = tf.add_paragraph()
        p_lbl.text = s_item["label"]
        p_lbl.font.name = "Arial"
        p_lbl.font.size = Pt(11)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = TEXT_DARK
        p_lbl.space_after = Pt(4)
        
        p_desc = tf.add_paragraph()
        p_desc.text = s_item["desc"]
        p_desc.font.name = "Calibri"
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_MUTED
        
    # Bottom Half: Two Comparative / Detailed Execution Containers
    bot_y = 3.9
    bot_h = 2.8
    bot_w = 5.67
    bot_gap = 0.39
    
    # Bottom Left Container: Field Validation Architecture
    c_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(bot_y), Inches(bot_w), Inches(bot_h))
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = CARD_BG
    c_left.line.color.rgb = BORDER_GRAY
    c_left.line.width = Pt(1)
    
    tx_bl = s3.shapes.add_textbox(Inches(1.05), Inches(bot_y + 0.2), Inches(bot_w - 0.5), Inches(bot_h - 0.4))
    tf_bl = tx_bl.text_frame
    tf_bl.word_wrap = True
    tf_bl.margin_top = tf_bl.margin_left = tf_bl.margin_bottom = tf_bl.margin_right = 0
    
    p_bl_t = tf_bl.paragraphs[0]
    p_bl_t.text = "Field-Tested Execution Pillars"
    p_bl_t.font.name = "Arial"
    p_bl_t.font.size = Pt(13)
    p_bl_t.font.bold = True
    p_bl_t.font.color.rgb = TEXT_DARK
    p_bl_t.space_after = Pt(8)
    
    pillars_exec = [
        "Hybrid Interoperability: Seamless bridge connecting on-prem SAP ECC / S/4HANA with BTP cloud services.",
        "Zero-Touch Lifecycle: Automated HR-driven onboarding, role mapping, and instantaneous offboarding revokes.",
        "Audit-Proof Integrity: Automated logging satisfying rigorous external SOX, ISO-27001, and GDPR compliance."
    ]
    for pe in pillars_exec:
        p_pe = tf_bl.add_paragraph()
        p_pe.text = "•  " + pe
        p_pe.font.name = "Calibri"
        p_pe.font.size = Pt(10)
        p_pe.font.color.rgb = TEXT_DARK
        p_pe.space_after = Pt(6)
        
    # Bottom Right Container: SAP Ecosystem Value
    c_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + bot_w + bot_gap), Inches(bot_y), Inches(bot_w), Inches(bot_h))
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = CARD_BG
    c_right.line.color.rgb = BORDER_GRAY
    c_right.line.width = Pt(1)
    
    tx_br = s3.shapes.add_textbox(Inches(0.8 + bot_w + bot_gap + 0.25), Inches(bot_y + 0.2), Inches(bot_w - 0.5), Inches(bot_h - 0.4))
    tf_br = tx_br.text_frame
    tf_br.word_wrap = True
    tf_br.margin_top = tf_br.margin_left = tf_br.margin_bottom = tf_br.margin_right = 0
    
    p_br_t = tf_br.paragraphs[0]
    p_br_t.text = "Value Proposition for SAP Product Strategy"
    p_br_t.font.name = "Arial"
    p_br_t.font.size = Pt(13)
    p_br_t.font.bold = True
    p_br_t.font.color.rgb = TEXT_DARK
    p_br_t.space_after = Pt(8)
    
    value_sap = [
        "Unlocks IAG Conversions: Solves complex enterprise objections by proving rapid, non-disruptive migration paths.",
        "Reduces Customer Churn: Eliminates governance friction during S/4HANA Cloud and RISE with SAP journeys.",
        "Telemetry & Insights: Real-world operational feedback on API edge cases and enterprise scalability."
    ]
    for vs in value_sap:
        p_vs = tf_br.add_paragraph()
        p_vs.text = "•  " + vs
        p_vs.font.name = "Calibri"
        p_vs.font.size = Pt(10)
        p_vs.font.color.rgb = TEXT_DARK
        p_vs.space_after = Pt(6)

    add_footer(s3, 3)

    # =========================================================================
    # SLIDE 4: Partnership & Pilot Proposal (Light Theme)
    # =========================================================================
    s4 = create_base_slide(is_dark=False)
    add_header(
        s4,
        "Partnership & Co-Innovation Proposal",
        "Establishing a structured collaboration framework between Lumbini Elite and the SAP GRC CPO Office."
    )
    
    # 3 Workstream Cards
    workstreams = [
        {
            "step": "PHASE 1",
            "title": "Joint Architecture Workshop",
            "timeline": "Weeks 1 – 2",
            "items": [
                "Roadmap deep dive on upcoming Cloud IAG and BTP governance features.",
                "Review joint reference architecture for Clean Core identity lifecycle.",
                "Define API extensibility standards and enterprise telemetry hooks."
            ]
        },
        {
            "step": "PHASE 2",
            "title": "Lighthouse Customer Pilot",
            "timeline": "Weeks 3 – 6",
            "items": [
                "Select 1–2 enterprise clients modernizing from on-prem GRC to IAG.",
                "Deploy co-innovated Clean Core baseline with zero modifications.",
                "Validate 90%+ cycle time reduction and capture telemetry."
            ]
        },
        {
            "step": "PHASE 3",
            "title": "Continuous Success Feedback Loops",
            "timeline": "Ongoing Track",
            "items": [
                "Direct bi-directional feedback loop into SAP GRC Product Management.",
                "Joint customer reference stories, whitepapers, and SAP Sapphire showcase.",
                "Co-developed standard packages accelerating RISE with SAP migrations."
            ]
        }
    ]
    
    ws_w = 3.65
    ws_gap = 0.39
    ws_y = 1.55
    ws_h = 3.7
    
    for i, ws in enumerate(workstreams):
        x = left_start + i * (ws_w + ws_gap)
        
        c_shape = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(ws_y), Inches(ws_w), Inches(ws_h))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = BORDER_GRAY
        c_shape.line.width = Pt(1)
        
        # Pill for Phase/Timeline
        pill = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.25), Inches(ws_y + 0.25), Inches(1.2), Inches(0.3))
        pill.fill.solid()
        pill.fill.fore_color.rgb = ACCENT_RED_BG
        pill.line.color.rgb = PRIMARY_RED
        pill.line.width = Pt(1)
        
        tx_pill = s4.shapes.add_textbox(Inches(x + 0.25), Inches(ws_y + 0.28), Inches(1.2), Inches(0.25))
        tf_pill = tx_pill.text_frame
        tf_pill.margin_top = tf_pill.margin_left = tf_pill.margin_bottom = tf_pill.margin_right = 0
        p_p = tf_pill.paragraphs[0]
        p_p.text = ws["step"]
        p_p.font.name = "Arial"
        p_p.font.size = Pt(9)
        p_p.font.bold = True
        p_p.font.color.rgb = PRIMARY_RED
        p_p.alignment = PP_ALIGN.CENTER
        
        # Timeline text
        tx_time = s4.shapes.add_textbox(Inches(x + 1.55), Inches(ws_y + 0.3), Inches(1.8), Inches(0.25))
        tf_time = tx_time.text_frame
        tf_time.margin_top = tf_time.margin_left = tf_time.margin_bottom = tf_time.margin_right = 0
        p_t = tf_time.paragraphs[0]
        p_t.text = ws["timeline"]
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(9.5)
        p_t.font.italic = True
        p_t.font.color.rgb = TEXT_MUTED
        
        # Card Body Text Box
        tx_box = s4.shapes.add_textbox(Inches(x + 0.25), Inches(ws_y + 0.75), Inches(ws_w - 0.5), Inches(ws_h - 0.9))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_bottom = tf.margin_right = 0
        
        p_title = tf.paragraphs[0]
        p_title.text = ws["title"]
        p_title.font.name = "Arial"
        p_title.font.size = Pt(13.5)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.space_after = Pt(12)
        
        for item in ws["items"]:
            pi = tf.add_paragraph()
            pi.text = "•  " + item
            pi.font.name = "Calibri"
            pi.font.size = Pt(10)
            pi.font.color.rgb = TEXT_DARK
            pi.space_after = Pt(7)
            
    # Bottom Callout / Proposed Immediate Next Step
    act_shape = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.5), Inches(11.733), Inches(1.15))
    act_shape.fill.solid()
    act_shape.fill.fore_color.rgb = CARD_BG
    act_shape.line.color.rgb = PRIMARY_RED
    act_shape.line.width = Pt(1.5)
    
    tx_act = s4.shapes.add_textbox(Inches(1.1), Inches(5.62), Inches(11.1), Inches(0.9))
    tf_act = tx_act.text_frame
    tf_act.word_wrap = True
    tf_act.margin_top = tf_act.margin_left = tf_act.margin_bottom = tf_act.margin_right = 0
    
    p_act1 = tf_act.paragraphs[0]
    p_act1.text = "Proposed Immediate Next Step: 45-Minute Working Session"
    p_act1.font.name = "Arial"
    p_act1.font.size = Pt(12)
    p_act1.font.bold = True
    p_act1.font.color.rgb = PRIMARY_RED
    p_act1.space_after = Pt(3)
    
    p_act2 = tf_act.add_paragraph()
    p_act2.text = "Align our Lead Architect with the SAP GRC Product Architecture team to scope the Phase 1 workshop agenda and establish technical criteria for the customer co-innovation pilot."
    p_act2.font.name = "Calibri"
    p_act2.font.size = Pt(11)
    p_act2.font.color.rgb = TEXT_DARK
    
    add_footer(s4, 4)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_presentation()
