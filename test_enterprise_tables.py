import os
import sys
import unittest
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

from design_system.table_engine import (
    TableArchetype,
    TableRowItem,
    EnterpriseTableSpec,
    TableDataClassifier,
    EnterpriseTableComposer,
    EnterpriseTableFactory
)
from ppt_mcp_server import (
    list_table_archetypes,
    render_enterprise_table
)


class TestEnterpriseTables(unittest.TestCase):
    OUTPUT_DECK = "test_all_12_tables_deck.pptx"
    MCP_OUTPUT_DECK = "test_mcp_table_deck.pptx"

    ALL_ARCHETYPES = [
        TableArchetype.ARCHITECTURE_COMPARISON,
        TableArchetype.REQUIREMENTS,
        TableArchetype.INTERFACE_CATALOGUE,
        TableArchetype.KPI_DEFINITIONS,
        TableArchetype.RISKS,
        TableArchetype.ASSUMPTIONS,
        TableArchetype.COMMERCIAL_PROPOSAL,
        TableArchetype.IMPLEMENTATION_SCOPE,
        TableArchetype.RESPONSIBILITY_MATRIX,
        TableArchetype.ROADMAP,
        TableArchetype.BUSINESS_BENEFITS,
        TableArchetype.DATA_MAPPINGS
    ]

    def tearDown(self):
        for p in [self.OUTPUT_DECK, self.MCP_OUTPUT_DECK]:
            if os.path.exists(p):
                try:
                    os.remove(p)
                except OSError:
                    pass

    def test_01_all_12_archetypes_spec_generation(self):
        """Verify specs for all 12 canonical enterprise table disciplines."""
        for arch in self.ALL_ARCHETYPES:
            spec = EnterpriseTableFactory.get_table_spec(arch, client_name="Boeing Aerospace")
            self.assertEqual(spec.archetype, arch)
            self.assertGreaterEqual(len(spec.headers), 4, f"Headers too short for {arch.name}")
            self.assertGreaterEqual(len(spec.rows), 4, f"Rows too few for {arch.name}")
            self.assertIsNotNone(spec.title)
            self.assertIsNotNone(spec.subtitle)

    def test_02_render_all_12_tables_deck(self):
        """Render a 12-slide presentation covering all 12 archetypes and audit cells, merges, and typography."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)

        for i, arch in enumerate(self.ALL_ARCHETYPES):
            spec = EnterpriseTableFactory.get_table_spec(arch, client_name="Lockheed Martin")
            spec.is_dark = (i % 2 == 1)
            slide = EnterpriseTableComposer.compose_and_render(prs, spec)
            self.assertIsNotNone(slide)

        self.assertEqual(len(prs.slides), 12)
        prs.save(self.OUTPUT_DECK)
        self.assertTrue(os.path.exists(self.OUTPUT_DECK))

        # Reopen and inspect table contents
        saved_prs = Presentation(self.OUTPUT_DECK)
        total_text_runs = 0
        total_tables = 0
        merged_section_rows_found = 0
        right_aligned_numeric_cells = 0
        center_aligned_code_cells = 0

        for slide_idx, slide in enumerate(saved_prs.slides):
            has_table = False
            for shape in slide.shapes:
                if shape.has_table:
                    has_table = True
                    total_tables += 1
                    tbl = shape.table

                    # Check each cell
                    for r_idx in range(len(tbl.rows)):
                        first_cell = tbl.cell(r_idx, 0)
                        if first_cell.span_width > 1:
                            merged_section_rows_found += 1

                        for c_idx in range(len(tbl.columns)):
                            cell = tbl.cell(r_idx, c_idx)
                            tf = cell.text_frame
                            for p in tf.paragraphs:
                                if p.alignment == PP_ALIGN.RIGHT:
                                    right_aligned_numeric_cells += 1
                                elif p.alignment == PP_ALIGN.CENTER:
                                    center_aligned_code_cells += 1

                                for r in p.runs:
                                    if r.font and r.font.name:
                                        total_text_runs += 1
                                        self.assertTrue(
                                            r.font.name.startswith("Inter"),
                                            f"Non-Inter font '{r.font.name}' found on slide {slide_idx}: '{r.text}'"
                                        )

            self.assertTrue(has_table, f"Slide {slide_idx} must contain a table shape")

        self.assertEqual(total_tables, 12, "Must contain exactly 12 tables")
        self.assertGreater(merged_section_rows_found, 0, "Must contain merged section header rows")
        self.assertGreater(right_aligned_numeric_cells, 10, "Must contain right-aligned numeric/financial cells")
        self.assertGreater(center_aligned_code_cells, 10, "Must contain center-aligned code/status cells")
        self.assertGreater(total_text_runs, 200, "Must audit over 200 text runs for typography")

    def test_03_mcp_tools_tables(self):
        """Verify list_table_archetypes and render_enterprise_table MCP tools."""
        list_res = list_table_archetypes()
        self.assertIn("archetypes", list_res)
        self.assertEqual(list_res["total_archetypes"], 13)

        # Test render tool with commercial proposal
        res = render_enterprise_table(
            path=self.MCP_OUTPUT_DECK,
            archetype="COMMERCIAL_PROPOSAL",
            client_name="Raytheon Defense",
            is_dark=False
        )
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["archetype"], "COMMERCIAL_PROPOSAL")
        self.assertEqual(res["total_slides"], 1)
        self.assertTrue(os.path.exists(self.MCP_OUTPUT_DECK))

        # Test append custom table
        custom_headers = ["COMPONENT", "ENVIRONMENT", "FAILOVER SLA", "MONTHLY COST"]
        custom_rows = [
            ["SECTION: 1. CORE CLUSTER", "", "", ""],
            ["HANA Cloud DB", "Multi-AZ AWS", "< 5 Minutes", "$12,400"],
            ["BTP Integration", "Global High Availability", "< 30 Seconds", "$4,800"]
        ]
        res2 = render_enterprise_table(
            path=self.MCP_OUTPUT_DECK,
            archetype="CUSTOM_TABLE",
            title="High-Availability Infrastructure Spec",
            client_name="Raytheon Defense",
            is_dark=True,
            custom_headers=custom_headers,
            custom_rows=custom_rows,
            callout_note="Disaster recovery automated via AWS Route53 health checks."
        )
        self.assertEqual(res2["status"], "success")
        self.assertEqual(res2["total_slides"], 2)


if __name__ == "__main__":
    unittest.main()
