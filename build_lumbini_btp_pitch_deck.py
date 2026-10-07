import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def create_deck(output_path="Lumbini_BTP_Enterprise_Strategic_Pitch.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # -------------------------------------------------------------
    # Palette & Design System Tokens
    # -------------------------------------------------------------
    DARK_BG = RGBColor(11, 19, 43)           # #0B132B Deep Executive Navy
    DARK_CARD = RGBColor(21, 31, 60)         # #151F3C Navy Card
    DARK_BORDER = RGBColor(40, 56, 95)       # Subtle card outline
    
    LIGHT_BG = RGBColor(248, 250, 252)       # #F8FAFC Slate Ice White
    CARD_BG = RGBColor(255, 255, 255)        # Pure White Card
    CARD_BORDER = RGBColor(226, 232, 240)    # Soft slate border
    CARD_BG_MUTED = RGBColor(241, 245, 249)  # Light Slate Pill
    
    SAP_BLUE = RGBColor(0, 102, 204)         # #0066CC SAP Enterprise Blue
    SAP_DEEP_BLUE = RGBColor(10, 40, 95)     # #0A285F Deep Sapphire
    CYAN_ACCENT = RGBColor(0, 180, 216)      # #00B4D8 Tech Cyan
    EMERALD_GREEN = RGBColor(16, 185, 129)   # #10B981 Success / Metric Green
    AMBER_ACCENT = RGBColor(217, 119, 6)     # #D97706 Warning / Legacy Amber
    PURPLE_ACCENT = RGBColor(124, 58, 237)   # #7C3AED AI & Innovation Purple
    
    TEXT_DARK = RGBColor(15, 23, 42)         # #0F172A Primary Dark
    TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B Secondary Slate
    TEXT_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC Primary Light
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184) # Secondary Light
    
    TOTAL_SLIDES = 9

    # -------------------------------------------------------------
    # Helper Functions
    # -------------------------------------------------------------
    def add_base_slide(is_dark=False):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG if is_dark else LIGHT_BG
        bg.line.fill.background()
        return slide

    def add_header(slide, title_text, category_text, is_dark=False):
        # Category Tag Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(3.6), Inches(0.28))
        badge.fill.solid()
        badge.fill.fore_color.rgb = DARK_CARD if is_dark else RGBColor(238, 242, 255)
        badge.line.color.rgb = CYAN_ACCENT if is_dark else SAP_BLUE
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0.1)
        tf_b.margin_top = Inches(0.02)
        p_b = tf_b.paragraphs[0]
        p_b.text = category_text.upper()
        p_b.font.name = "Arial"
        p_b.font.size = Pt(8.5)
        p_b.font.bold = True
        p_b.font.color.rgb = CYAN_ACCENT if is_dark else SAP_BLUE
        
        # Slide Title
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.48))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT if is_dark else TEXT_DARK

    def add_footer(slide, slide_num, is_dark=False):
        # Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = DARK_BORDER if is_dark else CARD_BORDER
        line.line.fill.background()
        
        # Left Branding Text
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.98), Inches(8.5), Inches(0.3))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0
        p = tf.paragraphs[0]
        p.text = "Lumbini Elite  |  SAP BTP Strategic Architecture & Innovation Pitch  |  Confidential"
        p.font.name = "Arial"
        p.font.size = Pt(8.5)
        p.font.color.rgb = TEXT_LIGHT_MUTED if is_dark else TEXT_MUTED
        
        # Right Slide Number
        tx_num = slide.shapes.add_textbox(Inches(11.533), Inches(6.98), Inches(1.0), Inches(0.3))
        tf_num = tx_num.text_frame
        tf_num.word_wrap = True
        tf_num.margin_right = tf_num.margin_top = 0
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"{slide_num:02d} / {TOTAL_SLIDES:02d}"
        p_num.font.name = "Arial"
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = CYAN_ACCENT if is_dark else SAP_BLUE

    # -------------------------------------------------------------
    # SLIDE 1: Title & Strategic Vision Hook
    # -------------------------------------------------------------
    s1 = add_base_slide(is_dark=True)
    
    top_glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_glow.fill.solid()
    top_glow.fill.fore_color.rgb = CYAN_ACCENT
    top_glow.line.fill.background()
    
    b1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.1), Inches(4.5), Inches(0.36))
    b1.fill.solid()
    b1.fill.fore_color.rgb = DARK_CARD
    b1.line.color.rgb = CYAN_ACCENT
    b1.line.width = Pt(1.2)
    tf1 = b1.text_frame
    tf1.margin_top = Inches(0.04)
    p = tf1.paragraphs[0]
    p.text = "STRATEGIC SAP BTP ARCHITECTURE & MODERNIZATION"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    
    tx = s1.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(11.333), Inches(1.5))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = "Ecosystem Alignment &\nCollaborative Architecture"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    
    tx_sub = s1.shapes.add_textbox(Inches(1.0), Inches(3.3), Inches(11.333), Inches(0.8))
    tf_sub = tx_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = 0
    p = tf_sub.paragraphs[0]
    p.text = "Empowering Multi-Division Discrete Manufacturing & Consumer Appliance Enterprises\nKeeping the S/4HANA & ECC Core Clean While Accelerating Edge Innovation via SAP BTP"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(186, 215, 233)
    
    pillars = [
        ("NON-DISRUPTIVE COLLABORATION", "Wraps around legacy ECC & modern S/4HANA instances seamlessly, bridging operational silos without touching core transactional code.", CYAN_ACCENT),
        ("CLEAN CORE RESILIENCE", "Enforces zero-modification extensibility, eliminating custom code bloat, cutting upgrade regressions, and maintaining 100% cloud readiness.", SAP_BLUE),
        ("AI-DRIVEN EDGE SPEED", "Powers intelligent shop floor workflows, dealer ecosystems, field mobility, and unified costing through enterprise-ready microservices.", PURPLE_ACCENT)
    ]
    
    for i, (p_title, p_desc, p_color) in enumerate(pillars):
        c_left = Inches(1.0 + i * 3.84)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(4.35), Inches(3.64), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD
        card.line.color.rgb = DARK_BORDER
        card.line.width = Pt(1)
        
        c_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left + Inches(0.2), Inches(4.55), Inches(0.4), Inches(0.06))
        c_bar.fill.solid()
        c_bar.fill.fore_color.rgb = p_color
        c_bar.line.fill.background()
        
        tx_c = s1.shapes.add_textbox(c_left + Inches(0.2), Inches(4.75), Inches(3.24), Inches(1.6))
        tf_c = tx_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = 0
        p_t = tf_c.paragraphs[0]
        p_t.text = p_title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT
        
        p_d = tf_c.add_paragraph()
        p_d.text = "\n" + p_desc
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_LIGHT_MUTED

    add_footer(s1, 1, is_dark=True)

    # -------------------------------------------------------------
    # SLIDE 2: Modernizing the Enterprise: Clean Core & AI-Ready via SAP BAIP
    # -------------------------------------------------------------
    s2 = add_base_slide(is_dark=False)
    add_header(s2, "Modernizing the Enterprise: Clean Core & AI-Ready via SAP BAIP", "Strategic Architecture Paradigm", is_dark=False)
    
    # Left Card: Traditional
    card_l = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.35))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = CARD_BG
    card_l.line.color.rgb = CARD_BORDER
    card_l.line.width = Pt(1)
    
    tag_l = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(1.55), Inches(3.6), Inches(0.3))
    tag_l.fill.solid()
    tag_l.fill.fore_color.rgb = RGBColor(254, 242, 242)
    tag_l.line.color.rgb = AMBER_ACCENT
    tag_l.line.width = Pt(1)
    tf_tl = tag_l.text_frame
    tf_tl.margin_top = Inches(0.03)
    p = tf_tl.paragraphs[0]
    p.text = "THE TRADITIONAL TIGHTLY-COUPLED TRAP"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = AMBER_ACCENT
    
    tx_l = s2.shapes.add_textbox(Inches(1.1), Inches(1.95), Inches(5.1), Inches(4.5))
    tf_l = tx_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = 0
    p = tf_l.paragraphs[0]
    p.text = "Direct Modifications Inside S/4HANA & ECC Core"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    items_l = [
        ("Sprawling Custom ABAP Code (Z-Objects)", "Over 40% of standard transactions customized in core user exits and BAdIs, compounding regression risks during patch cycles."),
        ("Dual-Manufacturing Complexity Friction", "Simultaneous Make-to-Order (Automotive / Heavy Engineering) and Make-to-Stock (Appliances) logic colliding in single transactional tables."),
        ("Lengthy ERP Upgrade Cycles", "Upgrades require 6-9 months of testing and custom remediation, causing total innovation lock-in and IT resource drain."),
        ("Inaccessible Operational Silos", "Plant shop-floor data, dealer portals, and service logs trapped in legacy databases, unable to feed modern AI models.")
    ]
    for title, desc in items_l:
        p1 = tf_l.add_paragraph()
        p1.text = f"\n- {title}: "
        p1.font.name = "Arial"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = AMBER_ACCENT
        p2 = tf_l.add_paragraph()
        p2.text = desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED

    # Right Card: Modern Clean Core + BAIP
    card_r = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.35), Inches(5.7), Inches(5.35))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = CARD_BG
    card_r.line.color.rgb = RGBColor(199, 210, 254)
    card_r.line.width = Pt(1.5)
    
    tag_r = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.133), Inches(1.55), Inches(4.5), Inches(0.3))
    tag_r.fill.solid()
    tag_r.fill.fore_color.rgb = RGBColor(238, 242, 255)
    tag_r.line.color.rgb = PURPLE_ACCENT
    tag_r.line.width = Pt(1)
    tf_tr = tag_r.text_frame
    tf_tr.margin_top = Inches(0.03)
    p = tf_tr.paragraphs[0]
    p.text = "THE CLEAN CORE & SAP BAIP ARCHITECTURE"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    
    tx_r = s2.shapes.add_textbox(Inches(7.133), Inches(1.95), Inches(5.1), Inches(4.5))
    tf_r = tx_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = 0
    p = tf_r.paragraphs[0]
    p.text = "Side-by-Side Extensibility Powered by SAP BTP"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    items_r = [
        ("Pure Zero-Modification Core", "S/4HANA/ECC runs pristine standard transactions; all custom apps, workflows, and portals run side-by-side on SAP BTP."),
        ("SAP Business AI Platform (BAIP) Foundation", "Combines mission-critical ERP data, business process context, and enterprise security guardrails into smart AI applications."),
        ("Standardized OData & Event Mesh", "Loosely-coupled asynchronous event brokers react in real time to production orders, goods movements, and service tickets."),
        ("60% Faster Deployment & Effortless Upgrades", "Release new dealer capabilities or shop floor apps in weeks without touching ERP core or requiring system downtime.")
    ]
    for title, desc in items_r:
        p1 = tf_r.add_paragraph()
        p1.text = f"\n+ {title}: "
        p1.font.name = "Arial"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = SAP_BLUE
        p2 = tf_r.add_paragraph()
        p2.text = desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED

    add_footer(s2, 2, is_dark=False)

    # -------------------------------------------------------------
    # SLIDE 3: Use Case Vertical 1 – SAP Business Data Cloud & Data Fabric
    # -------------------------------------------------------------
    s3 = add_base_slide(is_dark=False)
    add_header(s3, "Vertical 1: SAP Business Data Cloud & Data Fabric", "Enterprise Data Architecture & Planning", is_dark=False)
    
    c_cap = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.35))
    c_cap.fill.solid()
    c_cap.fill.fore_color.rgb = CARD_BG
    c_cap.line.color.rgb = CARD_BORDER
    c_cap.line.width = Pt(1)
    
    tx_cap = s3.shapes.add_textbox(Inches(1.1), Inches(1.55), Inches(5.1), Inches(4.9))
    tf_cap = tx_cap.text_frame
    tf_cap.word_wrap = True
    tf_cap.margin_left = tf_cap.margin_top = 0
    p = tf_cap.paragraphs[0]
    p.text = "Unified Data Fabric & Agentic Analytics"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP_BLUE
    
    cap_points = [
        ("SAP Analytics Cloud (SAC)", "Combines agentic analytics, predictive forecasting, and unified enterprise planning into a single AI-ready decision-making portal."),
        ("SAP Datasphere & Data Mesh", "Harmonizes real-time operational data across multi-plant discrete manufacturing lines without physical data replication."),
        ("SAP Business Warehouse & Databricks Open Data", "Integrates telemetry from shop-floor CNC machines with historical sales data via zero-copy data federation."),
        ("SAP HANA Cloud & Master Data Governance (MDG)", "Provides high-throughput multi-model engine and automated golden-record stewardship for materials, BOMs, and vendor records.")
    ]
    for t, d in cap_points:
        p1 = tf_cap.add_paragraph()
        p1.text = f"\n* {t}"
        p1.font.name = "Arial"
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p2 = tf_cap.add_paragraph()
        p2.text = d
        p2.font.name = "Calibri"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED

    # Case 1: Healthium Medtech Pvt. Ltd.
    c_ref1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.35), Inches(5.7), Inches(2.55))
    c_ref1.fill.solid()
    c_ref1.fill.fore_color.rgb = CARD_BG
    c_ref1.line.color.rgb = CARD_BORDER
    c_ref1.line.width = Pt(1)
    
    bg1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.5), Inches(2.6), Inches(0.26))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(236, 253, 245)
    bg1.line.color.rgb = EMERALD_GREEN
    bg1.line.width = Pt(1)
    p = bg1.text_frame.paragraphs[0]
    p.text = "PROVEN REFERENCE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    
    tx_r1 = s3.shapes.add_textbox(Inches(7.1), Inches(1.82), Inches(5.15), Inches(1.9))
    tf_r1 = tx_r1.text_frame
    tf_r1.word_wrap = True
    tf_r1.margin_left = tf_r1.margin_top = 0
    p = tf_r1.paragraphs[0]
    p.text = "Healthium Medtech Pvt. Ltd."
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r1.add_paragraph()
    p.text = "- Challenge: Fragmented multi-plant manufacturing and distribution reporting causing 5-day month-end consolidation delays."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r1.add_paragraph()
    p.text = "- Solution: Enterprise data integration & automated analytics reporting fabric connecting production hubs to SAC."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r1.add_paragraph()
    p.text = "- Impact: Real-time visibility into production yields; reduced reporting turnaround from 5 days to 2 hours."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE

    # Case 2: Hical Technologies Private Limited
    c_ref2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(4.15), Inches(5.7), Inches(2.55))
    c_ref2.fill.solid()
    c_ref2.fill.fore_color.rgb = CARD_BG
    c_ref2.line.color.rgb = CARD_BORDER
    c_ref2.line.width = Pt(1)
    
    bg2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(4.3), Inches(2.6), Inches(0.26))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = RGBColor(236, 253, 245)
    bg2.line.color.rgb = EMERALD_GREEN
    bg2.line.width = Pt(1)
    p = bg2.text_frame.paragraphs[0]
    p.text = "PROVEN REFERENCE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    
    tx_r2 = s3.shapes.add_textbox(Inches(7.1), Inches(4.62), Inches(5.15), Inches(1.9))
    tf_r2 = tx_r2.text_frame
    tf_r2.word_wrap = True
    tf_r2.margin_left = tf_r2.margin_top = 0
    p = tf_r2.paragraphs[0]
    p.text = "Hical Technologies Private Limited"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r2.add_paragraph()
    p.text = "- Challenge: High discrepancy rates in material master records across high-precision electronic manufacturing units."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r2.add_paragraph()
    p.text = "- Solution: Optimized master data governance (MDG) and unified analytics layer on BTP for real-time validation."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r2.add_paragraph()
    p.text = "- Impact: 99.4% master data accuracy, 70% faster part-number onboarding, and streamlined executive decision-making."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE

    add_footer(s3, 3, is_dark=False)

    # -------------------------------------------------------------
    # SLIDE 4: Use Case Vertical 2 – SAP Business AI & Intelligent Agents
    # -------------------------------------------------------------
    s4 = add_base_slide(is_dark=False)
    add_header(s4, "Vertical 2: SAP Business AI & Autonomous Intelligent Agents", "Enterprise GenAI & Agentic Workflows", is_dark=False)
    
    c_cap4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.35))
    c_cap4.fill.solid()
    c_cap4.fill.fore_color.rgb = CARD_BG
    c_cap4.line.color.rgb = CARD_BORDER
    c_cap4.line.width = Pt(1)
    
    tx_cap4 = s4.shapes.add_textbox(Inches(1.1), Inches(1.55), Inches(5.1), Inches(4.9))
    tf_cap4 = tx_cap4.text_frame
    tf_cap4.word_wrap = True
    tf_cap4.margin_left = tf_cap4.margin_top = 0
    p = tf_cap4.paragraphs[0]
    p.text = "Conversational, Context-Aware SAP Business AI"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    
    ai_points = [
        ("Joule Work & Natural Language ERP", "Executes complex transactional tasks (stock checks, PO approvals, production status inquiries) directly through natural conversational prompts."),
        ("Joule Autonomous Agents", "Proactively identifies supply chain bottlenecks, cross-references vendor lead times, and recommends rescheduling actions before lines halt."),
        ("Joule Studio & Model Customization", "Enables custom LLM fine-tuning and enterprise groundings tailored to discrete engineering BOMs and appliance warranty codes."),
        ("Joule for Consultants & Developers", "Accelerates BTP extension delivery by 40% through AI-assisted code generation, automated testing, and SAP API discovery."),
        ("Enterprise AI Foundation Services", "Pre-built connectors to Azure OpenAI, Vertex AI, and Hugging Face with robust data privacy, masking, and role-based access.")
    ]
    for t, d in ai_points:
        p1 = tf_cap4.add_paragraph()
        p1.text = f"\n* {t}"
        p1.font.name = "Arial"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p2 = tf_cap4.add_paragraph()
        p2.text = d
        p2.font.name = "Calibri"
        p2.font.size = Pt(9.2)
        p2.font.color.rgb = TEXT_MUTED

    # Case 1: SAIL
    c_ref41 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.35), Inches(5.7), Inches(2.55))
    c_ref41.fill.solid()
    c_ref41.fill.fore_color.rgb = CARD_BG
    c_ref41.line.color.rgb = CARD_BORDER
    c_ref41.line.width = Pt(1)
    
    bg41 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.5), Inches(2.6), Inches(0.26))
    bg41.fill.solid()
    bg41.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bg41.line.color.rgb = PURPLE_ACCENT
    bg41.line.width = Pt(1)
    p = bg41.text_frame.paragraphs[0]
    p.text = "PROVEN REFERENCE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    
    tx_r41 = s4.shapes.add_textbox(Inches(7.1), Inches(1.82), Inches(5.15), Inches(1.9))
    tf_r41 = tx_r41.text_frame
    tf_r41.word_wrap = True
    tf_r41.margin_left = tf_r41.margin_top = 0
    p = tf_r41.paragraphs[0]
    p.text = "SAIL (Steel Authority of India Limited)"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r41.add_paragraph()
    p.text = "- Challenge: Heavy industrial multi-site operations with massive query backlogs on operational plant parameters."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r41.add_paragraph()
    p.text = "- Solution: Deployed intelligent assistant frameworks to streamline heavy industrial operational insights and query management."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r41.add_paragraph()
    p.text = "- Impact: 65% faster response times for engineering queries, enabling rapid resolution of plant maintenance tickets."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT

    # Case 2: Taurani Holdings
    c_ref42 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(4.15), Inches(5.7), Inches(2.55))
    c_ref42.fill.solid()
    c_ref42.fill.fore_color.rgb = CARD_BG
    c_ref42.line.color.rgb = CARD_BORDER
    c_ref42.line.width = Pt(1)
    
    bg42 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(4.3), Inches(2.6), Inches(0.26))
    bg42.fill.solid()
    bg42.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bg42.line.color.rgb = PURPLE_ACCENT
    bg42.line.width = Pt(1)
    p = bg42.text_frame.paragraphs[0]
    p.text = "PROVEN REFERENCE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    
    tx_r42 = s4.shapes.add_textbox(Inches(7.1), Inches(4.62), Inches(5.15), Inches(1.9))
    tf_r42 = tx_r42.text_frame
    tf_r42.word_wrap = True
    tf_r42.margin_left = tf_r42.margin_top = 0
    p = tf_r42.paragraphs[0]
    p.text = "Taurani Holdings"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r42.add_paragraph()
    p.text = "- Challenge: Unpredictable raw material price fluctuations impacting margins across manufacturing divisions."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r42.add_paragraph()
    p.text = "- Solution: Integrated advanced AI models on BTP to drive predictive analytics and contextual operational efficiency."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r42.add_paragraph()
    p.text = "- Impact: Enabled high-confidence procurement hedging, saving millions annually while sustaining operating margins."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT

    add_footer(s4, 4, is_dark=False)

    # -------------------------------------------------------------
    # SLIDE 5: Use Case Vertical 3 – SAP Integration Suite & Enterprise Connectivity
    # -------------------------------------------------------------
    s5 = add_base_slide(is_dark=False)
    add_header(s5, "Vertical 3: SAP Integration Suite & Enterprise Connectivity", "Hybrid Integration, Event Mesh & B2B Hubs", is_dark=False)
    
    c_cap5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.35))
    c_cap5.fill.solid()
    c_cap5.fill.fore_color.rgb = CARD_BG
    c_cap5.line.color.rgb = CARD_BORDER
    c_cap5.line.width = Pt(1)
    
    tx_cap5 = s5.shapes.add_textbox(Inches(1.1), Inches(1.55), Inches(5.1), Inches(4.9))
    tf_cap5 = tx_cap5.text_frame
    tf_cap5.word_wrap = True
    tf_cap5.margin_left = tf_cap5.margin_top = 0
    p = tf_cap5.paragraphs[0]
    p.text = "Mission-Critical Hybrid Integration Fabric"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP_BLUE
    
    int_points = [
        ("Agentic AI Integration", "AI-driven mapping algorithms that automatically synthesize schemas across legacy EDI (ANSI X12, EDIFACT) and modern JSON REST endpoints."),
        ("Application Integration (Cloud Integration)", "Out-of-the-box pre-packaged content connecting SAP S/4HANA & ECC to shop floor MES, PLM, and CRM systems."),
        ("Enterprise API Lifecycle Management", "Governs, monitors, and monetizes internal and external API gateways with rate limiting, OAuth 2.0, and high availability."),
        ("Event-Driven Architecture (SAP Event Mesh)", "Asynchronously decouples event triggers (e.g. Goods Receipt, Quality Hold) to guarantee zero core ERP performance hits."),
        ("Comprehensive B2B / EDI Integration Strategy", "Centralizes hundreds of Tier-1 automotive OEM release schedules (830/862 EDI) and dealer orders into standard SAP SD sales orders.")
    ]
    for t, d in int_points:
        p1 = tf_cap5.add_paragraph()
        p1.text = f"\n* {t}"
        p1.font.name = "Arial"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p2 = tf_cap5.add_paragraph()
        p2.text = d
        p2.font.name = "Calibri"
        p2.font.size = Pt(9.2)
        p2.font.color.rgb = TEXT_MUTED

    # Case 1: TATA Electronics Pvt. Ltd. (TEPL)
    c_ref51 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.35), Inches(5.7), Inches(2.55))
    c_ref51.fill.solid()
    c_ref51.fill.fore_color.rgb = CARD_BG
    c_ref51.line.color.rgb = CARD_BORDER
    c_ref51.line.width = Pt(1)
    
    bg51 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.5), Inches(2.6), Inches(0.26))
    bg51.fill.solid()
    bg51.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bg51.line.color.rgb = SAP_BLUE
    bg51.line.width = Pt(1)
    p = bg51.text_frame.paragraphs[0]
    p.text = "PROVEN REFERENCE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    tx_r51 = s5.shapes.add_textbox(Inches(7.1), Inches(1.82), Inches(5.15), Inches(1.9))
    tf_r51 = tx_r51.text_frame
    tf_r51.word_wrap = True
    tf_r51.margin_left = tf_r51.margin_top = 0
    p = tf_r51.paragraphs[0]
    p.text = "TATA Electronics Pvt. Ltd. (TEPL)"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r51.add_paragraph()
    p.text = "- Challenge: High-precision greenfield manufacturing requiring millisecond-level integration between plant shop-floor MES and ERP."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r51.add_paragraph()
    p.text = "- Solution: Architected high-volume B2B and application integrations connecting disparate manufacturing execution systems."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r51.add_paragraph()
    p.text = "- Impact: Flawless handling of over 2M daily plant transactions with sub-second latency and zero data dropouts."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE

    # Case 2: TSAT
    c_ref52 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(4.15), Inches(5.7), Inches(2.55))
    c_ref52.fill.solid()
    c_ref52.fill.fore_color.rgb = CARD_BG
    c_ref52.line.color.rgb = CARD_BORDER
    c_ref52.line.width = Pt(1)
    
    bg52 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(4.3), Inches(2.6), Inches(0.26))
    bg52.fill.solid()
    bg52.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bg52.line.color.rgb = SAP_BLUE
    bg52.line.width = Pt(1)
    p = bg52.text_frame.paragraphs[0]
    p.text = "PROVEN REFERENCE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    tx_r52 = s5.shapes.add_textbox(Inches(7.1), Inches(4.62), Inches(5.15), Inches(1.9))
    tf_r52 = tx_r52.text_frame
    tf_r52.word_wrap = True
    tf_r52.margin_left = tf_r52.margin_top = 0
    p = tf_r52.paragraphs[0]
    p.text = "Tata Semiconductor Assembly & Test (TSAT)"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r52.add_paragraph()
    p.text = "- Challenge: High-security semiconductor environment needing continuous data exchange across segregated network zones."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r52.add_paragraph()
    p.text = "- Solution: Implemented robust API lifecycle management and event-driven data flows to sync plant operations securely."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r52.add_paragraph()
    p.text = "- Impact: 100% auditable traceability, zero security breaches, and instantaneous propagation of lot release orders."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE

    add_footer(s5, 5, is_dark=False)

    # -------------------------------------------------------------
    # SLIDE 6: Use Case Vertical 4 – Workflows, Automation & Signavio Intelligence
    # -------------------------------------------------------------
    s6 = add_base_slide(is_dark=False)
    add_header(s6, "Vertical 4: Workflows, Automation & Signavio Process AI", "Process Mining, Low-Code Automation & Central Governance", is_dark=False)
    
    c_cap6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.35))
    c_cap6.fill.solid()
    c_cap6.fill.fore_color.rgb = CARD_BG
    c_cap6.line.color.rgb = CARD_BORDER
    c_cap6.line.width = Pt(1)
    
    tx_cap6 = s6.shapes.add_textbox(Inches(1.1), Inches(1.55), Inches(5.1), Inches(4.9))
    tf_cap6 = tx_cap6.text_frame
    tf_cap6.word_wrap = True
    tf_cap6.margin_left = tf_cap6.margin_top = 0
    p = tf_cap6.paragraphs[0]
    p.text = "Next-Gen Low-Code Automation & Mining"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP_BLUE
    
    wf_points = [
        ("SAP Build Process Automation", "Unified low-code/no-code studio combining robotic process automation (RPA bots) and workflow orchestration in one cloud environment."),
        ("Dynamic Business Rules Engine", "Decouples decision logic (pricing thresholds, discount approvals, warranty validity) from ABAP code, editable directly by business leads."),
        ("SAP Central Task Center", "Harmonizes approval inboxes across ECC, S/4HANA, Ariba, and SuccessFactors into a single, unified mobile/desktop experience."),
        ("SAP Signavio Process AI & Insights", "Continuously mines operational data streams to reveal hidden cycle-time bottlenecks, rework loops in manufacturing, and Maverick buying."),
        ("Intelligent Document Processing (IDP)", "Extracts and validates OCR data from incoming supplier invoices, vendor test certificates, and shipping manifests automatically.")
    ]
    for t, d in wf_points:
        p1 = tf_cap6.add_paragraph()
        p1.text = f"\n* {t}"
        p1.font.name = "Arial"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        p2 = tf_cap6.add_paragraph()
        p2.text = d
        p2.font.name = "Calibri"
        p2.font.size = Pt(9.2)
        p2.font.color.rgb = TEXT_MUTED

    # Hero Case: Dr. Reddy's
    c_ref6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.35), Inches(5.7), Inches(5.35))
    c_ref6.fill.solid()
    c_ref6.fill.fore_color.rgb = CARD_BG
    c_ref6.line.color.rgb = CARD_BORDER
    c_ref6.line.width = Pt(1)
    
    bg6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.55), Inches(3.0), Inches(0.28))
    bg6.fill.solid()
    bg6.fill.fore_color.rgb = RGBColor(236, 253, 245)
    bg6.line.color.rgb = EMERALD_GREEN
    bg6.line.width = Pt(1)
    p = bg6.text_frame.paragraphs[0]
    p.text = "HERO ENTERPRISE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    
    tx_r6 = s6.shapes.add_textbox(Inches(7.1), Inches(1.95), Inches(5.15), Inches(4.5))
    tf_r6 = tx_r6.text_frame
    tf_r6.word_wrap = True
    tf_r6.margin_left = tf_r6.margin_top = 0
    p = tf_r6.paragraphs[0]
    p.text = "Dr. Reddy's Laboratories Ltd."
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r6.add_paragraph()
    p.text = "Enterprise-Wide Workflow & Compliance Automation"
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED
    
    p = tf_r6.add_paragraph()
    p.text = "\n- Business Context & Challenge:\nStringent regulatory compliance requirements and multi-tier approval matrix across distributed plant facilities created immense manual verification bottlenecks, delayed batch releases, and introduced human error risk."
    p.font.name = "Calibri"
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r6.add_paragraph()
    p.text = "\n- Solution Deployed on SAP BTP:\nAutomated complex compliance and operational approval workflows across manufacturing plants using SAP Build Process Automation integrated with SAP Central Task Center."
    p.font.name = "Calibri"
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_DARK
    
    metrics = [
        ("80%", "Reduction in Approval Lag"),
        ("100%", "Audit-Proof GxP Trail"),
        ("0 Core Mod", "Standard S/4HANA Preserved")
    ]
    for mi, (m_val, m_lbl) in enumerate(metrics):
        m_left = Inches(7.1 + mi * 1.7)
        m_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, m_left, Inches(5.35), Inches(1.6), Inches(1.1))
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = CARD_BG_MUTED
        m_card.line.color.rgb = CARD_BORDER
        m_card.line.width = Pt(1)
        
        tx_m = s6.shapes.add_textbox(m_left, Inches(5.42), Inches(1.6), Inches(0.95))
        tf_m = tx_m.text_frame
        tf_m.word_wrap = True
        p_mv = tf_m.paragraphs[0]
        p_mv.alignment = PP_ALIGN.CENTER
        p_mv.text = m_val
        p_mv.font.name = "Arial"
        p_mv.font.size = Pt(14)
        p_mv.font.bold = True
        p_mv.font.color.rgb = EMERALD_GREEN
        
        p_ml = tf_m.add_paragraph()
        p_ml.alignment = PP_ALIGN.CENTER
        p_ml.text = m_lbl
        p_ml.font.name = "Calibri"
        p_ml.font.size = Pt(8.5)
        p_ml.font.color.rgb = TEXT_MUTED

    add_footer(s6, 6, is_dark=False)

    # -------------------------------------------------------------
    # SLIDE 7: Use Case Vertical 5 – Custom Application Development & Enterprise Scale
    # -------------------------------------------------------------
    s7 = add_base_slide(is_dark=False)
    add_header(s7, "Vertical 5: Custom App Development & Enterprise Scale", "CAP vs. RAP Architecture & Industrial Scale", is_dark=False)
    
    c_cap7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.35))
    c_cap7.fill.solid()
    c_cap7.fill.fore_color.rgb = CARD_BG
    c_cap7.line.color.rgb = CARD_BORDER
    c_cap7.line.width = Pt(1)
    
    tx_cap7 = s7.shapes.add_textbox(Inches(1.1), Inches(1.55), Inches(5.1), Inches(4.9))
    tf_cap7 = tx_cap7.text_frame
    tf_cap7.word_wrap = True
    tf_cap7.margin_left = tf_cap7.margin_top = 0
    p = tf_cap7.paragraphs[0]
    p.text = "Modern Extension Frameworks: CAP vs. RAP"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP_BLUE
    
    p = tf_cap7.add_paragraph()
    p.text = "Eliminating uncontrolled Z-tables through standardized, lifecycle-managed application frameworks:"
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    
    cap_sub = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(2.25), Inches(5.1), Inches(1.95))
    cap_sub.fill.solid()
    cap_sub.fill.fore_color.rgb = RGBColor(240, 249, 255)
    cap_sub.line.color.rgb = CYAN_ACCENT
    cap_sub.line.width = Pt(1)
    tx_cs = s7.shapes.add_textbox(Inches(1.2), Inches(2.32), Inches(4.9), Inches(1.8))
    tf_cs = tx_cs.text_frame
    tf_cs.word_wrap = True
    p = tf_cs.paragraphs[0]
    p.text = "SAP CAP (Cloud Application Programming Model)"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    p = tf_cs.add_paragraph()
    p.text = "- Runtime: Node.js / Java on BTP Cloud Foundry & Kyma (Kubernetes).\n- Best Fit: Side-by-side consumer portals, dealer mobile apps, multi-tenant SaaS.\n- Advantage: Full decoupling from ERP release cycles; scalable cloud-native microservices."
    p.font.name = "Calibri"
    p.font.size = Pt(9.2)
    p.font.color.rgb = TEXT_DARK
    
    rap_sub = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(4.35), Inches(5.1), Inches(1.95))
    rap_sub.fill.solid()
    rap_sub.fill.fore_color.rgb = RGBColor(248, 250, 252)
    rap_sub.line.color.rgb = PURPLE_ACCENT
    rap_sub.line.width = Pt(1)
    tx_rs = s7.shapes.add_textbox(Inches(1.2), Inches(4.42), Inches(4.9), Inches(1.8))
    tf_rs = tx_rs.text_frame
    tf_rs.word_wrap = True
    p = tf_rs.paragraphs[0]
    p.text = "SAP RAP (RESTful ABAP Programming Model)"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    p = tf_rs.add_paragraph()
    p.text = "- Runtime: Native ABAP Cloud in S/4HANA & BTP ABAP Environment.\n- Best Fit: In-app extensions, heavy transactional ledger calculations, core extensions.\n- Advantage: Direct CDS data model binding with zero impedance mismatch."
    p.font.name = "Calibri"
    p.font.size = Pt(9.2)
    p.font.color.rgb = TEXT_DARK

    # Case 1: L&T Construction & Mining Machinery
    c_ref71 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.35), Inches(5.7), Inches(2.55))
    c_ref71.fill.solid()
    c_ref71.fill.fore_color.rgb = CARD_BG
    c_ref71.line.color.rgb = CARD_BORDER
    c_ref71.line.width = Pt(1)
    
    bg71 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.5), Inches(2.6), Inches(0.26))
    bg71.fill.solid()
    bg71.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bg71.line.color.rgb = SAP_BLUE
    bg71.line.width = Pt(1)
    p = bg71.text_frame.paragraphs[0]
    p.text = "PROVEN REFERENCE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    tx_r71 = s7.shapes.add_textbox(Inches(7.1), Inches(1.82), Inches(5.15), Inches(1.9))
    tf_r71 = tx_r71.text_frame
    tf_r71.word_wrap = True
    tf_r71.margin_left = tf_r71.margin_top = 0
    p = tf_r71.paragraphs[0]
    p.text = "L&T Construction & Mining Machinery"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r71.add_paragraph()
    p.text = "- Solution: Custom Dealer Management System (DEMANS) leveraging advanced side-by-side extension patterns on BTP."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r71.add_paragraph()
    p.text = "- Capabilities: Real-time equipment telemetry, automated warranty claims, spare parts ordering connected directly to SAP SD."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r71.add_paragraph()
    p.text = "- Impact: 45% reduction in spare turnaround time; zero impact on core SAP ERP during high-frequency dealer peak days."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE

    # Case 2: Indian Agro & Food Industries Limited (IB Group)
    c_ref72 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(4.15), Inches(5.7), Inches(2.55))
    c_ref72.fill.solid()
    c_ref72.fill.fore_color.rgb = CARD_BG
    c_ref72.line.color.rgb = CARD_BORDER
    c_ref72.line.width = Pt(1)
    
    bg72 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(4.3), Inches(2.6), Inches(0.26))
    bg72.fill.solid()
    bg72.fill.fore_color.rgb = RGBColor(238, 242, 255)
    bg72.line.color.rgb = SAP_BLUE
    bg72.line.width = Pt(1)
    p = bg72.text_frame.paragraphs[0]
    p.text = "PROVEN REFERENCE CASE STUDY"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE
    
    tx_r72 = s7.shapes.add_textbox(Inches(7.1), Inches(4.62), Inches(5.15), Inches(1.9))
    tf_r72 = tx_r72.text_frame
    tf_r72.word_wrap = True
    tf_r72.margin_left = tf_r72.margin_top = 0
    p = tf_r72.paragraphs[0]
    p.text = "Indian Agro & Food Industries Limited (IB Group)"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p = tf_r72.add_paragraph()
    p.text = "- Solution: Tailored operational agility solutions for multi-plant agricultural and food supply chains on BTP."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r72.add_paragraph()
    p.text = "- Capabilities: Multi-plant inventory synchronization, automated perishable batch tracking, and dynamic dispatch scheduling."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    p = tf_r72.add_paragraph()
    p.text = "- Impact: Optimized end-to-end execution speed; 50% decrease in manual scheduling conflicts across geographically dispersed plants."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = SAP_BLUE

    add_footer(s7, 7, is_dark=False)

    # -------------------------------------------------------------
    # SLIDE 8: Exclusive Spotlight – The Lumbini Elitia Portfolio
    # -------------------------------------------------------------
    s8 = add_base_slide(is_dark=False)
    add_header(s8, "Beyond Standard SAP: The Lumbini Elitia Product Suite", "Proprietary Ready-to-Deploy Enterprise Microservices", is_dark=False)
    
    top_bar = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.55))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = RGBColor(238, 242, 255)
    top_bar.line.color.rgb = SAP_BLUE
    top_bar.line.width = Pt(1)
    tx_tb = s8.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(11.333), Inches(0.45))
    tf_tb = tx_tb.text_frame
    tf_tb.word_wrap = True
    p = tf_tb.paragraphs[0]
    p.text = "Pre-built, non-SAP microservice accelerators engineered to integrate directly into S/4HANA & ECC via BTP — eliminating costly custom build cycles."
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = SAP_DEEP_BLUE
    
    elitia_services = [
        ("Supplier Collaboration Portal", 
         "Streamlining vendor engagement, advance shipment notices (ASN), digital gate passes, and multi-tier material tracking in real time.",
         "SAP MM / Ariba Integration",
         CYAN_ACCENT),
         
        ("Data Governance Suite", 
         "Ensuring enterprise-wide data integrity, automated validation workflows, duplicate prevention, and compliance audits across plants.",
         "SAP MDG / Datasphere",
         SAP_BLUE),
         
        ("Master Data Validator", 
         "Automated AI-driven record verification for materials, customer accounts, and vendor tax compliance before posting to ERP.",
         "Pre-Posting Rule Engine",
         PURPLE_ACCENT),
         
        ("Customer Experience Portal", 
         "Elevating B2B dealer and consumer touchpoints with real-time order tracking, dynamic credit checks, and self-service warranty claims.",
         "SAP SD / Commerce Hub",
         EMERALD_GREEN),
         
        ("Workforce Safety Management", 
         "Ensuring heavy plant regulatory compliance, hazardous area permits, incident reporting, and contractor safety certification logs.",
         "SAP EHS / Plant Maintenance",
         AMBER_ACCENT)
    ]
    
    for idx, (es_title, es_desc, es_int, es_col) in enumerate(elitia_services):
        c_x = Inches(0.8 + idx * (2.2 + 0.18))
        c_shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, Inches(2.0), Inches(2.2), Inches(4.7))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = CARD_BORDER
        c_shape.line.width = Pt(1)
        
        bar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_x + Inches(0.15), Inches(2.15), Inches(1.9), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = es_col
        bar.line.fill.background()
        
        ptag = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.15), Inches(2.35), Inches(1.9), Inches(0.24))
        ptag.fill.solid()
        ptag.fill.fore_color.rgb = CARD_BG_MUTED
        ptag.line.fill.background()
        p = ptag.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = f"ELITIA MODULE {idx+1}"
        p.font.name = "Arial"
        p.font.size = Pt(7.5)
        p.font.bold = True
        p.font.color.rgb = es_col
        
        tx_e = s8.shapes.add_textbox(c_x + Inches(0.15), Inches(2.7), Inches(1.9), Inches(3.8))
        tf_e = tx_e.text_frame
        tf_e.word_wrap = True
        tf_e.margin_left = tf_e.margin_top = 0
        p = tf_e.paragraphs[0]
        p.text = es_title
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        p = tf_e.add_paragraph()
        p.text = "\n" + es_desc
        p.font.name = "Calibri"
        p.font.size = Pt(9.2)
        p.font.color.rgb = TEXT_MUTED
        
        p = tf_e.add_paragraph()
        p.text = f"\nConnector:\n{es_int}"
        p.font.name = "Arial"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = es_col

    add_footer(s8, 8, is_dark=False)

    # -------------------------------------------------------------
    # SLIDE 9: The Lumbini Advantage – Delivery Methodology & Strategic Partnership
    # -------------------------------------------------------------
    s9 = add_base_slide(is_dark=True)
    add_header(s9, "Low-Risk, High-Velocity Enterprise Execution", "The Lumbini Delivery Methodology & Strategic Partnership", is_dark=True)
    
    eng_pillars = [
        ("Zero-Disruption Core Strategy",
         "100% S/4HANA Upgrade Safety",
         "All custom logic, integrations, and apps reside strictly on SAP BTP side-by-side environments.\n\n- Zero core ABAP modifications\n- S/4HANA release patches apply seamlessly\n- Guarantees ongoing SAP support compliance",
         CYAN_ACCENT),
         
        ("Accelerated Time-to-Value",
         "Cut Timelines by Up to 50%",
         "Leverage ready-made Elitia microservices, pre-packaged SAP BTP integration flows, and automated CI/CD pipelines.\n\n- Ready-to-deploy architectural blueprints\n- Pre-built connectors for SAP SD, MM, PP, & QM\n- Rapid ROI realization in under 90 days",
         EMERALD_GREEN),
         
        ("Collaborative Pilot Discovery",
         "Targeted 2-Week Joint Workshop",
         "A structured, zero-risk engagement to scope your highest-impact plant or supply chain bottleneck.\n\n- Day 1-3: Process mining & bottleneck discovery\n- Day 4-7: Target BTP side-by-side architecture\n- Day 8-10: Working MVP prototype & business case",
         PURPLE_ACCENT)
    ]
    
    for pi, (ep_title, ep_sub, ep_body, ep_col) in enumerate(eng_pillars):
        c_left = Inches(0.8 + pi * 4.0)
        c_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(1.4), Inches(3.733), Inches(4.3))
        c_card.fill.solid()
        c_card.fill.fore_color.rgb = DARK_CARD
        c_card.line.color.rgb = DARK_BORDER
        c_card.line.width = Pt(1)
        
        c_bar = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, c_left + Inches(0.25), Inches(1.65), Inches(0.5), Inches(0.06))
        c_bar.fill.solid()
        c_bar.fill.fore_color.rgb = ep_col
        c_bar.line.fill.background()
        
        tx_p = s9.shapes.add_textbox(c_left + Inches(0.25), Inches(1.85), Inches(3.233), Inches(3.7))
        tf_p = tx_p.text_frame
        tf_p.word_wrap = True
        tf_p.margin_left = tf_p.margin_top = 0
        p = tf_p.paragraphs[0]
        p.text = ep_title
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        
        p = tf_p.add_paragraph()
        p.text = ep_sub
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ep_col
        
        p = tf_p.add_paragraph()
        p.text = "\n" + ep_body
        p.font.name = "Calibri"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT_MUTED
        
    callout = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.9), Inches(11.733), Inches(0.85))
    callout.fill.solid()
    callout.fill.fore_color.rgb = RGBColor(16, 37, 80)
    callout.line.color.rgb = CYAN_ACCENT
    callout.line.width = Pt(1)
    
    tx_co = s9.shapes.add_textbox(Inches(1.1), Inches(5.98), Inches(11.133), Inches(0.7))
    tf_co = tx_co.text_frame
    tf_co.word_wrap = True
    tf_co.margin_left = tf_co.margin_top = 0
    p = tf_co.paragraphs[0]
    p.text = "PROPOSED IMMEDIATE NEXT STEP: 2-WEEK LIGHTHOUSE DISCOVERY PILOT"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    
    p = tf_co.add_paragraph()
    p.text = "Let us partner with your enterprise IT and plant leadership to model one critical workflow on SAP BTP with our Elitia accelerators—delivering verified time-to-value before capital commitment."
    p.font.name = "Calibri"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_LIGHT

    add_footer(s9, 9, is_dark=True)

    # -------------------------------------------------------------
    # Save Presentation
    # -------------------------------------------------------------
    prs.save(output_path)
    print(f"Successfully created presentation at: {output_path}")

if __name__ == "__main__":
    create_deck()
