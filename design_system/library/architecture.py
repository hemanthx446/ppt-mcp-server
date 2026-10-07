"""
Architecture Visual Primitives (Primitives 1–9).

Native, editable PowerPoint architecture diagrams for enterprise solutions,
manufacturing IT/OT, SAP integration, and cloud/edge topology.
"""

from dataclasses import dataclass, field
from typing import List, Optional, Tuple
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from ..typography import TypographySystem
from ..spacing import Margins, SpacingScale, GridCalculator
from ..color import Theme
from ..primitives import SurfacePrimitive, ConnectorPrimitive


@dataclass
class ArchTierData:
    tier_name: str
    subtitle: str
    subsystems: List[str]
    badge: Optional[str] = None
    protocol_to_next: Optional[str] = None


class SystemArchitecturePrimitive:
    """1. Multi-tier decoupled system architecture with protocol connectors and boundaries."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               tiers: List[ArchTierData], theme: Theme):
        count = len(tiers)
        cols = GridCalculator.get_columns(count, left=left, total_width=width, gap=0.55)

        for i, tier in enumerate(tiers):
            cx, cw = cols[i]
            # Outer boundary box
            SurfacePrimitive.render(slide, cx, top, cw, height, theme)

            # Header strip
            strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(top), Inches(cw), Inches(0.06))
            strip.fill.solid()
            strip.fill.fore_color.rgb = theme.border_accent
            strip.line.fill.background()

            tb = slide.shapes.add_textbox(Inches(cx + 0.16), Inches(top + 0.16), Inches(cw - 0.32), Inches(height - 0.32))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

            # Tier Name
            p1 = tf.paragraphs[0]
            TypographySystem.apply_to_paragraph(p1, TypographySystem.LABEL, tier.tier_name, theme.text_accent)

            # Subtitle
            p2 = tf.add_paragraph()
            TypographySystem.apply_to_paragraph(p2, TypographySystem.BODY_STRONG, tier.subtitle, theme.text_primary)
            p2.space_after = Pt(10.0)

            # Subsystems
            for sub in tier.subsystems[:4]:
                ps = tf.add_paragraph()
                ps.space_after = Pt(6.0)
                TypographySystem.apply_to_paragraph(ps, TypographySystem.BODY, f"•  {sub}", theme.text_muted)

            # Badge
            if tier.badge:
                pb = tf.add_paragraph()
                pb.space_after = Pt(4.0)
                TypographySystem.apply_to_paragraph(pb, TypographySystem.ANNOTATION, tier.badge, theme.status_success)

            # Connector Arrow
            if i < count - 1:
                arr_x = cx + cw + 0.10
                arr_y = top + (height / 2.0) - 0.15
                proto = tier.protocol_to_next if tier.protocol_to_next else "SYNC"
                ConnectorPrimitive.render_arrow(slide, arr_x, arr_y, 0.35, 0.28, "right", theme, label=proto)


class LayeredArchitecturePrimitive:
    """2. Horizontal layered technical stack (Presentation, Application, Computation, Storage)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               layers: List[Tuple[str, str, List[str]]], # [(layer_title, tech_stack, components)]
               theme: Theme):
        count = len(layers)
        rows = GridCalculator.get_rows(count, top=top, total_height=height, gap=0.18)

        for i, (l_title, tech_stack, comps) in enumerate(layers):
            ry, rh = rows[i]
            SurfacePrimitive.render(slide, left, ry, width, rh, theme, fill_color=theme.surface_alt if i % 2 == 1 else theme.surface)

            # Left tag
            tag_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(ry), Inches(2.2), Inches(rh))
            tag_box.fill.solid()
            tag_box.fill.fore_color.rgb = theme.surface
            tag_box.line.color.rgb = theme.border
            tf_t = tag_box.text_frame
            tf_t.word_wrap = True
            tf_t.margin_left = Inches(0.14)
            p_t1 = tf_t.paragraphs[0]
            TypographySystem.apply_to_paragraph(p_t1, TypographySystem.LABEL, l_title, theme.text_accent)
            p_t2 = tf_t.add_paragraph()
            TypographySystem.apply_to_paragraph(p_t2, TypographySystem.CAPTION, tech_stack, theme.text_muted)

            # Right content area
            tb = slide.shapes.add_textbox(Inches(left + 2.35), Inches(ry + 0.10), Inches(width - 2.5), Inches(rh - 0.20))
            tf_c = tb.text_frame
            tf_c.word_wrap = True
            tf_c.margin_left = tf_c.margin_top = 0
            comp_str = "   |   ".join(comps)
            p_c = tf_c.paragraphs[0]
            TypographySystem.apply_to_paragraph(p_c, TypographySystem.BODY, comp_str, theme.text_primary)


class EnterpriseToShopfloorPrimitive:
    """3. ISA-95 Manufacturing Hierarchy (Level 4 ERP -> Level 3 MOM/MES -> Level 2/1 SCADA/PLC)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               isa_levels: List[Tuple[str, str, List[str], str]], # [(level_name, domain, key_functions, protocol)]
               theme: Theme):
        count = len(isa_levels)
        rows = GridCalculator.get_rows(count, top=top, total_height=height, gap=0.16)

        for i, (lvl, domain, funcs, proto) in enumerate(isa_levels):
            ry, rh = rows[i]
            # Level container
            SurfacePrimitive.render(slide, left, ry, width, rh, theme)

            # Level badge
            b_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + 0.12), Inches(ry + 0.10), Inches(2.0), Inches(rh - 0.20))
            b_box.fill.solid()
            b_box.fill.fore_color.rgb = theme.surface_highlight
            b_box.line.color.rgb = theme.border_accent
            tf_b = b_box.text_frame
            tf_b.margin_left = Inches(0.10)
            p_b1 = tf_b.paragraphs[0]
            TypographySystem.apply_to_paragraph(p_b1, TypographySystem.LABEL, lvl, theme.text_accent)
            p_b2 = tf_b.add_paragraph()
            TypographySystem.apply_to_paragraph(p_b2, TypographySystem.CAPTION, domain, theme.text_primary)

            # Core functions
            f_box = slide.shapes.add_textbox(Inches(left + 2.25), Inches(ry + 0.10), Inches(width - 4.5), Inches(rh - 0.20))
            tf_f = f_box.text_frame
            tf_f.margin_left = tf_f.margin_top = 0
            p_f = tf_f.paragraphs[0]
            TypographySystem.apply_to_paragraph(p_f, TypographySystem.BODY, " • ".join(funcs[:3]), theme.text_secondary)

            # Protocol / interface tag
            p_box = slide.shapes.add_textbox(Inches(left + width - 2.1), Inches(ry + 0.10), Inches(1.9), Inches(rh - 0.20))
            tf_p = p_box.text_frame
            tf_p.margin_left = tf_p.margin_top = 0
            p_p = tf_p.paragraphs[0]
            TypographySystem.apply_to_paragraph(p_p, TypographySystem.ANNOTATION, f"INTERLOCK: {proto}", theme.status_warning)


class ApplicationLandscapePrimitive:
    """4. Categorized application domain map (Core ERP, Supply Chain, Operations, Analytics)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               domains: List[Tuple[str, List[Tuple[str, str]]]], # [(domain_name, [(app_name, role)])]
               theme: Theme):
        cols = GridCalculator.get_columns(len(domains), left=left, total_width=width, gap=0.22)
        for i, (d_name, apps) in enumerate(domains):
            cx, cw = cols[i]
            SurfacePrimitive.render(slide, cx, top, cw, height, theme)

            # Header
            h_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(top), Inches(cw), Inches(0.40))
            h_box.fill.solid()
            h_box.fill.fore_color.rgb = theme.surface_highlight
            h_box.line.color.rgb = theme.border
            tf_h = h_box.text_frame
            tf_h.margin_left = Inches(0.12)
            p_h = tf_h.paragraphs[0]
            TypographySystem.apply_to_paragraph(p_h, TypographySystem.LABEL, d_name, theme.text_accent)

            # Apps
            app_y = top + 0.50
            for app_title, app_role in apps[:5]:
                ab = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx + 0.12), Inches(app_y), Inches(cw - 0.24), Inches(0.68))
                ab.fill.solid()
                ab.fill.fore_color.rgb = theme.surface_alt
                ab.line.color.rgb = theme.border
                tf_a = ab.text_frame
                tf_a.margin_left = Inches(0.08)
                tf_a.margin_top = Inches(0.04)
                pa1 = tf_a.paragraphs[0]
                TypographySystem.apply_to_paragraph(pa1, TypographySystem.BODY_STRONG, app_title, theme.text_primary)
                pa2 = tf_a.add_paragraph()
                TypographySystem.apply_to_paragraph(pa2, TypographySystem.CAPTION, app_role, theme.text_muted)
                app_y += 0.74


class IntegrationArchitecturePrimitive:
    """5. Integration middleware architecture: Producer -> Event Queue / Broker -> Consumer."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               producer_info: Tuple[str, List[str]],
               broker_info: Tuple[str, List[str], str], # (name, capabilities, pattern)
               consumer_info: Tuple[str, List[str]],
               theme: Theme):
        # 3 main stages
        cols = GridCalculator.get_columns(3, left=left, total_width=width, gap=0.60)

        # Producer
        px, pw = cols[0]
        SurfacePrimitive.render(slide, px, top, pw, height, theme)
        tb_p = slide.shapes.add_textbox(Inches(px + 0.16), Inches(top + 0.16), Inches(pw - 0.32), Inches(height - 0.32))
        tf_p = tb_p.text_frame
        TypographySystem.apply_to_paragraph(tf_p.paragraphs[0], TypographySystem.LABEL, producer_info[0], theme.text_accent)
        for it in producer_info[1][:4]:
            p = tf_p.add_paragraph()
            p.space_after = Pt(6)
            TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {it}", theme.text_muted)

        # Connector 1
        ConnectorPrimitive.render_arrow(slide, px + pw + 0.10, top + height/2 - 0.15, 0.40, 0.30, "right", theme, label="EVENT PUSH")

        # Broker / Middleware
        bx, bw = cols[1]
        SurfacePrimitive.render(slide, bx, top, bw, height, theme, border_color=theme.border_accent)
        tb_b = slide.shapes.add_textbox(Inches(bx + 0.16), Inches(top + 0.16), Inches(bw - 0.32), Inches(height - 0.32))
        tf_b = tb_b.text_frame
        TypographySystem.apply_to_paragraph(tf_b.paragraphs[0], TypographySystem.LABEL, broker_info[0], theme.border_accent)
        for it in broker_info[1][:4]:
            p = tf_b.add_paragraph()
            p.space_after = Pt(6)
            TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {it}", theme.text_primary)
        p_pat = tf_b.add_paragraph()
        TypographySystem.apply_to_paragraph(p_pat, TypographySystem.ANNOTATION, f"PATTERN: {broker_info[2]}", theme.status_success)

        # Connector 2
        ConnectorPrimitive.render_arrow(slide, bx + bw + 0.10, top + height/2 - 0.15, 0.40, 0.30, "right", theme, label="FILTERED SUB")

        # Consumer
        cx, cw = cols[2]
        SurfacePrimitive.render(slide, cx, top, cw, height, theme)
        tb_c = slide.shapes.add_textbox(Inches(cx + 0.16), Inches(top + 0.16), Inches(cw - 0.32), Inches(height - 0.32))
        tf_c = tb_c.text_frame
        TypographySystem.apply_to_paragraph(tf_c.paragraphs[0], TypographySystem.LABEL, consumer_info[0], theme.text_accent)
        for it in consumer_info[1][:4]:
            p = tf_c.add_paragraph()
            p.space_after = Pt(6)
            TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {it}", theme.text_muted)


class DataFlowArchitecturePrimitive:
    """6. End-to-end data pipeline flow (Ingest -> Stream / Batch -> Warehouse / Lake -> Analytics)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               stages: List[Tuple[str, str, List[str]]], # [(stage_name, latency_label, components)]
               theme: Theme):
        cols = GridCalculator.get_columns(len(stages), left=left, total_width=width, gap=0.35)
        for i, (s_name, latency, comps) in enumerate(stages):
            cx, cw = cols[i]
            SurfacePrimitive.render(slide, cx, top, cw, height, theme)

            tb = slide.shapes.add_textbox(Inches(cx + 0.14), Inches(top + 0.14), Inches(cw - 0.28), Inches(height - 0.28))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = 0
            p1 = tf.paragraphs[0]
            TypographySystem.apply_to_paragraph(p1, TypographySystem.LABEL, s_name, theme.text_accent)
            p2 = tf.add_paragraph()
            TypographySystem.apply_to_paragraph(p2, TypographySystem.ANNOTATION, f"SLA: {latency}", theme.status_info)
            p2.space_after = Pt(8)
            for c in comps[:4]:
                p = tf.add_paragraph()
                p.space_after = Pt(4)
                TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {c}", theme.text_muted)

            if i < len(stages) - 1:
                ConnectorPrimitive.render_arrow(slide, cx + cw + 0.05, top + height/2 - 0.12, 0.25, 0.24, "right", theme)


class SecurityBoundaryPrimitive:
    """7. Network and security zones (Public Internet -> DMZ -> Enterprise Intranet -> Air-Gapped OT)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               zones: List[Tuple[str, str, List[str]]], # [(zone_name, trust_level, safeguards)]
               theme: Theme):
        cols = GridCalculator.get_columns(len(zones), left=left, total_width=width, gap=0.25)
        for i, (z_name, trust, guards) in enumerate(zones):
            cx, cw = cols[i]
            border_c = theme.status_critical if "public" in z_name.lower() else (theme.status_warning if "dmz" in z_name.lower() else theme.status_success)
            SurfacePrimitive.render(slide, cx, top, cw, height, theme, border_color=border_c, line_width_pt=1.5)

            tb = slide.shapes.add_textbox(Inches(cx + 0.15), Inches(top + 0.15), Inches(cw - 0.30), Inches(height - 0.30))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = 0
            p1 = tf.paragraphs[0]
            TypographySystem.apply_to_paragraph(p1, TypographySystem.LABEL, z_name, border_c)
            p2 = tf.add_paragraph()
            TypographySystem.apply_to_paragraph(p2, TypographySystem.CAPTION, f"TRUST: {trust}", theme.text_muted)
            p2.space_after = Pt(10)
            for g in guards[:4]:
                p = tf.add_paragraph()
                p.space_after = Pt(6)
                TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {g}", theme.text_primary)


class DeploymentArchitecturePrimitive:
    """8. Deployment topology (Multi-Region Cloud, On-Premise Gateways, Edge Clusters)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               cloud_specs: Tuple[str, List[str]],
               edge_specs: List[Tuple[str, List[str]]],
               theme: Theme):
        # 40% Cloud Left, 60% Edge Right
        (lx, lw), (rx, rw) = GridCalculator.get_split(left_ratio=0.40, left=left, total_width=width, gap=0.30)

        # Cloud Region
        SurfacePrimitive.render(slide, lx, top, lw, height, theme, border_color=theme.border_accent)
        tb_l = slide.shapes.add_textbox(Inches(lx + 0.16), Inches(top + 0.16), Inches(lw - 0.32), Inches(height - 0.32))
        tf_l = tb_l.text_frame
        TypographySystem.apply_to_paragraph(tf_l.paragraphs[0], TypographySystem.LABEL, cloud_specs[0], theme.border_accent)
        for s in cloud_specs[1][:5]:
            p = tf_l.add_paragraph()
            p.space_after = Pt(6)
            TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {s}", theme.text_primary)

        # Edge Cluster Nodes
        sub_rows = GridCalculator.get_rows(len(edge_specs), top=top, total_height=height, gap=0.15)
        for j, (node_name, details) in enumerate(edge_specs):
            ny, nh = sub_rows[j]
            SurfacePrimitive.render(slide, rx, ny, rw, nh, theme)
            tb_r = slide.shapes.add_textbox(Inches(rx + 0.16), Inches(ny + 0.10), Inches(rw - 0.32), Inches(nh - 0.20))
            tf_r = tb_r.text_frame
            TypographySystem.apply_to_paragraph(tf_r.paragraphs[0], TypographySystem.LABEL, node_name, theme.text_accent)
            p_det = tf_r.add_paragraph()
            TypographySystem.apply_to_paragraph(p_det, TypographySystem.BODY, " • ".join(details[:3]), theme.text_muted)


class CloudEdgeArchitecturePrimitive:
    """9. Cloud Control Plane vs Edge Runtime Node synchronization."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               cloud_capabilities: List[str],
               edge_capabilities: List[str],
               sync_mechanisms: List[str],
               theme: Theme):
        cols = GridCalculator.get_columns(3, left=left, total_width=width, gap=0.25)

        # Cloud
        cx, cw = cols[0]
        SurfacePrimitive.render(slide, cx, top, cw, height, theme)
        tb_c = slide.shapes.add_textbox(Inches(cx + 0.16), Inches(top + 0.16), Inches(cw - 0.32), Inches(height - 0.32))
        tf_c = tb_c.text_frame
        TypographySystem.apply_to_paragraph(tf_c.paragraphs[0], TypographySystem.LABEL, "CLOUD CONTROL PLANE", theme.text_accent)
        for c in cloud_capabilities[:4]:
            p = tf_c.add_paragraph()
            p.space_after = Pt(6)
            TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {c}", theme.text_primary)

        # Sync
        sx, sw = cols[1]
        SurfacePrimitive.render(slide, sx, top, sw, height, theme, fill_color=theme.surface_highlight)
        tb_s = slide.shapes.add_textbox(Inches(sx + 0.16), Inches(top + 0.16), Inches(sw - 0.32), Inches(height - 0.32))
        tf_s = tb_s.text_frame
        TypographySystem.apply_to_paragraph(tf_s.paragraphs[0], TypographySystem.LABEL, "HYBRID SYNC & CACHE", theme.border_accent)
        for s in sync_mechanisms[:4]:
            p = tf_s.add_paragraph()
            p.space_after = Pt(6)
            TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {s}", theme.text_muted)

        # Edge
        ex, ew = cols[2]
        SurfacePrimitive.render(slide, ex, top, ew, height, theme)
        tb_e = slide.shapes.add_textbox(Inches(ex + 0.16), Inches(top + 0.16), Inches(ew - 0.32), Inches(height - 0.32))
        tf_e = tb_e.text_frame
        TypographySystem.apply_to_paragraph(tf_e.paragraphs[0], TypographySystem.LABEL, "EDGE FACTORY RUNTIME", theme.text_accent)
        for e in edge_capabilities[:4]:
            p = tf_e.add_paragraph()
            p.space_after = Pt(6)
            TypographySystem.apply_to_paragraph(p, TypographySystem.BODY, f"•  {e}", theme.text_primary)
