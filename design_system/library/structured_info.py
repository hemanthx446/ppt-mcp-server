"""
Structured Information Visual Primitives (Primitives 51–56).

Native, editable PowerPoint tables and matrices for enterprise data,
solution comparison, commercial models, KPI tracking, action registers, and risk matrices.
"""

from dataclasses import dataclass
from typing import List, Tuple, Optional
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from ..typography import TypographySystem
from ..color import Theme
from ..primitives import TablePrimitive


class ProfessionalTablePrimitive:
    """51. Professional Enterprise Table with column weighting, zebra shading, and typography tokens."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               headers: List[str], rows: List[List[str]],
               theme: Theme, col_weights: Optional[List[float]] = None):
        TablePrimitive.render(slide, left, top, width, height, headers, rows, theme, col_weights)


class ComparisonMatrixPrimitive:
    """52. Multi-Option Comparison Matrix (Evaluation across architectural dimensions or vendor capabilities)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               options: List[str], # E.g. ["Option A: Custom BTP", "Option B: COTS MES", "Option C: Hybrid Edge"]
               criteria: List[Tuple[str, List[str]]], # [(criterion_name, [status_or_score_per_option])]
               theme: Theme):
        headers = ["EVALUATION CRITERIA"] + options
        col_w = [1.5] + [1.0] * len(options)

        r_count = len(criteria) + 1
        c_count = len(headers)
        tbl_shape = slide.shapes.add_table(r_count, c_count, Inches(left), Inches(top), Inches(width), Inches(height))
        tbl = tbl_shape.table

        total_weight = sum(col_w)
        for c_idx, weight in enumerate(col_w):
            tbl.columns[c_idx].width = Inches((weight / total_weight) * width)

        # Header Row
        for c_idx, h_text in enumerate(headers):
            cell = tbl.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme.surface_highlight
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if c_idx > 0:
                p.alignment = PP_ALIGN.CENTER
            TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_HEADER, h_text, theme.text_accent if c_idx == 0 else theme.text_primary)

        # Body Rows
        for r_idx, (crit_name, opt_vals) in enumerate(criteria, start=1):
            row_fill = theme.surface_alt if (r_idx % 2 == 1) else theme.surface

            # Criterion Label
            c0 = tbl.cell(r_idx, 0)
            c0.fill.solid()
            c0.fill.fore_color.rgb = row_fill
            tf0 = c0.text_frame
            tf0.word_wrap = True
            TypographySystem.apply_to_paragraph(tf0.paragraphs[0], TypographySystem.TABLE_BODY, crit_name, theme.text_primary)

            for c_idx, val in enumerate(opt_vals[:len(options)], start=1):
                cell = tbl.cell(r_idx, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_fill
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER

                val_str = str(val).strip()
                # Status styling
                if any(k in val_str.upper() for k in ["FULL", "HIGH", "YES", "✓", "RECOMMENDED"]):
                    TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val_str, theme.status_success)
                elif any(k in val_str.upper() for k in ["PARTIAL", "MED", "MODERATE", "~"]):
                    TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val_str, theme.status_warning)
                elif any(k in val_str.upper() for k in ["NONE", "LOW", "NO", "✗", "POOR"]):
                    TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val_str, theme.status_critical)
                else:
                    TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_BODY, val_str, theme.text_primary)


class CommercialTablePrimitive:
    """53. Commercial / Investment Model Table with milestones, billing %, amounts, and summary totals."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               milestones: List[Tuple[str, str, str, str, str]], # [(phase_no, deliverable, timing, billing_pct, amount_str)]
               total_amount: str,
               theme: Theme):
        headers = ["PHASE / MILESTONE", "KEY DELIVERABLE & ACCEPTANCE", "TIMELINE", "BILLING %", "INVESTMENT"]
        weights = [1.2, 2.4, 0.9, 0.7, 1.0]

        r_count = len(milestones) + 2 # +1 header, +1 total row
        c_count = 5
        tbl_shape = slide.shapes.add_table(r_count, c_count, Inches(left), Inches(top), Inches(width), Inches(height))
        tbl = tbl_shape.table

        tot_w = sum(weights)
        for c_idx, w in enumerate(weights):
            tbl.columns[c_idx].width = Inches((w / tot_w) * width)

        # Header Row
        for c_idx, h_text in enumerate(headers):
            cell = tbl.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme.surface_highlight
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if c_idx >= 2:
                p.alignment = PP_ALIGN.RIGHT
            TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_HEADER, h_text, theme.text_accent)

        # Milestone Rows
        for r_idx, (p_no, deliv, tim, pct, amt) in enumerate(milestones, start=1):
            row_fill = theme.surface_alt if (r_idx % 2 == 1) else theme.surface
            row_vals = [p_no, deliv, tim, pct, amt]

            for c_idx, val in enumerate(row_vals):
                cell = tbl.cell(r_idx, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_fill
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if c_idx >= 2:
                    p.alignment = PP_ALIGN.RIGHT
                TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_BODY, val, theme.text_primary)

        # Total Row
        tot_row_idx = len(milestones) + 1
        for c_idx in range(c_count):
            cell = tbl.cell(tot_row_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme.surface_highlight
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if c_idx == 0:
                TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, "TOTAL ENGAGEMENT", theme.text_accent)
            elif c_idx == 3:
                p.alignment = PP_ALIGN.RIGHT
                TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, "100%", theme.text_accent)
            elif c_idx == 4:
                p.alignment = PP_ALIGN.RIGHT
                TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, total_amount, theme.status_success)


class KPITablePrimitive:
    """54. KPI Governance Table with baselines, targets, measurement frequency, and status pills."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               metrics: List[Tuple[str, str, str, str, str, str, str]], # [(metric_name, category, baseline, target, cadence, owner, status)]
               theme: Theme):
        headers = ["METRIC NAME", "CATEGORY", "BASELINE", "TARGET", "CADENCE", "OWNER", "STATUS"]
        weights = [1.8, 1.1, 0.8, 0.8, 0.8, 1.0, 0.9]

        r_count = len(metrics) + 1
        c_count = 7
        tbl_shape = slide.shapes.add_table(r_count, c_count, Inches(left), Inches(top), Inches(width), Inches(height))
        tbl = tbl_shape.table

        tot_w = sum(weights)
        for c_idx, w in enumerate(weights):
            tbl.columns[c_idx].width = Inches((w / tot_w) * width)

        # Header Row
        for c_idx, h_text in enumerate(headers):
            cell = tbl.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme.surface_highlight
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if c_idx in [2, 3, 4, 6]:
                p.alignment = PP_ALIGN.CENTER
            TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_HEADER, h_text, theme.text_accent)

        # Data Rows
        for r_idx, (m_name, cat, base, tgt, cad, own, stat) in enumerate(metrics, start=1):
            row_fill = theme.surface_alt if (r_idx % 2 == 1) else theme.surface
            vals = [m_name, cat, base, tgt, cad, own, stat]

            for c_idx, val in enumerate(vals):
                cell = tbl.cell(r_idx, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_fill
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if c_idx in [2, 3, 4, 6]:
                    p.alignment = PP_ALIGN.CENTER

                if c_idx == 6: # Status
                    s_up = str(val).upper()
                    if "ON TRACK" in s_up or "GOOD" in s_up or "PASS" in s_up:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, f"● {val}", theme.status_success)
                    elif "RISK" in s_up or "WARN" in s_up:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, f"▲ {val}", theme.status_warning)
                    else:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, f"■ {val}", theme.status_critical)
                elif c_idx == 0:
                    TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.text_primary)
                else:
                    TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_BODY, val, theme.text_primary)


class ActionRegisterPrimitive:
    """55. Action Item Register (Action ID, Description, Workstream, Priority, Owner, Due Date, Status)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               actions: List[Tuple[str, str, str, str, str, str, str]], # [(act_id, desc, workstream, priority, owner, due_date, status)]
               theme: Theme):
        headers = ["ID", "ACTION ITEM DESCRIPTION", "WORKSTREAM", "PRIORITY", "OWNER", "DUE DATE", "STATUS"]
        weights = [0.7, 2.5, 1.1, 0.8, 1.0, 0.9, 0.9]

        r_count = len(actions) + 1
        c_count = 7
        tbl_shape = slide.shapes.add_table(r_count, c_count, Inches(left), Inches(top), Inches(width), Inches(height))
        tbl = tbl_shape.table

        tot_w = sum(weights)
        for c_idx, w in enumerate(weights):
            tbl.columns[c_idx].width = Inches((w / tot_w) * width)

        for c_idx, h_text in enumerate(headers):
            cell = tbl.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme.surface_highlight
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if c_idx in [0, 3, 5, 6]:
                p.alignment = PP_ALIGN.CENTER
            TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_HEADER, h_text, theme.text_accent)

        for r_idx, (aid, desc, ws, prio, own, due, stat) in enumerate(actions, start=1):
            row_fill = theme.surface_alt if (r_idx % 2 == 1) else theme.surface
            vals = [aid, desc, ws, prio, own, due, stat]

            for c_idx, val in enumerate(vals):
                cell = tbl.cell(r_idx, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_fill
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if c_idx in [0, 3, 5, 6]:
                    p.alignment = PP_ALIGN.CENTER

                if c_idx == 3: # Priority
                    p_up = str(val).upper()
                    if "HIGH" in p_up or "P1" in p_up:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.status_critical)
                    elif "MED" in p_up or "P2" in p_up:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.status_warning)
                    else:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_BODY, val, theme.text_muted)
                elif c_idx == 6: # Status
                    s_up = str(val).upper()
                    if "DONE" in s_up or "CLOSED" in s_up:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.status_success)
                    elif "PROGRESS" in s_up:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.text_accent)
                    else:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_BODY, val, theme.status_warning)
                else:
                    TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_BODY, val, theme.text_primary)


class RiskRegisterPrimitive:
    """56. Enterprise Risk Register (Risk ID, Risk Event, Impact, Likelihood, Risk Score, Mitigation, Owner)."""

    @staticmethod
    def render(slide, left: float, top: float, width: float, height: float,
               risks: List[Tuple[str, str, str, str, str, str, str]], # [(risk_id, risk_event, severity, likelihood, score_level, mitigation, owner)]
               theme: Theme):
        headers = ["ID", "RISK EVENT / VULNERABILITY", "SEVERITY", "LIKELIHOOD", "SCORE", "MITIGATION STRATEGY", "OWNER"]
        weights = [0.7, 2.3, 0.8, 0.8, 0.7, 2.1, 0.9]

        r_count = len(risks) + 1
        c_count = 7
        tbl_shape = slide.shapes.add_table(r_count, c_count, Inches(left), Inches(top), Inches(width), Inches(height))
        tbl = tbl_shape.table

        tot_w = sum(weights)
        for c_idx, w in enumerate(weights):
            tbl.columns[c_idx].width = Inches((w / tot_w) * width)

        for c_idx, h_text in enumerate(headers):
            cell = tbl.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme.surface_highlight
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if c_idx in [0, 2, 3, 4]:
                p.alignment = PP_ALIGN.CENTER
            TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_HEADER, h_text, theme.text_accent)

        for r_idx, (rid, rev, sev, lik, sc, mit, own) in enumerate(risks, start=1):
            row_fill = theme.surface_alt if (r_idx % 2 == 1) else theme.surface
            vals = [rid, rev, sev, lik, sc, mit, own]

            for c_idx, val in enumerate(vals):
                cell = tbl.cell(r_idx, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_fill
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if c_idx in [0, 2, 3, 4]:
                    p.alignment = PP_ALIGN.CENTER

                if c_idx == 4: # Score level
                    sc_up = str(val).upper()
                    if "HIGH" in sc_up or "CRITICAL" in sc_up:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.status_critical)
                    elif "MED" in sc_up:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.status_warning)
                    else:
                        TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.status_success)
                elif c_idx == 1:
                    TypographySystem.apply_to_paragraph(p, TypographySystem.BODY_STRONG, val, theme.text_primary)
                else:
                    TypographySystem.apply_to_paragraph(p, TypographySystem.TABLE_BODY, val, theme.text_primary)
