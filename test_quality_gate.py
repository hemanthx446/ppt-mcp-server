import os
import sys
import unittest
from pptx import Presentation
from pptx.util import Inches

from design_system.quality_gate import (
    QualityStatus,
    QualityDimensionAudit,
    QualityGateReport,
    PresentationQualityGate
)
from design_system.narrative_intelligence import (
    NarrativeFramework,
    NarrativeIntelligenceEngine
)
from ppt_mcp_server import (
    audit_presentation_quality,
    export_presentation_with_quality_gate
)


class TestPresentationQualityGate(unittest.TestCase):
    CLEAN_DECK = "test_clean_deck.pptx"
    DEFECTIVE_DECK = "test_defective_deck.pptx"
    REMEDIATED_DECK = "test_remediated_deck.pptx"

    def tearDown(self):
        for p in [self.CLEAN_DECK, self.DEFECTIVE_DECK, self.REMEDIATED_DECK]:
            if os.path.exists(p):
                try:
                    os.remove(p)
                except OSError:
                    pass

    def test_01_clean_presentation_passes_quality_gate(self):
        """Verify that a standard presentation generated via the engine passes all quality gates."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)

        # Generate a 4-slide Solution Selling deck (Why -> What -> How -> Value)
        NarrativeIntelligenceEngine.render_narrative_deck(
            prs=prs,
            framework=NarrativeFramework.WHY_WHAT_HOW_VALUE,
            topic="Autonomous Manufacturing Execution",
            client_name="Boeing Aerospace"
        )
        prs.save(self.CLEAN_DECK)
        self.assertTrue(os.path.exists(self.CLEAN_DECK))

        # Audit clean presentation
        report = PresentationQualityGate.audit_and_remediate(Presentation(self.CLEAN_DECK), auto_remediate=False)
        self.assertEqual(report.status, QualityStatus.PASSED)
        self.assertTrue(report.is_export_authorized)
        self.assertGreaterEqual(report.overall_quality_score, 90)
        self.assertEqual(report.typography_compliance_percent, 100.0)
        self.assertLessEqual(report.card_usage_ratio, 0.25)
        self.assertEqual(len(report.violations_detected), 0)

    def test_02_detect_and_auto_remediate_defects(self):
        """Verify that quality gate flags non-Inter typography and auto-remediates it."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Inject intentional defect: non-Inter font run (Arial)
        tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(5.0), Inches(2.0))
        p = tb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = "This text intentionally uses Arial font."
        r.font.name = "Arial"

        prs.save(self.DEFECTIVE_DECK)
        self.assertTrue(os.path.exists(self.DEFECTIVE_DECK))

        # 1. Audit without auto-remediate -> should flag violations
        raw_report = PresentationQualityGate.audit_and_remediate(Presentation(self.DEFECTIVE_DECK), auto_remediate=False)
        self.assertGreater(len(raw_report.violations_detected), 0)
        self.assertIn("non-Inter", raw_report.violations_detected[0])

        # 2. Audit with auto-remediate -> should fix font to Inter and authorize export
        fixed_prs = Presentation(self.DEFECTIVE_DECK)
        fix_report = PresentationQualityGate.audit_and_remediate(fixed_prs, auto_remediate=True)
        self.assertEqual(fix_report.status, QualityStatus.REMEDIATED_AND_PASSED)
        self.assertTrue(fix_report.is_export_authorized)
        self.assertGreater(len(fix_report.auto_remediations_applied), 0)

        fixed_prs.save(self.REMEDIATED_DECK)

        # Inspect saved file to confirm font was corrected
        reopened = Presentation(self.REMEDIATED_DECK)
        for s in reopened.slides:
            for shape in s.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        for r in p.runs:
                            self.assertTrue(r.font.name.startswith("Inter"), f"Font was not remediated: {r.font.name}")

    def test_03_mcp_quality_gate_tools(self):
        """Verify FastMCP audit_presentation_quality and export_presentation_with_quality_gate tools."""
        # Create a clean deck first
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)
        NarrativeIntelligenceEngine.render_narrative_deck(
            prs=prs,
            framework=NarrativeFramework.TELL_SHOW_TELL,
            topic="Clean-Core ERP Architecture",
            client_name="Lockheed Martin"
        )
        prs.save(self.CLEAN_DECK)

        # Audit tool
        audit_res = audit_presentation_quality(self.CLEAN_DECK)
        self.assertEqual(audit_res["status"], "success")
        self.assertEqual(audit_res["gate_status"], "PASSED")
        self.assertTrue(audit_res["is_export_authorized"])
        self.assertEqual(audit_res["typography_compliance"], "100.0%")

        # Export tool
        export_res = export_presentation_with_quality_gate(
            path=self.CLEAN_DECK,
            output_path=self.REMEDIATED_DECK,
            auto_remediate=True
        )
        self.assertEqual(export_res["status"], "success")
        self.assertTrue(export_res["is_export_authorized"])
        self.assertTrue(os.path.exists(self.REMEDIATED_DECK))

    def test_04_fail_and_remediate_architecture_only_as_text(self):
        """Verify that architecture represented only as text triggers fail gate and auto-remediates to diagram."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
        t_box.text_frame.text = "Target System Architecture"
        b_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.0))
        b_box.text_frame.text = "• ERP Core: S/4HANA Clean Core\n• Integration: SAP BTP Event Mesh\n• MES Layer: Digital Manufacturing Cloud"

        # Audit without remediation -> FAIL
        report_fail = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=False)
        self.assertEqual(report_fail.status, QualityStatus.FAILED)
        self.assertFalse(report_fail.is_export_authorized)
        self.assertTrue(any("architecture only as text" in v for v in report_fail.violations_detected))

        # Audit with remediation -> REMEDIATED_AND_PASSED
        report_pass = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=True)
        self.assertEqual(report_pass.status, QualityStatus.REMEDIATED_AND_PASSED)
        self.assertTrue(report_pass.is_export_authorized)
        self.assertTrue(any("architecture" in r.lower() for r in report_pass.auto_remediations_applied))

    def test_05_fail_and_remediate_process_only_as_bullets(self):
        """Verify that process workflow represented only as bullets triggers fail gate and auto-remediates."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
        t_box.text_frame.text = "Operational Workflow Process"
        b_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.0))
        b_box.text_frame.text = "1. Order Release: Confirmed in ERP\n2. Dispatch: Machine assigned\n3. Execution: WIP tracked"

        report_fail = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=False)
        self.assertEqual(report_fail.status, QualityStatus.FAILED)
        self.assertTrue(any("process only as bullet points" in v for v in report_fail.violations_detected))

        report_pass = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=True)
        self.assertEqual(report_pass.status, QualityStatus.REMEDIATED_AND_PASSED)
        self.assertTrue(report_pass.is_export_authorized)

    def test_06_fail_and_remediate_filler_content(self):
        """Verify that slides containing filler content (Lorem Ipsum/TBD) fail gate and auto-remediate."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
        t_box.text_frame.text = "Executive Overview"
        b_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.0))
        b_box.text_frame.text = "Lorem Ipsum dolor sit amet, TBD architecture specification."

        report_fail = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=False)
        self.assertEqual(report_fail.status, QualityStatus.FAILED)
        self.assertTrue(any("filler content" in v for v in report_fail.violations_detected))

        report_pass = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=True)
        self.assertEqual(report_pass.status, QualityStatus.REMEDIATED_AND_PASSED)
        self.assertTrue(report_pass.is_export_authorized)
        self.assertNotIn("lorem ipsum", b_box.text_frame.text.lower())

    def test_07_fail_and_remediate_dense_text(self):
        """Verify that excessively dense text blocks fail gate and auto-remediate."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
        t_box.text_frame.text = "Operational Strategy"
        b_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.0))
        b_box.text_frame.text = " ".join(["operational-word"] * 90)

        report_fail = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=False)
        self.assertEqual(report_fail.status, QualityStatus.FAILED)
        self.assertTrue(any("text density is excessive" in v for v in report_fail.violations_detected))

        report_pass = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=True)
        self.assertEqual(report_pass.status, QualityStatus.REMEDIATED_AND_PASSED)
        self.assertTrue(report_pass.is_export_authorized)

    def test_08_fail_and_remediate_dashboard_only_kpis(self):
        """Verify that a dashboard with only KPI cards and no charts/tables fails gate and auto-remediates."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
        t_box.text_frame.text = "Executive Metrics Dashboard"
        kpi_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(3.0), Inches(1.5))
        kpi_box.text_frame.text = "94.2%\nOEE Efficiency"

        report_fail = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=False)
        self.assertEqual(report_fail.status, QualityStatus.FAILED)
        self.assertTrue(any("isolated KPI cards" in v for v in report_fail.violations_detected))

        report_pass = PresentationQualityGate.audit_and_remediate(prs, auto_remediate=True)
        self.assertEqual(report_pass.status, QualityStatus.REMEDIATED_AND_PASSED)
        self.assertTrue(report_pass.is_export_authorized)
        self.assertTrue(any(s.has_table for s in slide.shapes))


if __name__ == "__main__":
    unittest.main()
