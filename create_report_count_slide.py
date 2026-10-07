import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_report_count_slide(output_path="Hical_CXO_Cockpits_Report_Count_Slide.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # -------------------------------------------------------------
    # Palette & Design Tokens
    # -------------------------------------------------------------
    DARK_BG = RGBColor(11, 19, 43)           # #0B132B Deep Executive Navy
    DARK_CARD = RGBColor(21, 31, 60)         # #151F3C Navy Card
    DARK_BORDER = RGBColor(40, 56, 95)       # Navy Outline
    
    LIGHT_BG = RGBColor(248, 250, 252)       # #F8FAFC Slate Ice White
    CARD_BG = RGBColor(255, 255, 255)        # Pure White
    CARD_BORDER = RGBColor(226, 232, 240)    # Soft slate border
    CARD_BG_MUTED = RGBColor(241, 245, 249)  # Light Slate Pill
    
    SAP_BLUE = RGBColor(0, 102, 204)         # #0066CC Enterprise SAP Blue
    CYAN_ACCENT = RGBColor(0, 180, 216)      # #00B4D8 Tech Cyan
    GOLD_ACCENT = RGBColor(217, 119, 6)      # #D97706 Lumbini Amber/Gold
    GOLD_LIGHT_BG = RGBColor(254, 243, 199)  # Gold Tint
    
    EMERALD_GREEN = RGBColor(16, 185, 129)   # #10B981 Success / On-Track
    GREEN_LIGHT_BG = RGBColor(236, 253, 245) # Soft Green Tint
    AMBER_WARN = RGBColor(245, 158, 11)      # #F59E0B Warning
    AMBER_LIGHT_BG = RGBColor(255, 251, 235) # Soft Amber Tint
    
    PURPLE_ACCENT = RGBColor(124, 58, 237)   # AI / Costing Purple
    PURPLE_LIGHT_BG = RGBColor(245, 243, 255)
    
    TEXT_DARK = RGBColor(15, 23, 42)         # #0F172A Primary Dark
    TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B Secondary Slate
    TEXT_LIGHT = RGBColor(248, 250, 252)     # #F8FAFC Off-White
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()

    # Header Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(5.0), Inches(0.25))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(238, 242, 255)
    badge.line.color.rgb = SAP_BLUE
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = Inches(0.12)
    tf_b.margin_top = Inches(0.01)
    p_b = tf_b.paragraphs[0]
    p_b.text = "DELIVERABLE MANIFEST | 18 TOTAL EXECUTIVE REPORTS".upper()
    p_b.font.name = "Arial"
    p_b.font.size = Pt(8.5)
    p_b.font.bold = True
    p_b.font.color.rgb = SAP_BLUE

    # Title & Subtitle
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.62), Inches(11.733), Inches(0.40))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "Cockpit Report Inventory: Itemized Deliverables by Cockpit"
    p.font.name = "Arial"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    tx_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.02), Inches(11.733), Inches(0.26))
    tf_sub = tx_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
    ps = tf_sub.paragraphs[0]
    ps.text = "Itemized Architecture: Exactly 6 Dedicated Production Reports Per Cockpit Across Finance, Sourcing, and Manufacturing"
    ps.font.name = "Arial"
    ps.font.size = Pt(10)
    ps.font.italic = True
    ps.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # 4 Quick Metric Summary Badges (Top Row)
    # -------------------------------------------------------------
    kpis = [
        ("TOTAL COCKPIT SUITES", "3 Executive Suites", SAP_BLUE, RGBColor(238, 242, 255)),
        ("TOTAL DELIVERABLE REPORTS", "18 Production Reports", GOLD_ACCENT, GOLD_LIGHT_BG),
        ("REPORTS PER COCKPIT", "6 Reports Each (4 Core + 2 Strategic)", EMERALD_GREEN, GREEN_LIGHT_BG),
        ("UNIT COMMERCIAL ALIGNMENT", "~₹1.67L / Report (₹10.0L / Suite)", PURPLE_ACCENT, PURPLE_LIGHT_BG),
    ]
    
    kpi_w = 2.75
    kpi_gap = 0.244
    for i, (k_lbl, k_val, k_clr, k_bg) in enumerate(kpis):
        k_left = 0.8 + i * (kpi_w + kpi_gap)
        card_k = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(k_left), Inches(1.36), Inches(kpi_w), Inches(0.56))
        card_k.fill.solid()
        card_k.fill.fore_color.rgb = k_bg
        card_k.line.color.rgb = k_clr
        card_k.line.width = Pt(1)
        tf_k = card_k.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_right = Inches(0.12)
        tf_k.margin_top = Inches(0.04)
        
        p1 = tf_k.paragraphs[0]
        p1.text = k_lbl
        p1.font.name = "Arial"
        p1.font.size = Pt(7.2)
        p1.font.bold = True
        p1.font.color.rgb = k_clr
        
        p2 = tf_k.add_paragraph()
        p2.text = k_val
        p2.font.name = "Arial"
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # 3 Main Cockpit Columns (6 Reports each)
    # -------------------------------------------------------------
    cockpits_data = [
        {
            "header": "COCKPIT 1: CASH FLOW & FULFILLMENT",
            "badge": "6 Dedicated Reports (4 Core + 2 Strategic) | Finance & SCM",
            "border": SAP_BLUE,
            "header_bg": RGBColor(238, 245, 255),
            "reports": [
                ("1.1 Scheduled Shipments & Dispatch Pipeline", "Line-item tracking of customer orders vs. dispatch dates, OTIF %, and backlog velocity."),
                ("1.2 Expected Inbound Materials & Bay Staging", "Open vendor PO lines, in-transit ASN visibility, and dock arrival staging readiness."),
                ("1.3 Supplier Payables (AP Due) & Liability Aging", "Payables aging (30/60/90 days), cash discount capture, and vendor credit exposure."),
                ("1.4 Customer Receivables (AR Due) & Cash Realization", "Open AR aging, customer tier DSO optimization, and unbilled defense milestone claims."),
                ("1.5 [Strategic] Net Liquidity Runway Simulator", "Multi-scenario cash burn forecasting (AR vs. AP) under 30/60/90-day stress to safeguard covenants."),
                ("1.6 [Strategic] Late Delivery Risk & LD Penalty Tracker", "Algorithmic penalty scoring correlating routing delays with contract LD clauses (protecting 1–3% margin).")
            ]
        },
        {
            "header": "COCKPIT 2: PROCUREMENT 52-WEEK PIPELINE",
            "badge": "6 Dedicated Reports (4 Core + 2 Strategic) | Sourcing & Job-Work",
            "border": CYAN_ACCENT,
            "header_bg": RGBColor(235, 251, 255),
            "reports": [
                ("2.1 52-Week Sourcing Pipeline & PO Releases", "Forward rolling demand vs. confirmed POs, uncommitted PR gaps, and purchase budget adherence."),
                ("2.2 Component-to-Part Traceability & Program Lineage", "Multi-tier BOM lineage mapping open PO lines directly to SFG/FG and end-customer defense programs."),
                ("2.3 Vendor Concentration & Lead-Time Volatility", "Herfindahl spend concentration (HHI), quoted vs. actual lead-time drift, and supplier scorecards."),
                ("2.4 Subcontractor Job-Work & Special Stock 'O'", "Materials issued to platers/heattreaters (541/542), vendor possession aging, outside processing valuation."),
                ("2.5 [Strategic] Subcontractor TAT & Stage Yield/Scrap", "Outside cycle times (TAT) and stage yield loss at external job-workers to stop untracked material bleed."),
                ("2.6 [Strategic] Critical Path Component Shortage Predictor", "Predictive heuristic flagging items with lead times >16 wks before buffer stock breaches safety thresholds.")
            ]
        },
        {
            "header": "COCKPIT 3: MANUFACTURING & BATCH COSTING",
            "badge": "6 Dedicated Reports (4 Core + 2 Strategic) | Operations & Quality",
            "border": PURPLE_ACCENT,
            "header_bg": RGBColor(248, 244, 255),
            "reports": [
                ("3.1 Pre-Release MRP Kitting Readiness Gate", "Automated 100% component availability validation prior to releasing production orders (RESB/AFPO)."),
                ("3.2 Batch Costing Variance Engine (Actual vs. Std BOM)", "Material price/usage variance, confirmed machine/labor hours (AFRU) vs. standard routing."),
                ("3.3 Work-Center Utilization (OEE) & Line Throughput", "CNC machine capacity loading, queue bottleneck telemetry across winding, machining, and testing."),
                ("3.4 Shift Operations & Daily Floor Telemetry", "Hourly operator confirmation velocity, machine setup vs. run ratios, and QA dispatch readiness."),
                ("3.5 [Strategic] Cost of Poor Quality (COPQ) Scrap Leakage", "Inspection defect logging (551/552) attributed to work centers, tool sets, and shifts (2–4% margin recovery)."),
                ("3.6 [Strategic] WIP Valuation & In-Process Ageing Matrix", "Financial valuation of stalled semi-finished stock categorized by dwelling age (<15, 16–30, 31–60, >60 days).")
            ]
        }
    ]

    col_w = 3.75
    col_gap = 0.24
    col_top = 2.02
    col_h = 4.45

    for c_idx, data in enumerate(cockpits_data):
        c_left = 0.8 + c_idx * (col_w + col_gap)
        
        # Outer Card Container
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(col_top), Inches(col_w), Inches(col_h))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = data["border"]
        card.line.width = Pt(1.5)
        
        # Header banner inside column
        hb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left + 0.06), Inches(col_top + 0.06), Inches(col_w - 0.12), Inches(0.52))
        hb.fill.solid()
        hb.fill.fore_color.rgb = data["header_bg"]
        hb.line.color.rgb = data["border"]
        hb.line.width = Pt(1)
        tf_hb = hb.text_frame
        tf_hb.word_wrap = True
        tf_hb.margin_left = tf_hb.margin_right = Inches(0.06)
        tf_hb.margin_top = Inches(0.03)
        
        p_h1 = tf_hb.paragraphs[0]
        p_h1.text = data["header"]
        p_h1.font.name = "Arial"
        p_h1.font.size = Pt(8.2)
        p_h1.font.bold = True
        p_h1.font.color.rgb = data["border"]
        
        p_h2 = tf_hb.add_paragraph()
        p_h2.text = data["badge"]
        p_h2.font.name = "Arial"
        p_h2.font.size = Pt(7.0)
        p_h2.font.bold = True
        p_h2.font.color.rgb = TEXT_DARK

        # Report Items Box
        rep_box = slide.shapes.add_textbox(Inches(c_left + 0.05), Inches(col_top + 0.64), Inches(col_w - 0.10), Inches(3.80))
        tf_rep = rep_box.text_frame
        tf_rep.word_wrap = True
        tf_rep.margin_left = tf_rep.margin_right = tf_rep.margin_top = tf_rep.margin_bottom = 0

        for r_num, (r_name, r_desc) in enumerate(data["reports"]):
            p_r = tf_rep.paragraphs[0] if r_num == 0 else tf_rep.add_paragraph()
            p_r.space_before = Pt(3) if r_num > 0 else Pt(0)
            p_r.space_after = Pt(2)
            
            is_strat = "[Strategic]" in r_name
            
            run_title = p_r.add_run()
            run_title.text = "• " + r_name + "\n"
            run_title.font.name = "Arial"
            run_title.font.size = Pt(7.3)
            run_title.font.bold = True
            run_title.font.color.rgb = data["border"] if not is_strat else GOLD_ACCENT
            
            run_desc = p_r.add_run()
            run_desc.text = "   " + r_desc
            run_desc.font.name = "Arial"
            run_desc.font.size = Pt(6.8)
            run_desc.font.bold = False
            run_desc.font.color.rgb = TEXT_MUTED if not is_strat else TEXT_DARK

    # -------------------------------------------------------------
    # Bottom Architecture & SLA Guarantee Banner
    # -------------------------------------------------------------
    banner_b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.56), Inches(11.733), Inches(0.38))
    banner_b.fill.solid()
    banner_b.fill.fore_color.rgb = RGBColor(241, 245, 249)
    banner_b.line.color.rgb = SAP_BLUE
    banner_b.line.width = Pt(1)
    
    tf_bb = banner_b.text_frame
    tf_bb.word_wrap = True
    tf_bb.margin_left = tf_bb.margin_right = Inches(0.15)
    tf_bb.margin_top = Inches(0.04)
    p_bb = tf_bb.paragraphs[0]
    
    r_b1 = p_bb.add_run()
    r_b1.text = "COMPLETE 4-LAYER TURNKEY STACK PER REPORT: "
    r_b1.font.bold = True
    r_b1.font.size = Pt(7.5)
    r_b1.font.color.rgb = SAP_BLUE
    
    r_b2 = p_bb.add_run()
    r_b2.text = "All 18 reports include: ① SAP HANA Push-Down Calculation Views  |  ② Power BI DirectQuery Interactive Semantic Models  |  ③ ACDOCA / MATDOC GL Reconciliation  |  ④ Dynamic Role-Based Security (RLS) & UAT Sign-Off"
    r_b2.font.bold = False
    r_b2.font.size = Pt(7.2)
    r_b2.font.color.rgb = TEXT_DARK

    # Footer
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.02), Inches(11.733), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = CARD_BORDER
    line.line.fill.background()
    
    tx_f = slide.shapes.add_textbox(Inches(0.8), Inches(7.06), Inches(8.5), Inches(0.28))
    tf_f = tx_f.text_frame
    tf_f.margin_left = tf_f.margin_top = 0
    pf = tf_f.paragraphs[0]
    pf.text = "Lumbini Elite Solutions  |  Hical Technologies Strategic CXO Dashboards Proposal  |  Confidential"
    pf.font.name = "Arial"
    pf.font.size = Pt(8.5)
    pf.font.color.rgb = TEXT_MUTED
    
    tx_fn = slide.shapes.add_textbox(Inches(11.533), Inches(7.06), Inches(1.0), Inches(0.28))
    tf_fn = tx_fn.text_frame
    tf_fn.margin_right = tf_fn.margin_top = 0
    pfn = tf_fn.paragraphs[0]
    pfn.alignment = PP_ALIGN.RIGHT
    pfn.text = "07 / 13"
    pfn.font.name = "Arial"
    pfn.font.size = Pt(8.5)
    pfn.font.bold = True
    pfn.font.color.rgb = SAP_BLUE

    prs.save(output_path)
    print(f"Standalone slide presentation successfully generated: {output_path}")

if __name__ == "__main__":
    build_report_count_slide()
