import os
import sys
import unittest
from pptx import Presentation
from pptx.util import Inches

from design_system.dashboard import (
    DashboardArchetype,
    DashboardQuestionClassifier,
    ExecutiveCockpitSpec,
    ExecutiveDashboardComposer
)
from ppt_mcp_server import (
    list_dashboard_archetypes,
    render_executive_dashboard
)


class TestExecutiveDashboard(unittest.TestCase):
    OUTPUT_DECK = "test_executive_dashboards_deck.pptx"
    SINGLE_DECK = "test_single_dashboard.pptx"

    CANONICAL_QUESTIONS = [
        "Are we shipping what we promised?",
        "Where is production constrained?",
        "What is driving poor quality?",
        "Where is working capital trapped?",
        "Which suppliers are creating risk?",
        "How much is quality costing us?"
    ]

    EXPECTED_ARCHETYPES = [
        DashboardArchetype.SHIPPING_PROMISE,
        DashboardArchetype.PRODUCTION_CONSTRAINTS,
        DashboardArchetype.QUALITY_DRIVERS,
        DashboardArchetype.WORKING_CAPITAL_WIP,
        DashboardArchetype.SUPPLIER_RISK,
        DashboardArchetype.COPQ_FINANCIAL
    ]

    def tearDown(self):
        for path in [self.OUTPUT_DECK, self.SINGLE_DECK]:
            if os.path.exists(path):
                try:
                    os.remove(path)
                except OSError:
                    pass

    def test_01_question_classifier(self):
        """Verify all 6 canonical questions classify accurately into their archetypes."""
        for q, expected in zip(self.CANONICAL_QUESTIONS, self.EXPECTED_ARCHETYPES):
            detected = DashboardQuestionClassifier.classify(q)
            self.assertEqual(detected, expected, f"Failed for question: '{q}' -> got {detected}, expected {expected}")

        # Also test variations
        self.assertEqual(DashboardQuestionClassifier.classify("What is our customer OTIF delivery rate?"), DashboardArchetype.SHIPPING_PROMISE)
        self.assertEqual(DashboardQuestionClassifier.classify("Where are the plant bottlenecks and capacity overloads?"), DashboardArchetype.PRODUCTION_CONSTRAINTS)
        self.assertEqual(DashboardQuestionClassifier.classify("Why is scrap and rework spiking on Line 2?"), DashboardArchetype.QUALITY_DRIVERS)
        self.assertEqual(DashboardQuestionClassifier.classify("How much WIP inventory aging do we have?"), DashboardArchetype.WORKING_CAPITAL_WIP)
        self.assertEqual(DashboardQuestionClassifier.classify("Which vendor lead time drift creates stockout risk?"), DashboardArchetype.SUPPLIER_RISK)
        self.assertEqual(DashboardQuestionClassifier.classify("What is our monthly COPQ and cost of poor quality?"), DashboardArchetype.COPQ_FINANCIAL)
        self.assertEqual(DashboardQuestionClassifier.classify("How do we optimize overall operational margin?"), DashboardArchetype.CUSTOM_COCKPIT)

    def test_02_build_from_question_specifications(self):
        """Verify that specifications synthesized from questions contain all 4 analytical zones."""
        for q in self.CANONICAL_QUESTIONS:
            spec = ExecutiveDashboardComposer.build_from_question(q)
            self.assertIsNotNone(spec.kpi_strip, f"KPI strip missing for {q}")
            self.assertGreaterEqual(len(spec.kpi_strip), 3, f"KPI strip too small for {q}")
            self.assertIsNotNone(spec.trend_periods, f"Trend periods missing for {q}")
            self.assertIsNotNone(spec.trend_series, f"Trend series missing for {q}")
            self.assertIsNotNone(spec.breakdown_data, f"Breakdown data missing for {q}")
            self.assertIsNotNone(spec.action_table_headers, f"Action table headers missing for {q}")
            self.assertIsNotNone(spec.action_table_rows, f"Action table rows missing for {q}")

    def test_03_render_all_canonical_cockpits_deck(self):
        """Render a 6-slide deck covering all canonical questions and verify typography and structure."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)

        for i, q in enumerate(self.CANONICAL_QUESTIONS):
            spec = ExecutiveDashboardComposer.build_from_question(q, client_name="Boeing Aerospace")
            spec.is_dark = (i % 2 == 1)  # alternate light and dark themes
            slide = ExecutiveDashboardComposer.compose_and_render(prs, spec)
            self.assertIsNotNone(slide)

        self.assertEqual(len(prs.slides), 6)
        prs.save(self.OUTPUT_DECK)
        self.assertTrue(os.path.exists(self.OUTPUT_DECK))

        # Reopen saved presentation and audit contents
        saved_prs = Presentation(self.OUTPUT_DECK)
        self.assertEqual(len(saved_prs.slides), 6)

        total_text_runs = 0
        total_tables = 0
        total_charts = 0

        for slide_idx, slide in enumerate(saved_prs.slides):
            has_table = False
            has_chart = False
            found_zone_labels = []

            for shape in slide.shapes:
                if shape.has_table:
                    has_table = True
                    total_tables += 1
                if shape.has_chart:
                    has_chart = True
                    total_charts += 1

                # Check fonts across all text runs
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        p_text = p.text.strip()
                        if p_text.startswith("1. WHAT CHANGED?") or p_text.startswith("2. WHY DOES IT MATTER?") or \
                           p_text.startswith("3. WHERE IS THE PROBLEM?") or p_text.startswith("4. WHAT SHOULD WE ACT ON?"):
                            found_zone_labels.append(p_text[:16])

                        for r in p.runs:
                            if r.font and r.font.name:
                                total_text_runs += 1
                                self.assertTrue(
                                    r.font.name.startswith("Inter"),
                                    f"Non-Inter font '{r.font.name}' found on slide {slide_idx}: '{r.text}'"
                                )

            self.assertTrue(has_table, f"Slide {slide_idx} should contain an action table")
            self.assertTrue(has_chart, f"Slide {slide_idx} should contain a trend chart")
            self.assertGreaterEqual(len(found_zone_labels), 3, f"Slide {slide_idx} missing question zone labels")

        self.assertGreater(total_text_runs, 100, "Should have audited over 100 text runs")
        self.assertEqual(total_tables, 6, "Each of the 6 slides must have an action table")
        self.assertEqual(total_charts, 6, "Each of the 6 slides must have a native chart")

    def test_04_mcp_tool_list_archetypes(self):
        """Test list_dashboard_archetypes MCP tool."""
        res = list_dashboard_archetypes()
        self.assertIn("total_archetypes", res)
        self.assertEqual(res["total_archetypes"], 7)
        self.assertEqual(len(res["archetypes"]), 7)
        arch_names = [a["archetype"] for a in res["archetypes"]]
        self.assertIn("SHIPPING_PROMISE", arch_names)
        self.assertIn("PRODUCTION_CONSTRAINTS", arch_names)
        self.assertIn("QUALITY_DRIVERS", arch_names)
        self.assertIn("WORKING_CAPITAL_WIP", arch_names)
        self.assertIn("SUPPLIER_RISK", arch_names)
        self.assertIn("COPQ_FINANCIAL", arch_names)
        self.assertIn("CUSTOM_COCKPIT", arch_names)

    def test_05_mcp_tool_render_executive_dashboard(self):
        """Test render_executive_dashboard MCP tool end-to-end."""
        res = render_executive_dashboard(
            path=self.SINGLE_DECK,
            business_question="Are we shipping what we promised?",
            key_takeaway="Tier-1 wiring harness stockout resolved; OTIF recovering to 92% next cycle.",
            client_name="Lockheed Martin",
            is_dark=False
        )
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["archetype"], "SHIPPING_PROMISE")
        self.assertEqual(res["total_slides"], 1)
        self.assertEqual(len(res["analytical_zones_rendered"]), 4)
        self.assertTrue(os.path.exists(self.SINGLE_DECK))

        # Append a second slide for production constraints
        res2 = render_executive_dashboard(
            path=self.SINGLE_DECK,
            business_question="Where is production constrained?",
            client_name="Lockheed Martin",
            is_dark=True
        )
        self.assertEqual(res2["status"], "success")
        self.assertEqual(res2["archetype"], "PRODUCTION_CONSTRAINTS")
        self.assertEqual(res2["total_slides"], 2)


if __name__ == "__main__":
    unittest.main()
