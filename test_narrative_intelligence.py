import os
import sys
import unittest
from pptx import Presentation
from pptx.util import Inches

from design_system.narrative_intelligence import (
    NarrativeFramework,
    AudienceCognitiveObjective,
    NarrativeSlideIntent,
    NarrativeIntelligenceEngine
)
from ppt_mcp_server import (
    list_narrative_frameworks,
    generate_narrative_presentation
)


class TestNarrativeIntelligence(unittest.TestCase):
    TELL_SHOW_TELL_DECK = "test_tell_show_tell_deck.pptx"
    PRIDE_PURPOSE_DECK = "test_pride_purpose_deck.pptx"
    SOLUTION_SELLING_DECK = "test_solution_selling_deck.pptx"
    DEMO_DECK = "test_demo_deck.pptx"
    MCP_DECK = "test_mcp_narrative_deck.pptx"

    def tearDown(self):
        for p in [self.TELL_SHOW_TELL_DECK, self.PRIDE_PURPOSE_DECK, self.SOLUTION_SELLING_DECK, self.DEMO_DECK, self.MCP_DECK]:
            if os.path.exists(p):
                try:
                    os.remove(p)
                except OSError:
                    pass

    def test_01_narrative_planning_and_cognitive_objectives(self):
        """Verify that all 4 frameworks generate coherent arcs with internal cognitive reasoning."""
        frameworks = [
            (NarrativeFramework.TELL_SHOW_TELL, 3),
            (NarrativeFramework.PRIDE_PURPOSE_DESTINATION, 6),
            (NarrativeFramework.WHY_WHAT_HOW_VALUE, 4),
            (NarrativeFramework.EXECUTIVE_DEMONSTRATION, 6)
        ]

        for fw, expected_slides in frameworks:
            intents = NarrativeIntelligenceEngine.plan_narrative_deck(
                framework=fw,
                topic="Autonomous Shop Floor Intelligence",
                client_name="Boeing Aerospace"
            )
            self.assertEqual(len(intents), expected_slides, f"Framework {fw.name} has incorrect slide count")

            for intent in intents:
                self.assertIsNotNone(intent.slide_title)
                self.assertIsNotNone(intent.slide_subtitle)
                self.assertIsNotNone(intent.cognitive_objective)
                # Verify cognitive objectives are rich and complete
                cog = intent.cognitive_objective
                self.assertTrue(len(cog.understand) > 5)
                self.assertTrue(len(cog.believe) > 5)
                self.assertTrue(len(cog.decide) > 5)
                self.assertTrue(len(cog.do) > 5)

    def test_02_render_tell_show_tell_deck(self):
        """Render TELL -> SHOW -> TELL and verify no internal reasoning is exposed on slides."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)

        intents = NarrativeIntelligenceEngine.render_narrative_deck(
            prs=prs,
            framework=NarrativeFramework.TELL_SHOW_TELL,
            topic="Clean-Core SAP S/4HANA & MES Integration",
            client_name="Lockheed Martin"
        )
        self.assertEqual(len(prs.slides), 3)
        prs.save(self.TELL_SHOW_TELL_DECK)
        self.assertTrue(os.path.exists(self.TELL_SHOW_TELL_DECK))

        # Reopen and audit slides
        saved_prs = Presentation(self.TELL_SHOW_TELL_DECK)
        forbidden_meta_words = ["cognitive_objective", "audience needs to", "decide or do", "internal reasoning"]

        total_text_runs = 0
        for slide_idx, slide in enumerate(saved_prs.slides):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        text_lower = p.text.lower()
                        for f_word in forbidden_meta_words:
                            self.assertNotIn(f_word, text_lower, f"Internal reasoning leaked on slide {slide_idx}: '{p.text}'")

                        for r in p.runs:
                            if r.font and r.font.name:
                                total_text_runs += 1
                                self.assertTrue(
                                    r.font.name.startswith("Inter"),
                                    f"Non-Inter font '{r.font.name}' found on slide {slide_idx}: '{r.text}'"
                                )

        self.assertGreater(total_text_runs, 50, "Must audit text runs across all 3 slides")

    def test_03_mcp_narrative_tools(self):
        """Verify FastMCP list_narrative_frameworks and generate_narrative_presentation tools."""
        list_res = list_narrative_frameworks()
        self.assertEqual(list_res["total_frameworks"], 4)
        fw_ids = [f["framework_id"] for f in list_res["frameworks"]]
        self.assertIn("TELL_SHOW_TELL", fw_ids)
        self.assertIn("PRIDE_PURPOSE_DESTINATION", fw_ids)
        self.assertIn("WHY_WHAT_HOW_VALUE", fw_ids)
        self.assertIn("EXECUTIVE_DEMONSTRATION", fw_ids)

        # Generate Solution Selling (4 slides)
        gen_res = generate_narrative_presentation(
            path=self.MCP_DECK,
            framework="WHY_WHAT_HOW_VALUE",
            topic="Plant Digital Transformation",
            client_name="Raytheon Defense"
        )
        self.assertEqual(gen_res["status"], "success")
        self.assertEqual(gen_res["total_slides"], 4)
        self.assertEqual(gen_res["framework"], "WHY_WHAT_HOW_VALUE")
        self.assertTrue(os.path.exists(self.MCP_DECK))


if __name__ == "__main__":
    unittest.main()
