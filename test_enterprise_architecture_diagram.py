import os
import sys
import unittest
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

from design_system.architecture_engine import (
    ArchElementType,
    FlowDirection,
    ArchElement,
    ArchFlow,
    ArchLayer,
    ArchAnnotation,
    ArchitectureDiagramSpec,
    EnterpriseArchitectureComposer,
    ArchitectureBlueprintFactory
)
from ppt_mcp_server import (
    list_architecture_templates,
    render_enterprise_architecture_diagram
)


class TestEnterpriseArchitectureDiagram(unittest.TestCase):
    OUTPUT_DECK = "test_architecture_engine_deck.pptx"
    MCP_OUTPUT_DECK = "test_mcp_arch_deck.pptx"

    def tearDown(self):
        for p in [self.OUTPUT_DECK, self.MCP_OUTPUT_DECK]:
            if os.path.exists(p):
                try:
                    os.remove(p)
                except OSError:
                    pass

    def test_01_blueprint_factory_specs(self):
        """Verify pre-built canonical architecture blueprints."""
        clean_core = ArchitectureBlueprintFactory.clean_core_sap_to_shopfloor()
        self.assertEqual(len(clean_core.layers), 5)
        self.assertIn("ENTERPRISE / ERP", clean_core.layers[0].name)
        self.assertIn("INTEGRATION", clean_core.layers[1].name)
        self.assertIn("MES", clean_core.layers[2].name)
        self.assertIn("EDGE", clean_core.layers[3].name)
        self.assertIn("SHOP FLOOR", clean_core.layers[4].name)
        self.assertGreaterEqual(len(clean_core.annotations), 4)

        # Verify elements contain diverse types: Data stores, APIs, Devices, Actors
        elem_types = []
        for l in clean_core.layers:
            for e in l.elements:
                elem_types.append(e.element_type)
        self.assertIn(ArchElementType.SYSTEM, elem_types)
        self.assertIn(ArchElementType.DATA_STORE, elem_types)
        self.assertIn(ArchElementType.INTERFACE_API, elem_types)
        self.assertIn(ArchElementType.DEVICE, elem_types)
        self.assertIn(ArchElementType.ACTOR, elem_types)

        # Zero trust blueprint
        zero_trust = ArchitectureBlueprintFactory.zero_trust_ot_security_perimeter()
        self.assertEqual(len(zero_trust.layers), 4)
        self.assertGreaterEqual(len(zero_trust.annotations), 3)

    def test_02_render_and_shape_diversity(self):
        """Render architecture blueprint and verify native PowerPoint shapes."""
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.500)

        # Slide 1: Clean Core ISA-95 (Dark)
        spec1 = ArchitectureBlueprintFactory.clean_core_sap_to_shopfloor(client_name="Boeing Aerospace")
        slide1 = EnterpriseArchitectureComposer.compose_and_render(prs, spec1)
        self.assertIsNotNone(slide1)

        # Slide 2: Zero-Trust Security Perimeter (Light)
        spec2 = ArchitectureBlueprintFactory.zero_trust_ot_security_perimeter(client_name="Raytheon Defense")
        spec2.is_dark = False
        slide2 = EnterpriseArchitectureComposer.compose_and_render(prs, spec2)
        self.assertIsNotNone(slide2)

        prs.save(self.OUTPUT_DECK)
        self.assertTrue(os.path.exists(self.OUTPUT_DECK))

        # Reopen and inspect shapes
        saved_prs = Presentation(self.OUTPUT_DECK)
        self.assertEqual(len(saved_prs.slides), 2)

        found_cylinders = 0
        found_hexagons = 0
        found_ovals = 0
        found_arrows = 0
        found_rectangles = 0
        total_text_runs = 0
        protocol_labels_found = []

        for slide_idx, slide in enumerate(saved_prs.slides):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        if "INTERFACE:" in p.text or "PAYLOAD:" in p.text:
                            protocol_labels_found.append(p.text)
                        for r in p.runs:
                            if r.font and r.font.name:
                                total_text_runs += 1
                                self.assertTrue(
                                    r.font.name.startswith("Inter"),
                                    f"Non-Inter font '{r.font.name}' found on slide {slide_idx}: '{r.text}'"
                                )

                # Inspect shape types (auto_shape_type)
                try:
                    stype = shape.auto_shape_type
                    if stype == MSO_SHAPE.CAN:
                        found_cylinders += 1
                    elif stype == MSO_SHAPE.HEXAGON:
                        found_hexagons += 1
                    elif stype == MSO_SHAPE.OVAL:
                        found_ovals += 1
                    elif stype in (MSO_SHAPE.DOWN_ARROW, MSO_SHAPE.UP_DOWN_ARROW):
                        found_arrows += 1
                    elif stype == MSO_SHAPE.RECTANGLE:
                        found_rectangles += 1
                except (AttributeError, ValueError):
                    pass

        # Assert shape diversity: not just cards!
        self.assertGreater(found_cylinders, 0, "Architecture diagram must contain data store cylinders (MSO_SHAPE.CAN)")
        self.assertGreater(found_hexagons, 0, "Architecture diagram must contain interface hexagons (MSO_SHAPE.HEXAGON)")
        self.assertGreater(found_arrows, 0, "Architecture diagram must contain inter-layer flow arrows")
        self.assertGreater(found_rectangles, 10, "Architecture diagram must contain layer and system boundaries")
        self.assertGreater(len(protocol_labels_found), 4, "Must contain explicit INTERFACE and PAYLOAD movement labels")
        self.assertGreater(total_text_runs, 80, "Must audit text runs for typography")

    def test_03_mcp_tools_architecture(self):
        """Verify FastMCP list_architecture_templates and render_enterprise_architecture_diagram tools."""
        templates_res = list_architecture_templates()
        self.assertIn("templates", templates_res)
        self.assertGreaterEqual(templates_res["total_templates"], 3)

        # Test render tool with Clean Core template
        render_res = render_enterprise_architecture_diagram(
            path=self.MCP_OUTPUT_DECK,
            title="Clean-Core S/4HANA to MES Integration Architecture",
            template_id="CLEAN_CORE_SAP_TO_SHOPFLOOR",
            client_name="Lockheed Martin",
            is_dark=True
        )
        self.assertEqual(render_res["status"], "success")
        self.assertEqual(render_res["total_layers"], 5)
        self.assertEqual(render_res["total_slides"], 1)
        self.assertTrue(os.path.exists(self.MCP_OUTPUT_DECK))

        # Test append custom architecture slide
        custom_layers = [
            {
                "name": "CLOUD ANALYTICS & PLANNING",
                "level_tag": "TIER 1 • CLOUD",
                "trust_zone": "Corporate VPC",
                "elements": [
                    {"name": "SAP IBP Supply Chain", "type": "SYSTEM", "role": "Demand & Supply Planning"},
                    {"name": "Global Inventory Lake", "type": "DATA_STORE", "tech": "Snowflake / Delta Lake"}
                ],
                "flow_to_next": {
                    "protocol": "REST / JSON & Webhook",
                    "data_object": "ApprovedProductionSchedule_v1",
                    "direction": "DOWN",
                    "badge": "Nightly Batch"
                }
            },
            {
                "name": "PLANT SCHEDULING & DISPATCH",
                "level_tag": "TIER 2 • PLANT",
                "trust_zone": "Plant Network",
                "elements": [
                    {"name": "Capacity Dispatch Engine", "type": "SYSTEM", "role": "Finite Bottleneck Scheduling"},
                    {"name": "Line Lead Dispatcher", "type": "ACTOR", "role": "Approval Persona"}
                ]
            }
        ]
        custom_annotations = [
            {"category": "Finite Capacity", "principle": "Dispatch schedules account for tool changeovers and operator constraints."}
        ]

        render_res2 = render_enterprise_architecture_diagram(
            path=self.MCP_OUTPUT_DECK,
            title="Custom Plant Capacity & Scheduling Architecture",
            template_id="CUSTOM_ARCHITECTURE",
            client_name="Lockheed Martin",
            is_dark=False,
            custom_layers=custom_layers,
            custom_annotations=custom_annotations
        )
        self.assertEqual(render_res2["status"], "success")
        self.assertEqual(render_res2["total_slides"], 2)
        self.assertEqual(render_res2["total_layers"], 2)


if __name__ == "__main__":
    unittest.main()
