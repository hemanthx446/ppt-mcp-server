import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_presentation(output_path="CEO_Speech_Deck.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # -------------------------------------------------------------
    # Palette & Design Tokens
    # -------------------------------------------------------------
    DARK_BG = RGBColor(18, 18, 18)           # #121212 Dark Charcoal
    DARK_CARD = RGBColor(28, 28, 30)         # #1C1C1E Dark Slate Card
    DARK_BORDER = RGBColor(48, 48, 52)       # Subtle card outline
    
    LIGHT_BG = RGBColor(248, 249, 250)       # #F8F9FA Ice White
    CARD_BG = RGBColor(255, 255, 255)        # Crisp White Container
    BORDER_GRAY = RGBColor(226, 232, 240)    # Soft slate border
    
    PRIMARY_RED = RGBColor(227, 6, 19)       # #E30613 Bold Strategic Red
    ACCENT_RED_BG = RGBColor(254, 242, 242)  # Light Red Wash for badges/callouts
    
    TEXT_DARK = RGBColor(26, 32, 44)         # #1A202C Executive Dark
    TEXT_MUTED = RGBColor(100, 116, 139)     # Slate Muted Subtitle
    TEXT_LIGHT = RGBColor(248, 250, 252)     # High-contrast Light
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184)
    
    TOTAL_SLIDES = 3

    def create_base_slide(is_dark=False):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG if is_dark else LIGHT_BG
        bg.line.fill.background()
        return slide

    def add_header(slide, title_text, subtitle_text):
        # Red Accent Bar
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.55), Inches(0.08), Inches(0.55))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = PRIMARY_RED
        stripe.line.fill.background()
        
        # Title
        tx_box = slide.shapes.add_textbox(Inches(1.05), Inches(0.48), Inches(11.4), Inches(0.45))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        # Subtitle
        tx_sub = slide.shapes.add_textbox(Inches(1.05), Inches(0.98), Inches(11.4), Inches(0.35))
        tf_sub = tx_sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = "Calibri"
        p_sub.font.size = Pt(12)
        p_sub.font.italic = True
        p_sub.font.color.rgb = TEXT_MUTED

    def add_footer(slide, slide_num):
        # Footer Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_GRAY
        line.line.fill.background()
        
        # Left Text
        tx_l = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(8.5), Inches(0.3))
        tf_l = tx_l.text_frame
        tf_l.margin_left = tf_l.margin_top = tf_l.margin_bottom = tf_l.margin_right = 0
        p_l = tf_l.paragraphs[0]
        p_l.text = "Lumbini Elite Solutions  |  CEO Strategic Briefing  |  Confidential"
        p_l.font.name = "Calibri"
        p_l.font.size = Pt(9)
        p_l.font.color.rgb = TEXT_MUTED
        
        # Right Page Number
        tx_r = slide.shapes.add_textbox(Inches(11.033), Inches(6.95), Inches(1.5), Inches(0.3))
        tf_r = tx_r.text_frame
        tf_r.margin_left = tf_r.margin_top = tf_r.margin_bottom = tf_r.margin_right = 0
        p_r = tf_r.paragraphs[0]
        p_r.text = f"Slide {slide_num} of {TOTAL_SLIDES}"
        p_r.font.name = "Calibri"
        p_r.font.size = Pt(9)
        p_r.font.color.rgb = TEXT_MUTED
        p_r.alignment = PP_ALIGN.RIGHT

    # =========================================================================
    # SLIDE 1: Executive Vision & Strategic Resilience (Dark Theme)
    # =========================================================================
    s1 = create_base_slide(is_dark=True)
    
    # Top Tag Pill
    tag_bg = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.15), Inches(3.8), Inches(0.38))
    tag_bg.fill.solid()
    tag_bg.fill.fore_color.rgb = DARK_CARD
    tag_bg.line.color.rgb = PRIMARY_RED
    tag_bg.line.width = Pt(1)
    
    tx_tag = s1.shapes.add_textbox(Inches(1.1), Inches(1.22), Inches(3.6), Inches(0.3))
    tf_tag = tx_tag.text_frame
    tf_tag.margin_top = tf_tag.margin_left = tf_tag.margin_bottom = tf_tag.margin_right = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "CEO STRATEGIC BRIEFING"
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = PRIMARY_RED
    
    # Main Title
    tx_title = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(1.6))
    tf_title = tx_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_top = tf_title.margin_left = tf_title.margin_bottom = tf_title.margin_right = 0
    p1 = tf_title.paragraphs[0]
    p1.text = "Architecting Enterprise Resilience"
    p1.font.name = "Arial"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.space_after = Pt(8)
    
    p2 = tf_title.add_paragraph()
    p2.text = "Security, Scale, and Strategic Value Realization"
    p2.font.name = "Arial"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY_RED
    
    # Red Horizontal Accent Line
    div_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(3.65), Inches(2.2), Inches(0.04))
    div_line.fill.solid()
    div_line.fill.fore_color.rgb = PRIMARY_RED
    div_line.line.fill.background()
    
    # Executive Positioning Card
    exec_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.9), Inches(11.3), Inches(0.98))
    exec_card.fill.solid()
    exec_card.fill.fore_color.rgb = DARK_CARD
    exec_card.line.color.rgb = DARK_BORDER
    exec_card.line.width = Pt(1)
    
    tx_exec = s1.shapes.add_textbox(Inches(1.25), Inches(4.02), Inches(10.8), Inches(0.75))
    tf_exec = tx_exec.text_frame
    tf_exec.word_wrap = True
    tf_exec.margin_top = tf_exec.margin_left = tf_exec.margin_bottom = tf_exec.margin_right = 0
    
    p_ex1 = tf_exec.paragraphs[0]
    p_ex1.text = "Partnering with Lumbini Elite to protect multi-billion-dollar assets and future-proof digital transformations."
    p_ex1.font.name = "Arial"
    p_ex1.font.size = Pt(13)
    p_ex1.font.bold = True
    p_ex1.font.color.rgb = TEXT_LIGHT
    p_ex1.space_after = Pt(4)
    
    p_ex2 = tf_exec.add_paragraph()
    p_ex2.text = "Executive Briefing presented by Hemanth Kumar, Principal Consultant  |  Lumbini Elite Solutions"
    p_ex2.font.name = "Calibri"
    p_ex2.font.size = Pt(11)
    p_ex2.font.color.rgb = TEXT_LIGHT_MUTED
    
    # Bottom 3 Executive Value Containers
    pills = [
        ("Institutional Protection", "Safeguarding core business continuity across multi-billion-dollar operational assets."),
        ("Transformation Velocity", "De-risking Cloud & AI adoption while driving down operational friction and overhead."),
        ("Audit Predictability", "Guaranteeing zero-defect regulatory governance under continuous global oversight.")
    ]
    card_w = 3.55
    gap = 0.32
    start_x = 1.0
    for i, (p_title, p_desc) in enumerate(pills):
        x = start_x + i * (card_w + gap)
        c_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.2), Inches(card_w), Inches(1.4))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = DARK_CARD
        c_shape.line.color.rgb = DARK_BORDER
        c_shape.line.width = Pt(1)
        
        tx_c = s1.shapes.add_textbox(Inches(x + 0.2), Inches(5.35), Inches(card_w - 0.4), Inches(1.1))
        tf_c = tx_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_top = tf_c.margin_left = tf_c.margin_bottom = tf_c.margin_right = 0
        
        pc1 = tf_c.paragraphs[0]
        pc1.text = p_title
        pc1.font.name = "Arial"
        pc1.font.size = Pt(12.5)
        pc1.font.bold = True
        pc1.font.color.rgb = PRIMARY_RED
        pc1.space_after = Pt(4)
        
        pc2 = tf_c.add_paragraph()
        pc2.text = p_desc
        pc2.font.name = "Calibri"
        pc2.font.size = Pt(10.5)
        pc2.font.color.rgb = TEXT_LIGHT_MUTED
        
    s1.notes_slide.notes_text_frame.text = (
        "SLIDE 1: EXECUTIVE VISION & STRATEGIC RESILIENCE\n\n"
        "Talking Points for Hemanth Kumar (Spoken to CEO):\n"
        "- Opening Hook: 'As CEO, your imperative is balancing strategic velocity with absolute institutional protection. When undergoing core digital transformation, governance cannot become an operational brake.'\n"
        "- Core Positioning: 'At Lumbini Elite, we partner with C-suites to turn security and compliance into an architectural accelerator—protecting your multi-billion-dollar ERP assets while de-risking Cloud and AI value realization.'\n"
        "- Dialogue Focus: 'Today's brief discussion centers on three CEO priorities: quantifiable financial efficiency, total audit predictability, and a seamless growth path into Clean Core and autonomous AI.'"
    )

    # =========================================================================
    # SLIDE 2: Bottom-Line Impact & Risk Mitigation (Light Theme)
    # =========================================================================
    s2 = create_base_slide(is_dark=False)
    add_header(
        s2,
        "Proven Financial and Operational ROI at Scale",
        "Delivering quantifiable business continuity, cost reduction, and compliance certainty across global operations."
    )
    
    # 3 High-Impact Executive Metric Containers
    ceo_metrics = [
        {
            "stat": "60%",
            "pillar": "Operational Efficiency",
            "context": "Maintenance Overhead Reduction",
            "desc": "Slashed recurring administrative overhead across 20,000+ global users via enterprise role rationalization, composite pruning, and automated re-certification.",
            "roi": "Bottom-Line ROI: Millions saved in recurring support and licensing waste."
        },
        {
            "stat": "< 2 Hours",
            "pillar": "Speed-to-Market & Velocity",
            "context": "Access & SLA Compression",
            "desc": "Compressed critical system access and firefighter elevation provisioning turnaround from 5 business days down to under 2 hours.",
            "roi": "Operational ROI: Zero commercial project delays or developer idle time."
        },
        {
            "stat": "100%",
            "pillar": "Audit Certainty & Risk",
            "context": "Compliance & Penalty Elimination",
            "desc": "Unblemished 100% compliance pass rate under SOX-404, ISO-27001, and strict international intellectual property export regulations.",
            "roi": "Governance ROI: Complete mitigation of multi-million-dollar penalty risks."
        }
    ]
    
    col_w = 3.65
    col_gap = 0.39
    left_start = 0.8
    card_y = 1.55
    card_h = 4.2
    
    for i, m_item in enumerate(ceo_metrics):
        x = left_start + i * (col_w + col_gap)
        
        c_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(card_y), Inches(col_w), Inches(card_h))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = BORDER_GRAY
        c_shape.line.width = Pt(1)
        
        # Red Header Strip
        top_bar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(card_y), Inches(col_w), Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = PRIMARY_RED
        top_bar.line.fill.background()
        
        tx_box = s2.shapes.add_textbox(Inches(x + 0.25), Inches(card_y + 0.22), Inches(col_w - 0.5), Inches(card_h - 0.44))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_bottom = tf.margin_right = 0
        
        # Big Stat
        p_stat = tf.paragraphs[0]
        p_stat.text = m_item["stat"]
        p_stat.font.name = "Arial"
        p_stat.font.size = Pt(30)
        p_stat.font.bold = True
        p_stat.font.color.rgb = PRIMARY_RED
        p_stat.space_after = Pt(3)
        
        # Pillar
        p_pil = tf.add_paragraph()
        p_pil.text = m_item["pillar"]
        p_pil.font.name = "Arial"
        p_pil.font.size = Pt(12.5)
        p_pil.font.bold = True
        p_pil.font.color.rgb = TEXT_DARK
        
        # Context Subtitle
        p_ctx = tf.add_paragraph()
        p_ctx.text = m_item["context"]
        p_ctx.font.name = "Calibri"
        p_ctx.font.size = Pt(10)
        p_ctx.font.bold = True
        p_ctx.font.color.rgb = PRIMARY_RED
        p_ctx.space_after = Pt(10)
        
        # Description
        p_desc = tf.add_paragraph()
        p_desc.text = m_item["desc"]
        p_desc.font.name = "Calibri"
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_DARK
        p_desc.space_after = Pt(12)
        
        # ROI Callout Box within card
        p_roi = tf.add_paragraph()
        p_roi.text = m_item["roi"]
        p_roi.font.name = "Calibri"
        p_roi.font.size = Pt(9.5)
        p_roi.font.italic = True
        p_roi.font.bold = True
        p_roi.font.color.rgb = PRIMARY_RED
        
    # Bottom Executive Takeaway Banner
    takeaway_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.75))
    takeaway_shape.fill.solid()
    takeaway_shape.fill.fore_color.rgb = ACCENT_RED_BG
    takeaway_shape.line.color.rgb = PRIMARY_RED
    takeaway_shape.line.width = Pt(1)
    
    tx_tk = s2.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.55))
    tf_tk = tx_tk.text_frame
    tf_tk.word_wrap = True
    tf_tk.margin_top = tf_tk.margin_left = tf_tk.margin_bottom = tf_tk.margin_right = 0
    ptk = tf_tk.paragraphs[0]
    ptk.text = "CEO Value Realization: "
    ptk.font.name = "Arial"
    ptk.font.size = Pt(11)
    ptk.font.bold = True
    ptk.font.color.rgb = PRIMARY_RED
    
    run_tk = ptk.add_run()
    run_tk.text = "Every 1% reduction in manual authorization overhead unlocks enterprise agility. We convert regulatory friction into automated operating leverage that flows directly to the bottom line."
    run_tk.font.name = "Calibri"
    run_tk.font.size = Pt(10.5)
    run_tk.font.bold = False
    run_tk.font.color.rgb = TEXT_DARK
    
    add_footer(s2, 2)
    
    s2.notes_slide.notes_text_frame.text = (
        "SLIDE 2: BOTTOM-LINE IMPACT & RISK MITIGATION\n\n"
        "Talking Points for Hemanth Kumar (Spoken to CEO):\n"
        "- Operational Efficiency: 'We've proven across 20,000+ global users that rationalizing security architecture drops annual maintenance overhead by 60%. This isn't just an IT metric; it directly recovers millions in wasted engineering hours.'\n"
        "- Strategic Velocity: 'In fast-moving enterprises, waiting 5 days for critical project or firefighter access stalls commercial initiatives. By compressing SLAs to under 2 hours, your teams operate at full market speed.'\n"
        "- Audit Assurance: 'Under rigorous SOX, ISO, and IP oversight, our client landscapes hold a 100% first-time audit pass rate. We completely eliminate recurring audit findings and regulatory fines.'"
    )

    # =========================================================================
    # SLIDE 3: The Strategic Growth Partnership (Light Theme)
    # =========================================================================
    s3 = create_base_slide(is_dark=False)
    add_header(
        s3,
        "De-Risking the Future: Cloud, IAG, and AI Readiness",
        "Aligning enterprise governance with long-term digital growth and Clean Core agility."
    )
    
    growth_pillars = [
        {
            "tag": "ASSET PROTECTION",
            "title": "Clean Core Protection",
            "sub": "Zero-Modification Standard",
            "bullets": [
                "Zero core ABAP modifications during S/4HANA transitions.",
                "Prevents custom code bloat from creating technical debt.",
                "Ensures seamless future upgrade cycles without business downtime."
            ]
        },
        {
            "tag": "TIME-TO-VALUE",
            "title": "Accelerated IAG Onboarding",
            "sub": "40% Faster Cloud Migration",
            "bullets": [
                "Proprietary pre-cloud remediation toolkits scrub legacy role catalogs.",
                "Automates SoD rationalization prior to cloud transition.",
                "Cuts customer cloud onboarding and deployment timelines by up to 40%."
            ]
        },
        {
            "tag": "STRATEGIC CALL TO ACTION",
            "title": "Executive Alignment Pilot",
            "sub": "30-Day Proof of Value",
            "bullets": [
                "Streamlined 30-day executive alignment and targeted lighthouse pilot.",
                "Validates quantifiable security and cost benefits with zero risk.",
                "Produces a clear Board-ready ROI roadmap for full-scale rollout."
            ]
        }
    ]
    
    for i, gp in enumerate(growth_pillars):
        x = left_start + i * (col_w + col_gap)
        
        c_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(card_y), Inches(col_w), Inches(3.7))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = BORDER_GRAY
        c_shape.line.width = Pt(1)
        
        # Pill Tag
        pill = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.25), Inches(card_y + 0.25), Inches(1.8), Inches(0.3))
        pill.fill.solid()
        pill.fill.fore_color.rgb = ACCENT_RED_BG
        pill.line.color.rgb = PRIMARY_RED
        pill.line.width = Pt(1)
        
        tx_pill = s3.shapes.add_textbox(Inches(x + 0.25), Inches(card_y + 0.28), Inches(1.8), Inches(0.25))
        tf_pill = tx_pill.text_frame
        tf_pill.margin_top = tf_pill.margin_left = tf_pill.margin_bottom = tf_pill.margin_right = 0
        p_p = tf_pill.paragraphs[0]
        p_p.text = gp["tag"]
        p_p.font.name = "Arial"
        p_p.font.size = Pt(8.5)
        p_p.font.bold = True
        p_p.font.color.rgb = PRIMARY_RED
        p_p.alignment = PP_ALIGN.CENTER
        
        # Card Body
        tx_box = s3.shapes.add_textbox(Inches(x + 0.25), Inches(card_y + 0.72), Inches(col_w - 0.5), Inches(3.7 - 0.85))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_bottom = tf.margin_right = 0
        
        p_title = tf.paragraphs[0]
        p_title.text = gp["title"]
        p_title.font.name = "Arial"
        p_title.font.size = Pt(13)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.space_after = Pt(2)
        
        p_sub = tf.add_paragraph()
        p_sub.text = gp["sub"]
        p_sub.font.name = "Calibri"
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = PRIMARY_RED
        p_sub.space_after = Pt(12)
        
        for item in gp["bullets"]:
            pi = tf.add_paragraph()
            pi.text = "-  " + item
            pi.font.name = "Calibri"
            pi.font.size = Pt(10)
            pi.font.color.rgb = TEXT_DARK
            pi.space_after = Pt(7)
            
    # Bottom Action Callout Container: Proposed Executive Next Step
    act_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.5), Inches(11.733), Inches(1.15))
    act_shape.fill.solid()
    act_shape.fill.fore_color.rgb = CARD_BG
    act_shape.line.color.rgb = PRIMARY_RED
    act_shape.line.width = Pt(1.5)
    
    tx_act = s3.shapes.add_textbox(Inches(1.1), Inches(5.62), Inches(11.1), Inches(0.9))
    tf_act = tx_act.text_frame
    tf_act.word_wrap = True
    tf_act.margin_top = tf_act.margin_left = tf_act.margin_bottom = tf_act.margin_right = 0
    
    p_act1 = tf_act.paragraphs[0]
    p_act1.text = "Proposed Executive Next Step: 30-Day Lighthouse Pilot"
    p_act1.font.name = "Arial"
    p_act1.font.size = Pt(12)
    p_act1.font.bold = True
    p_act1.font.color.rgb = PRIMARY_RED
    p_act1.space_after = Pt(3)
    
    p_act2 = tf_act.add_paragraph()
    p_act2.text = "Initiate a low-friction 30-day alignment on a targeted business unit to deploy our pre-cloud remediation toolkits, prove 40% onboarding compression, and present a definitive ROI business case to your executive committee."
    p_act2.font.name = "Calibri"
    p_act2.font.size = Pt(11)
    p_act2.font.color.rgb = TEXT_DARK
    
    add_footer(s3, 3)
    
    s3.notes_slide.notes_text_frame.text = (
        "SLIDE 3: THE STRATEGIC GROWTH PARTNERSHIP\n\n"
        "Talking Points for Hemanth Kumar (Spoken to CEO):\n"
        "- Clean Core Agility: 'Digital transformations often fail because legacy ERP systems get bogged down with custom ABAP modifications. We preserve a pure Clean Core on SAP BTP, guaranteeing upgrade agility and business continuity.'\n"
        "- Accelerated Time-to-Value: 'Migrating to Cloud IAG shouldn't take 12 months of manual triage. Our pre-cloud remediation accelerators clean the slate first, compressing your transition timeline by 40%.'\n"
        "- The Ask / Next Step: 'We propose a low-risk, 30-day lighthouse pilot on a single critical business unit. We will prove the 40% acceleration and quantifiable cost savings firsthand, giving you concrete data before committing broader enterprise resources.'"
    )

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_presentation()
